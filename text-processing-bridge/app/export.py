"""
処理結果を配布用ファイル(txt・md・csv・docx・xlsx)のバイト列へ変換する。

Open WebUIのPipe/Toolは管理画面へ1ファイル貼り付けで配布され、テストを
書けない。ファイル生成のロジックはここへ集約し、Pipe側は
/v1/render/export を呼んでバイト列を受け取るだけにする。
"""

from __future__ import annotations

import csv
import io
import re

FORMAT_MIME = {
    "txt": "text/plain; charset=utf-8",
    "md": "text/markdown; charset=utf-8",
    "docx": (
        "application/vnd.openxmlformats-officedocument"
        ".wordprocessingml.document"
    ),
    "csv": "text/csv; charset=utf-8",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}


def _strip_markdown(value: str) -> str:
    value = re.sub(r"\*\*(.+?)\*\*", r"\1", value)
    value = re.sub(r"__(.+?)__", r"\1", value)
    value = re.sub(r"`(.+?)`", r"\1", value)
    return value


def _protect_cell(value: str) -> str:
    # 表計算ソフトが数式として解釈しないようにする。
    return "'" + value if value.startswith(("=", "+", "-", "@")) else value


def _to_rows(content: str) -> list[list[str]]:
    """本文を行単位の表へ落とす。空行は行番号を進めずに詰める。"""
    lines = [
        line.rstrip()
        for line in content.replace("\r\n", "\n").replace("\r", "\n").split("\n")
    ]
    return [[str(i), line] for i, line in enumerate(
        [l for l in lines if l.strip()], start=1
    )]


def build_docx(content: str, title: str = "") -> bytes:
    from docx import Document
    from docx.shared import Pt

    document = Document()
    normal = document.styles["Normal"]
    normal.font.name = "Yu Gothic"
    normal.font.size = Pt(10.5)

    if title.strip():
        document.add_heading(_strip_markdown(title.strip()), level=0)

    for line in content.replace("\r\n", "\n").replace("\r", "\n").split("\n"):
        heading = re.match(r"^(#{1,6})\s+(.+)$", line)
        bold_heading = re.match(r"^\*\*(.+?)\*\*$", line.strip())
        bullet = re.match(r"^\s*[-*+]\s+(.+)$", line)
        numbered = re.match(r"^\s*\d+[.)]\s+(.+)$", line)

        if heading:
            document.add_heading(
                _strip_markdown(heading.group(2)),
                level=min(len(heading.group(1)), 6),
            )
        elif bold_heading:
            document.add_heading(_strip_markdown(bold_heading.group(1)), level=1)
        elif bullet:
            document.add_paragraph(
                _strip_markdown(bullet.group(1)), style="List Bullet"
            )
        elif numbered:
            document.add_paragraph(
                _strip_markdown(numbered.group(1)), style="List Number"
            )
        else:
            document.add_paragraph(_strip_markdown(line))

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def build_csv(content: str) -> bytes:
    buffer = io.StringIO(newline="")
    writer = csv.writer(buffer, lineterminator="\r\n")
    writer.writerow(["行", "内容"])
    for number, line in _to_rows(content):
        writer.writerow([number, _protect_cell(line)])
    return buffer.getvalue().encode("utf-8-sig")


def build_xlsx(content: str, sheet_name: str) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = (re.sub(r"[\\/*?:\[\]]", "_", sheet_name or "結果") or "結果")[:31]
    sheet.freeze_panes = "A2"
    sheet.append(["行", "内容"])
    for cell in sheet[1]:
        cell.font = Font(bold=True)

    rows = _to_rows(content)
    for number, line in rows:
        sheet.append([int(number), _protect_cell(line)])

    sheet.column_dimensions[get_column_letter(1)].width = 6
    longest = max((len(line) for _, line in rows), default=20)
    sheet.column_dimensions[get_column_letter(2)].width = min(
        max(longest + 2, 20), 100
    )
    for row in sheet.iter_rows(min_row=2, min_col=2, max_col=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def render_output(output_format: str, content: str, title: str) -> tuple[bytes, str]:
    """指定形式のバイト列とMIMEタイプを返す。"""
    if output_format == "docx":
        return build_docx(content, title), FORMAT_MIME["docx"]
    if output_format == "xlsx":
        return build_xlsx(content, title), FORMAT_MIME["xlsx"]
    if output_format == "csv":
        return build_csv(content), FORMAT_MIME["csv"]
    if output_format == "md":
        return content.encode("utf-8"), FORMAT_MIME["md"]
    return content.encode("utf-8-sig"), FORMAT_MIME["txt"]
