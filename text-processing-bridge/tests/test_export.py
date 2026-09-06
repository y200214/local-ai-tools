from __future__ import annotations

import base64
import io

from app.export import (
    FORMAT_MIME,
    _protect_cell,
    _to_rows,
    build_csv,
    build_docx,
    build_xlsx,
    render_output,
)


def test_render_output_txt_has_bom() -> None:
    # メモ帳やWordで文字化けしないよう、txtはBOM付きUTF-8で出す。
    data, mime = render_output("txt", "こんにちは", "")

    assert data.startswith(b"\xef\xbb\xbf")
    assert mime == FORMAT_MIME["txt"]


def test_render_output_md_is_plain_utf8() -> None:
    data, mime = render_output("md", "# 見出し", "")

    assert data == "# 見出し".encode("utf-8")
    assert mime == FORMAT_MIME["md"]


def test_to_rows_skips_blank_lines_without_gaps() -> None:
    rows = _to_rows("1行目\n\n2行目\r\n3行目")

    assert rows == [["1", "1行目"], ["2", "2行目"], ["3", "3行目"]]


def test_protect_cell_escapes_formula_prefix() -> None:
    assert _protect_cell("=SUM(A1)") == "'=SUM(A1)"
    assert _protect_cell("通常の文") == "通常の文"


def test_build_csv_has_header_and_rows() -> None:
    text = build_csv("要点1\n要点2").decode("utf-8-sig")

    lines = text.strip().split("\r\n")
    assert lines[0] == "行,内容"
    assert lines[1] == "1,要点1"
    assert lines[2] == "2,要点2"


def test_build_xlsx_roundtrip() -> None:
    from openpyxl import load_workbook

    data = build_xlsx("要点1\n要点2", "テスト結果")
    sheet = load_workbook(io.BytesIO(data)).active

    assert sheet.title == "テスト結果"
    assert sheet["A1"].value == "行"
    assert sheet["B2"].value == "要点1"
    assert sheet["B3"].value == "要点2"


def test_build_docx_roundtrip_with_headings_and_bullets() -> None:
    from docx import Document

    data = build_docx("# 概要\n- 箇条書き1\n本文です", title="要約 2026")
    document = Document(io.BytesIO(data))
    texts = [p.text for p in document.paragraphs]

    assert "要約 2026" in texts
    assert "概要" in texts
    assert "箇条書き1" in texts
    assert "本文です" in texts


def test_export_endpoint_returns_decodable_base64(api_client) -> None:
    response = api_client.post(
        "/v1/render/export",
        json={"content": "本文", "format": "txt", "title": ""},
    )

    assert response.status_code == 200
    body = response.json()
    assert body["mime"] == FORMAT_MIME["txt"]
    decoded = base64.b64decode(body["data_b64"])
    assert decoded.decode("utf-8-sig") == "本文"


def test_export_endpoint_rejects_unknown_format(api_client) -> None:
    response = api_client.post(
        "/v1/render/export",
        json={"content": "本文", "format": "pdf"},
    )

    assert response.status_code == 422
