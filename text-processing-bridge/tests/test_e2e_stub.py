"""
LLMスタブによる通しテスト(文字起こし→構造化→docx描画)。

実LLMへ繋がずに、構造化APIの出力JSONをそのまま描画APIへ渡し、
生成docxが開ける・議題と発言が残る・数値が改変されない・
出典マーカーが塗られる、までを一気に保証する。

スタブの回答はLLMの契約(instructions.py: 話者は文字起こしに明記が
ある場合だけ抽出し、推測しない)に従う内容にする。話者が明記された
文字起こしには話者付きの回答を、明記の無い文字起こしには話者空欄の
回答を対にし、契約違反の出力例をテストへ持ち込まない。

実テンプレート(templates/)は機密で読めないため、render_minutes が
前提とする行構成(本文表=tables[1]・21行・本文セル3段落)だけを
python-docx で合成して使う。テンプレート仕様が変わったときは
このテストの _build_template_bytes も合わせて直すこと。
"""

from __future__ import annotations

import base64
import io

from docx import Document
from docx.enum.text import WD_COLOR_INDEX

from app.local_workflow import WorkflowRunResult
from app.main import app, get_workflow_client


class FakeWorkflowClient:
    def __init__(self, results: list[WorkflowRunResult]) -> None:
        self.results = results
        self.calls: list[tuple[str, dict[str, object]]] = []

    async def run(
        self, workflow_name: str, inputs: dict[str, object]
    ) -> WorkflowRunResult:
        self.calls.append((workflow_name, inputs))
        return self.results[len(self.calls) - 1]


# 架空の次第と文字起こし。実在の会議データはテストに使わない(data/は触らない)。
SAMPLE_AGENDA = """
令和8年度 第1回 経営企画会議 次第
日 時：令和8年6月9日（火）17:00～
場 所：大会議室
議 題
1 病院の経営状況について
① 月次収支の報告 ・・・・ 資料1
経営企画課 山田 課長補佐
"""

ROSTER = "出席者：乙野教授、山田課長補佐"

# 話者が明記された文字起こし。LLMが話者を「抽出」できる唯一の形。
TRANSCRIPT_WITH_SPEAKERS = (
    "それでは山田課長補佐から資料1の説明をお願いします。\n"
    "乙野先生：病床稼働率が目標の93％を下回っているのではないでしょうか。\n"
    "山田課長補佐：6月の稼働率は92.5％でした。来週の経営会議までに詳細を共有します。"
)

# 話者の明記が無い文字起こし。ここから話者を当てるのは捏造になる。
TRANSCRIPT_WITHOUT_SPEAKERS = (
    "それでは資料1の説明をお願いします。"
    "病床稼働率が目標の93％を下回っているのではないでしょうか。"
    "6月の稼働率は92.5％でした。来週の経営会議までに詳細を共有します。"
)


def _build_template_bytes() -> bytes:
    """render_minutes の前提行構成だけを持つ合成テンプレートを作る。"""
    document = Document()
    document.add_paragraph("○○定例会議 議事録")
    document.add_paragraph("日時：")
    document.add_paragraph("場所：")
    document.add_paragraph("出席者：")
    document.add_table(rows=1, cols=1)  # tables[0] はヘッダ表(描画では使わない)
    body = document.add_table(rows=21, cols=2)
    cell = body.rows[1].cells[0]
    cell.add_paragraph("")  # ＜質疑応答・意見等＞見出し用の雛形
    cell.add_paragraph("")  # 発言行用の雛形
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _all_texts(document: Document) -> str:
    parts = [paragraph.text for paragraph in document.paragraphs]
    parts.append(_table_texts(document))
    return "\n".join(parts)


def _table_texts(document: Document) -> str:
    """本文表だけの文面。ヘッダの名簿(出席者：…)を含めない検証に使う。"""
    parts = []
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.extend(paragraph.text for paragraph in cell.paragraphs)
    return "\n".join(parts)


def _table_highlights(document: Document) -> set:
    highlights = set()
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        if run.font.highlight_color is not None:
                            highlights.add(run.font.highlight_color)
    return highlights


def _render_docx(api_client, document_data: dict) -> Document:
    rendered = api_client.post(
        "/v1/render/minutes_docx",
        json={
            "document": document_data,
            "template_b64": base64.b64encode(_build_template_bytes()).decode(
                "ascii"
            ),
        },
    )
    assert rendered.status_code == 200
    docx_bytes = base64.b64decode(rendered.json()["data_b64"])
    # 生成物がdocxとして開けること自体も検証のうち
    return Document(io.BytesIO(docx_bytes))


def test_stub_e2e_explicit_speakers_are_normalized(api_client) -> None:
    """明示話者あり: LLMが抽出した呼称が名簿表記へ揃い、緑マーカーで描かれる。"""
    fake = FakeWorkflowClient(
        [
            # 1周目: 発言の振り分け。話者は文字起こしに明記された表記のまま返る
            WorkflowRunResult(
                text=(
                    # 数値の羅列だけの断定調発言は「報告の読み上げ」として
                    # 落とされる仕様(merge.looks_like_report)のため、
                    # 質問→数値入り回答の対にして「回答の数値は残る」経路を固定する
                    "##項目 1-1\n"
                    "##報告 山田課長補佐\n"
                    "##資料 資料1\n"
                    "##発言 乙野先生 | 病床稼働率が目標の93％を下回っているのではないか。\n"
                    "##発言 山田課長補佐 | 6月の稼働率は92.5％だった。来週の経営会議までに詳細を共有する。\n"
                ),
                workflow_run_id="run-1",
            ),
            # 2周目(強制チェック): 追加なし
            WorkflowRunResult(text="", workflow_run_id="run-2"),
        ]
    )
    app.dependency_overrides[get_workflow_client] = lambda: fake

    structured = api_client.post(
        "/v1/text/minutes_structured",
        json={
            "text": TRANSCRIPT_WITH_SPEAKERS,
            "agenda_text": SAMPLE_AGENDA,
            "roster_text": ROSTER,
        },
    )
    assert structured.status_code == 200

    result = _render_docx(api_client, structured.json()["document"])
    text = _all_texts(result)

    # 次第由来の議題が見出しとして流れている
    assert "病院の経営状況について" in text
    # 発言者の呼称が名簿の表記(乙野先生→乙野教授)へ寄っている
    assert "乙野教授" in _table_texts(result)
    # 発言中の数値が欠落・改変されていない
    assert "92.5" in text
    assert "93" in text
    # 出典マーカー: 文字起こしから特定できた話者は緑
    assert WD_COLOR_INDEX.BRIGHT_GREEN in _table_highlights(result)


def test_stub_e2e_unknown_speakers_stay_blank(api_client) -> None:
    """明示話者なし: 名簿があっても話者を捏造せず、空欄(赤マーカー)で描かれる。"""
    fake = FakeWorkflowClient(
        [
            # 話者の明記が無いので、契約に従い話者は空欄で返る
            WorkflowRunResult(
                text=(
                    "##項目 1-1\n"
                    "##発言  | 病床稼働率が目標の93％を下回っているのではないか。\n"
                    "##発言  | 6月の稼働率は92.5％だった。来週の経営会議までに詳細を共有する。\n"
                ),
                workflow_run_id="run-1",
            ),
            WorkflowRunResult(text="", workflow_run_id="run-2"),
        ]
    )
    app.dependency_overrides[get_workflow_client] = lambda: fake

    structured = api_client.post(
        "/v1/text/minutes_structured",
        json={
            "text": TRANSCRIPT_WITHOUT_SPEAKERS,
            "agenda_text": SAMPLE_AGENDA,
            "roster_text": ROSTER,
        },
    )
    assert structured.status_code == 200

    result = _render_docx(api_client, structured.json()["document"])
    body_text = _table_texts(result)

    # 名簿の名前が本文の発言者として湧いて出ないこと
    # (ヘッダの出席者欄には名簿として載るため、本文表だけを見る)
    assert "乙野" not in body_text
    # 発言そのものと数値は残っている
    assert "92.5" in body_text
    # 不明話者は空欄+赤マーカー(手入力の目印)で描かれる
    assert WD_COLOR_INDEX.RED in _table_highlights(result)
