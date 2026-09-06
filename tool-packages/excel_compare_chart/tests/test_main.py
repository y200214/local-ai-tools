"""Excel比較グラフの単体テスト。

**実物の帳票の作りに合わせて確かめる**(2026-08-31に実物で作り直した)。
見出しが1行目にない・1枚に表が複数ある・項目名が複数列に分かれている、
という形を架空データで再現している。実データは使わない。

「推測しないこと」と「聞かれたことに具体的に答えること」を中心に見る。
"""

from __future__ import annotations

import pytest
from openpyxl import Workbook

import main


def real_shaped_book() -> Workbook:
    """実物と同じ作りの架空ブック。

    ・1〜9行目は表題や注意書き(数値なし)
    ・11行目が見出し、H〜K列が年度
    ・項目名はB列(大項目)とC列(内訳)に分かれる
    ・24行目から2つ目の表
    """
    book = Workbook()
    sheet = book.active
    sheet.title = "科別（案）"
    sheet["B2"] = "令和８年度 第１四半期(架空)"
    sheet["B4"] = "診療科名 ："

    def section(start: int, title: str, rows: list[tuple[str, str, list[int]]]) -> None:
        sheet.cell(row=start, column=2, value=title)
        sheet.cell(row=start + 1, column=2, value="項目")
        for offset, name in enumerate(["H30年度", "R6年度", "R7年度", "R8年度"]):
            sheet.cell(row=start + 1, column=8 + offset, value=name)
        for index, (major, minor, values) in enumerate(rows):
            row = start + 2 + index
            if major:
                sheet.cell(row=row, column=2, value=major)
            if minor:
                sheet.cell(row=row, column=3, value=minor)
            for offset, value in enumerate(values):
                sheet.cell(row=row, column=8 + offset, value=value)

    section(10, "(1)入院関係の実績", [
        ("①収入（百万円）", "", [1070, 1270, 1335, 1395]),
        ("", "1 手技料収入", [640, 760, 800, 835]),
        ("", "2 薬剤収入", [260, 315, 330, 345]),
        ("②単価（円）", "", [82700, 89400, 91800, 94000]),
    ])
    section(24, "(2)外来関係の実績", [
        ("①収入（百万円）", "", [970, 1140, 1205, 1255]),
        ("", "1 手技料収入", [610, 710, 750, 780]),
        ("②単価（円）", "", [27800, 31000, 31800, 32500]),
    ])
    return book


# ---------------------------------------------------------------------------
# 実物の作りを読めるか
# ---------------------------------------------------------------------------
def test_見出しが1行目でなくても表を見つける() -> None:
    sheet = real_shaped_book().active
    blocks = main.find_blocks(sheet)
    assert len(blocks) == 2
    assert blocks[0].header_row == 11  # 1行目ではない
    assert [main._column_letter(c) for c in blocks[0].columns] == ["H", "I", "J", "K"]


def test_表の名前を拾う() -> None:
    blocks = main.find_blocks(real_shaped_book().active)
    assert blocks[0].title == "(1)入院関係の実績"
    assert blocks[1].title == "(2)外来関係の実績"


def test_見出し名を読む() -> None:
    sheet = real_shaped_book().active
    block = main.find_blocks(sheet)[0]
    assert block.header(sheet, 8) == "H30年度"
    assert "H列「H30年度」" in block.choices(sheet)


def test_複数列にまたがる項目名をつなぐ() -> None:
    sheet = real_shaped_book().active
    block = main.find_blocks(sheet)[0]
    labels = [label for label, _ in main.series_of(sheet, block, 8)]
    assert "①収入（百万円）" in labels
    assert "1 手技料収入" in labels  # 内訳はC列にある


def test_エラー値や単位つきの数値を読み分ける() -> None:
    assert main._as_number("#REF!") is None
    assert main._as_number("1,234") == 1234
    assert main._as_number("45件") == 45
    assert main._as_number(None) is None


# ---------------------------------------------------------------------------
# 聞かれたことに具体的に答えるか
# ---------------------------------------------------------------------------
def test_数値でない列を指定されたらそう答える() -> None:
    # 「BとC列を比較して」に対し「どの表ですか」と返さない(実際にそうなっていた)
    sheet = real_shaped_book().active
    blocks = main.find_blocks(sheet)
    with pytest.raises(main.NeedsUserInput) as caught:
        main.check_named_columns("BとC列を比較してください", blocks, sheet)
    message = str(caught.value)
    assert "数値が入っていません" in message
    assert "H列「H30年度」" in message  # 使える列を示す


def test_使える列の指定なら先に進む() -> None:
    sheet = real_shaped_book().active
    blocks = main.find_blocks(sheet)
    main.check_named_columns("H列とI列を比較して", blocks, sheet)  # 例外が出ない


def test_表が複数なら見本つきで聞き返す() -> None:
    sheet = real_shaped_book().active
    blocks = main.find_blocks(sheet)
    with pytest.raises(main.NeedsUserInput) as caught:
        main.pick_block("グラフにして", blocks, sheet)
    message = str(caught.value)
    assert "(1)入院関係の実績" in message and "(2)外来関係の実績" in message
    assert "書き方の例" in message  # 行き止まりにしない
    assert "H列とI列" in message


def test_表を名前で指定できる() -> None:
    sheet = real_shaped_book().active
    blocks = main.find_blocks(sheet)
    block = main.pick_block("『(2)外来関係の実績』のH列とI列を比較して", blocks, sheet)
    assert block.title == "(2)外来関係の実績"


def test_表が1つなら聞き返さない() -> None:
    book = Workbook()
    sheet = book.active
    sheet["B1"] = "項目"
    sheet["C1"] = "前年"
    sheet["D1"] = "今年"
    for index, (name, before, after) in enumerate(
        [("内科", 10, 12), ("外科", 20, 18)], start=2
    ):
        sheet.cell(row=index, column=2, value=name)
        sheet.cell(row=index, column=3, value=before)
        sheet.cell(row=index, column=4, value=after)
    blocks = main.find_blocks(sheet)
    assert len(blocks) == 1
    assert main.pick_columns("比較して", blocks[0], sheet) == [3, 4]


def test_列を見出し名で指定できる() -> None:
    sheet = real_shaped_book().active
    block = main.find_blocks(sheet)[0]
    assert main.pick_columns("『H30年度』と『R8年度』を比べて", block, sheet) == [8, 11]


def test_数値列が多くて指定が無ければ選択肢を出す() -> None:
    sheet = real_shaped_book().active
    block = main.find_blocks(sheet)[0]
    with pytest.raises(main.NeedsUserInput) as caught:
        main.pick_columns("比較して", block, sheet)
    assert "H列「H30年度」" in str(caught.value)


# ---------------------------------------------------------------------------
# 突き合わせと出力
# ---------------------------------------------------------------------------
def test_共通する項目だけを並べ食い違いを報告する() -> None:
    labels, left, right, mismatches = main.align(
        [("内科", 10.0), ("外科", 20.0), ("小児科", 30.0)],
        [("内科", 15.0), ("外科", 18.0)],
    )
    assert labels == ["内科", "外科"]
    assert left == [10.0, 20.0] and right == [15.0, 18.0]
    assert any("小児科" in note for note in mismatches)


def test_グラフを描ける(tmp_path) -> None:
    target = tmp_path / "図.png"
    main.draw_chart(["①収入（百万円）", "1 手技料収入"], [1070.0, 640.0], [1395.0, 835.0],
                    "H30年度", "R8年度", str(target))
    assert target.stat().st_size > 1000


def test_記録用Excelに元データと条件を残す(tmp_path) -> None:
    target = tmp_path / "記録.xlsx"
    main.write_record(["①収入（百万円）"], [1070.0], [1395.0], "H30年度", "R8年度",
                      "架空.xlsx", "比較して", "(1)入院関係の実績", str(target))
    from openpyxl import load_workbook

    book = load_workbook(target)
    assert book["比較データ"].cell(row=2, column=4).value == 325.0
    conditions = {row[0]: row[1] for row in book["条件"].iter_rows(values_only=True)}
    assert conditions["使った表"] == "(1)入院関係の実績"
    assert "LLM" in conditions["計算方法"]


def test_チェック行のTrueFalseで表が割れない() -> None:
    """実物の帳票には True/False が並ぶ確認行がある。

    これを「文字が並ぶ行=見出し」と数えると、1つの表が2つに割れて
    表名がどちらもシート名になり、聞き返しが役に立たなくなる(実物で発生)。
    """
    book = Workbook()
    sheet = book.active
    sheet["B1"] = "科"
    sheet["C1"] = "前年"
    sheet["D1"] = "今年"
    for index, name in enumerate(["内科", "外科", "小児科"], start=2):
        sheet.cell(row=index, column=2, value=name)
        sheet.cell(row=index, column=3, value=index * 10)
        sheet.cell(row=index, column=4, value=index * 11)
    # 確認用の行(数値ではなく真偽値が並ぶ)
    sheet.cell(row=6, column=3, value=False)
    sheet.cell(row=6, column=4, value=False)
    for index, name in enumerate(["眼科", "皮膚科"], start=8):
        sheet.cell(row=index, column=2, value=name)
        sheet.cell(row=index, column=3, value=index * 10)
        sheet.cell(row=index, column=4, value=index * 11)

    blocks = main.find_blocks(sheet)
    assert len(blocks) == 1
    assert "眼科" in [label for label, _ in main.series_of(sheet, blocks[0], 3)]


def test_真偽値やエラー値は項目名にしない() -> None:
    assert main._text_of(False) == ""
    assert main._text_of("#REF!") == ""
    assert main._text_of(" 内科 ") == "内科"


# ---------------------------------------------------------------------------
# グラフの読みやすさ
# ---------------------------------------------------------------------------
def test_項目が多いと横棒にして幅が広がりすぎない(tmp_path) -> None:
    """診療科ごとの比較は30項目を超える。縦棒のままでは画面に収まらない。"""
    import struct

    labels = [f"第{index}診療科" for index in range(30)]
    values = [float(index + 1) for index in range(30)]
    target = tmp_path / "図.png"
    main.draw_chart(labels, values, values, "R7年度", "R8年度", str(target))
    width, height = struct.unpack(">II", target.read_bytes()[16:24])
    assert width < 2000, f"横に広がりすぎている: {width}px"
    assert height > width, "項目が多いときは縦に伸ばす(横棒)"


def test_飛び抜けて大きい合計行はグラフから外す() -> None:
    rows = [("病院全体", 1000.0, 1100.0), ("内科", 10.0, 12.0), ("外科", 20.0, 18.0)]
    kept, removed = main.split_total_rows(rows)
    assert [row[0] for row in removed] == ["病院全体"]
    assert [row[0] for row in kept] == ["内科", "外科"]


def test_合計らしくても桁が近ければ外さない() -> None:
    # 「合計」でも他と並ぶ大きさなら、勝手に消さない
    rows = [("合計", 30.0, 30.0), ("内科", 20.0, 22.0), ("外科", 25.0, 24.0)]
    kept, removed = main.split_total_rows(rows)
    assert removed == [] and len(kept) == 3


def test_総合診療科を合計と見間違えない() -> None:
    assert main._looks_like_total("病院全体") is True
    assert main._looks_like_total("合計") is True
    assert main._looks_like_total("総診 総合診療科") is False
    assert main._looks_like_total("会計課") is False
