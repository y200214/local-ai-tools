"""決定事項・宿題一覧の単体テスト。

**LLMを呼ばずに、検証の部分だけを確かめる。**
このツールの肝は「LLMの答えを信用しないこと」なので、
作り話・数値の捏造・推測で埋めた担当を落とせるかを、その場で作った
架空のLLM応答で試す。
"""

from __future__ import annotations

import main

TRANSCRIPT = [
    (1, "甲野課長: 来月から発注窓口を総務課に一本化することを決定します。"),
    (2, "甲野課長: 丙野主任は、4月30日までに各部署へ新しい手順を周知してください。"),
    (3, "乙野: 様式の見直しも必要だと思います。"),
]


def test_根拠が実在すれば採用する() -> None:
    items = [{
        "kind": "決定",
        "content": "発注窓口を総務課に一本化する",
        "owner": "", "due": "",
        "evidence": "来月から発注窓口を総務課に一本化することを決定します。",
    }]
    verified, rejected = main.verify_items(items, TRANSCRIPT)
    assert rejected == []
    assert verified[0]["line"] == 1
    assert verified[0]["kind"] == "決定"


def test_文字起こしに無い発言を根拠にしたら捨てる() -> None:
    # LLMが作り話をした場合。黙って採らない
    items = [{
        "kind": "決定", "content": "予算を倍にする", "owner": "", "due": "",
        "evidence": "来年度の予算を倍額にすることで合意しました。",
    }]
    verified, rejected = main.verify_items(items, TRANSCRIPT)
    assert verified == []
    assert any("根拠の発言" in reason for reason in rejected)


def test_根拠に無い数値が入っていたら捨てる() -> None:
    items = [{
        "kind": "宿題", "content": "5月20日までに周知する", "owner": "", "due": "",
        "evidence": "丙野主任は、4月30日までに各部署へ新しい手順を周知してください。",
    }]
    verified, rejected = main.verify_items(items, TRANSCRIPT)
    assert verified == []
    assert any("数値" in reason for reason in rejected)


def test_根拠にある数値ならそのまま通す() -> None:
    items = [{
        "kind": "宿題", "content": "4月30日までに各部署へ周知する",
        "owner": "丙野主任", "due": "4月30日",
        "evidence": "丙野主任は、4月30日までに各部署へ新しい手順を周知してください。",
    }]
    verified, rejected = main.verify_items(items, TRANSCRIPT)
    assert rejected == []
    assert verified[0]["owner"] == "丙野主任"
    assert verified[0]["due"] == "4月30日"


def test_根拠に書かれていない担当は要確認にする() -> None:
    # LLMが推測で埋めた担当を、そのまま信じない
    items = [{
        "kind": "決定", "content": "発注窓口を一本化する",
        "owner": "乙野", "due": "来月中",
        "evidence": "来月から発注窓口を総務課に一本化することを決定します。",
    }]
    verified, _ = main.verify_items(items, TRANSCRIPT)
    assert verified[0]["owner"] == main.UNKNOWN
    assert verified[0]["due"] == main.UNKNOWN


def test_種別や内容が空なら捨てる() -> None:
    items = [
        {"kind": "雑談", "content": "あれこれ", "evidence": "様式の見直しも必要だと思います。"},
        {"kind": "決定", "content": "", "evidence": "様式の見直しも必要だと思います。"},
    ]
    verified, rejected = main.verify_items(items, TRANSCRIPT)
    assert verified == []
    assert len(rejected) == 2


def test_短すぎる引用はどの行にも当てない() -> None:
    # 「はい」程度の引用はどの行にも当たってしまう。根拠として認めない
    assert main.find_evidence("はい", TRANSCRIPT) is None


def test_JSON配列を取り出す() -> None:
    raw = 'こちらが結果です。\n```json\n[{"kind": "決定", "content": "あ"}]\n```\nご確認ください。'
    assert main.parse_items(raw) == [{"kind": "決定", "content": "あ"}]


def test_JSONで無ければ空にする() -> None:
    # 形式を外した答えを無理に読まない(読めたふりをしない)
    assert main.parse_items("すみません、抽出できませんでした。") == []
    assert main.parse_items("[壊れた") == []


def test_Excelを書き出せる(tmp_path) -> None:
    target = tmp_path / "一覧.xlsx"
    main.write_excel(
        [{"kind": "決定", "content": "一本化する", "owner": "要確認",
          "due": "要確認", "line": 1, "evidence": TRANSCRIPT[0][1]}],
        str(target),
    )
    from openpyxl import load_workbook

    sheet = load_workbook(target).active
    assert [cell.value for cell in sheet[1]][:4] == ["種別", "内容", "担当", "期限"]
    assert sheet.cell(row=2, column=1).value == "決定"


def test_項目が0件でも見出しだけの表を作る(tmp_path) -> None:
    # 何も採用できなくても、確認できる形の成果物は返す
    target = tmp_path / "空.xlsx"
    main.write_excel([], str(target))
    from openpyxl import load_workbook

    assert load_workbook(target).active.max_row == 1
