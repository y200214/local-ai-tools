"""決定事項・宿題の一覧を作る。文字起こしから抽出し、Excelにまとめる。

**ローカルLLMに探させるが、その答えを信用しない。**
LLMが挙げた項目は、必ず元の文字起こしと機械的に突き合わせてから採用する。

  1. LLMが「決定事項」「宿題」の候補と、その**根拠になった発言**を挙げる
  2. 根拠の発言が**本当に文字起こしにあるか**を照合する。無ければ捨てる
  3. 担当者・期限は、**その根拠の発言に書かれている場合だけ**採用する。
     書かれていなければ「要確認」にする(推測して埋めない)
  4. 内容に、根拠の発言に無い数値が混じっていたら捨てる(数値の捏造よけ)

捨てた候補は黙って消さず、「処理できなかったもの」として件数と理由を返す。
"""

from __future__ import annotations

import argparse
import json
import os
import re
import unicodedata
import urllib.request

MODEL = "gemma4:26b"
OLLAMA_URL = "http://127.0.0.1:11434/api/chat"
# LLMの答えが形式を外したときの作り直し回数。無限に粘らない
MAX_ATTEMPTS = 3
# 根拠として認める最短の長さ。短すぎる引用はどの行にも当たってしまう
MIN_EVIDENCE_CHARS = 8
UNKNOWN = "要確認"

PROMPT = """次の会議の文字起こしから、決定事項と宿題(持ち帰り事項)を抜き出してください。

必ずJSONの配列だけを返してください。説明や前置きは書かないでください。
各要素は次の形にしてください。

  {"kind": "決定" または "宿題",
   "content": "決まったこと・やること",
   "owner": "担当者(書かれていなければ空文字)",
   "due": "期限(書かれていなければ空文字)",
   "evidence": "根拠になった発言をそのまま引用"}

守ること:
- evidence は文字起こしにある文をそのまま写してください。要約しないでください
- 担当者・期限は、その発言に書かれている場合だけ書いてください。**推測しないでください**
- 決定でも宿題でもない雑談は入れないでください

文字起こし:
"""


def _normalized(text: str) -> str:
    """突き合わせ用に表記を均す。"""
    value = unicodedata.normalize("NFKC", text)
    return re.sub(r"[\s、。,\.!?！？「」『』()()]+", "", value)


def call_llm(transcript: str, model: str = MODEL, timeout: int = 300) -> str:
    """ローカルLLMへ1往復。応答本文を返す。"""
    payload = json.dumps({
        "model": model,
        "stream": False,
        "think": False,  # 思考を出すモデルだと本文が空で返ることがある
        "options": {"temperature": 0, "num_ctx": 16384},
        "messages": [{"role": "user", "content": PROMPT + transcript}],
    }, ensure_ascii=False).encode("utf-8")
    request = urllib.request.Request(
        OLLAMA_URL, data=payload,
        headers={"Content-Type": "application/json"}, method="POST",
    )
    with urllib.request.urlopen(request, timeout=timeout) as response:
        body = json.loads(response.read().decode("utf-8"))
    return (body.get("message") or {}).get("content", "")


def parse_items(raw: str) -> list[dict]:
    """LLMの応答からJSON配列を取り出す。前後に説明が付いていても拾う。"""
    text = raw.strip()
    if text.startswith("```"):
        text = re.sub(r"^```[a-zA-Z]*\s*|\s*```$", "", text).strip()
    start, end = text.find("["), text.rfind("]")
    if start == -1 or end == -1 or end < start:
        return []
    try:
        parsed = json.loads(text[start:end + 1])
    except ValueError:
        return []
    if not isinstance(parsed, list):
        return []
    return [item for item in parsed if isinstance(item, dict)]


def find_evidence(evidence: str, lines: list[tuple[int, str]]) -> tuple[int, str] | None:
    """根拠の引用が実在する行を探す。見つからなければ None。"""
    key = _normalized(evidence)
    if len(key) < MIN_EVIDENCE_CHARS:
        return None
    for number, line in lines:
        normalized = _normalized(line)
        if key in normalized or (len(normalized) >= MIN_EVIDENCE_CHARS and normalized in key):
            return number, line
    return None


def verify_items(
    items: list[dict], lines: list[tuple[int, str]]
) -> tuple[list[dict], list[str]]:
    """LLMの候補を1件ずつ検める。採用したものと、捨てた理由を返す。"""
    verified: list[dict] = []
    rejected: list[str] = []

    for index, item in enumerate(items, start=1):
        kind = str(item.get("kind") or "").strip()
        content = str(item.get("content") or "").strip()
        evidence = str(item.get("evidence") or "").strip()

        if kind not in ("決定", "宿題") or not content:
            rejected.append(f"{index}件目: 種別か内容が空でした")
            continue

        found = find_evidence(evidence, lines)
        if found is None:
            # 文字起こしに無い発言を根拠にしている(作り話の可能性)
            rejected.append(f"{index}件目: 根拠の発言が文字起こしに見つかりませんでした")
            continue
        number, source = found

        # 内容に、根拠の発言に無い数値が混じっていたら採らない
        source_digits = set(re.findall(r"\d+", unicodedata.normalize("NFKC", source)))
        content_digits = set(re.findall(r"\d+", unicodedata.normalize("NFKC", content)))
        invented = content_digits - source_digits
        if invented:
            rejected.append(
                f"{index}件目: 根拠に無い数値が入っていました({'、'.join(sorted(invented))})"
            )
            continue

        # 担当・期限は根拠の発言に書かれている場合だけ採る。無ければ埋めない
        normalized_source = _normalized(source)
        owner = str(item.get("owner") or "").strip()
        due = str(item.get("due") or "").strip()
        owner = owner if owner and _normalized(owner) in normalized_source else UNKNOWN
        due = due if due and _normalized(due) in normalized_source else UNKNOWN

        verified.append({
            "kind": kind, "content": content, "owner": owner, "due": due,
            "line": number, "evidence": source.strip(),
        })

    return verified, rejected


def write_excel(items: list[dict], target: str) -> None:
    """一覧をExcelへ書き出す。項目が0件でも見出しだけの表を作る。"""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font

    book = Workbook()
    sheet = book.active
    sheet.title = "決定事項・宿題"

    headers = ["種別", "内容", "担当", "期限", "根拠(行)", "根拠の発言"]
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True)

    for item in items:
        sheet.append([
            item["kind"], item["content"], item["owner"], item["due"],
            item["line"], item["evidence"],
        ])

    widths = [8, 44, 12, 14, 10, 50]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[sheet.cell(row=1, column=index).column_letter].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    book.save(target)


def main() -> int:
    parser = argparse.ArgumentParser(description="決定事項・宿題の一覧を作る")
    parser.add_argument("--request", required=True)
    args = parser.parse_args()

    request = json.loads(open(args.request, encoding="utf-8").read())
    job_dir = os.path.dirname(args.request)

    # 読み取りはコア側の共通部品に任せる(D-5)。Word・字幕・Excelなど、
    # 対応形式が増えてもここは変わらない。読めないときは理由と直し方が返る
    import toolpack_textio

    try:
        text, _ = toolpack_textio.read_request_input(request, job_dir)
    except toolpack_textio.UnreadableFile as error:
        print(json.dumps({
            "status": "user_error", "message": str(error),
            "files": [], "skipped": [], "notes": [],
        }, ensure_ascii=False))
        return 0

    lines = [
        (number, line) for number, line in enumerate(text.splitlines(), start=1)
        if line.strip()
    ]
    if not lines:
        print(json.dumps({
            "status": "user_error", "message": "文字起こしの中身がありません。",
            "files": [], "skipped": [], "notes": [],
        }, ensure_ascii=False))
        return 0

    notes: list[str] = []
    items: list[dict] = []
    last_error = ""
    for attempt in range(1, MAX_ATTEMPTS + 1):
        try:
            raw = call_llm(text)
        except OSError as error:
            print(json.dumps({
                "status": "user_error",
                "message": (
                    "ローカルLLMが応答しませんでした。"
                    "時間をおいて試すか、管理者へ連絡してください。"
                ),
                "files": [], "skipped": [], "notes": [],
            }, ensure_ascii=False))
            return 0
        items = parse_items(raw)
        if items:
            if attempt > 1:
                notes.append(f"LLMの答えが形式を外したため {attempt} 回目で作り直しました")
            break
        last_error = "決められた形式で答えませんでした"
    else:
        notes.append(f"LLMが{MAX_ATTEMPTS}回とも{last_error}")

    verified, rejected = verify_items(items, lines)

    target_name = "決定事項・宿題一覧.xlsx"
    write_excel(verified, os.path.join(job_dir, request["outdir"], target_name))

    decisions = sum(1 for item in verified if item["kind"] == "決定")
    homework = len(verified) - decisions
    unknown_owner = sum(1 for item in verified if item["owner"] == UNKNOWN)
    unknown_due = sum(1 for item in verified if item["due"] == UNKNOWN)

    message = (
        f"決定事項 {decisions} 件、宿題 {homework} 件を一覧にしました。\n"
        "根拠の発言が文字起こしにあるものだけを載せています。"
    )
    skipped = list(rejected)
    if not verified:
        skipped.append("採用できる項目がありませんでした(元の文字起こしをご確認ください)")
    if unknown_owner:
        notes.append(f"担当が書かれていない項目が {unknown_owner} 件あります(「{UNKNOWN}」)")
    if unknown_due:
        notes.append(f"期限が書かれていない項目が {unknown_due} 件あります(「{UNKNOWN}」)")

    print(json.dumps({
        "status": "ok", "message": message, "files": [target_name],
        "skipped": skipped, "notes": notes,
    }, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
