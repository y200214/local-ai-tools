"""
title: universal_file_export
id: universal_file_export
author: Y_KAZUKI
description: Export text and table data as TXT, MD, DOCX, CSV, or XLSX files and attach them to the current Open WebUI message.
required_open_webui_version: 0.6.0
requirements: python-docx,openpyxl,httpx
version: 2.0.0
license: MIT
"""

from __future__ import annotations

import csv
import inspect
import io
import json
import re
from pathlib import Path
from typing import Any, Awaitable, Callable, Literal, Optional

import httpx
from docx import Document
from docx.shared import Pt
from fastapi import Request
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from pydantic import BaseModel, Field


EventEmitter = Callable[[dict[str, Any]], Awaitable[None]]
DocumentFormat = Literal["txt", "md", "docx"]
TableFormat = Literal["csv", "xlsx"]


class Tools:
    class Valves(BaseModel):
        prefer_internal_upload: bool = Field(
            default=True,
            description=(
                "Open WebUI内部のファイル登録処理を優先して使用します。"
                "通常は有効のままにしてください。"
            ),
        )
        open_webui_base_url: str = Field(
            default="",
            description=(
                "内部登録が利用できない場合だけ使うOpen WebUI APIの接続先。"
                "通常は空欄。例: http://127.0.0.1:8080"
            ),
        )
        api_key: str = Field(
            default="",
            description=(
                "HTTP APIへのフォールバック時に使うOpen WebUI APIキー。"
                "内部登録を使う通常構成では空欄で構いません。"
            ),
        )
        timeout_seconds: int = Field(
            default=120,
            ge=10,
            le=600,
            description="HTTPフォールバック時のタイムアウト秒数。",
        )

    def __init__(self) -> None:
        self.valves = self.Valves()
        self.citation = False

    async def export_document(
        self,
        content: str,
        filename: str = "output",
        file_format: DocumentFormat = "docx",
        title: str = "",
        __request__: Optional[Request] = None,
        __user__: Optional[dict[str, Any]] = None,
        __event_emitter__: Optional[EventEmitter] = None,
    ) -> str:
        """
        完成済みの文章をTXT、Markdown、Wordのいずれかで出力します。
        contentを要約、省略、改変せず、完成した全文を渡してください。

        :param content: ファイルへ保存する完成済みの全文。
        :param filename: 拡張子を除いたファイル名。拡張子を含めても自動調整します。
        :param file_format: txt、md、docxのいずれか。
        :param title: Word文書の先頭に追加する任意のタイトル。不要なら空文字。
        :return: ファイル出力結果。
        """
        if not isinstance(content, str) or not content.strip():
            return "エラー: 出力する文章が空です。"
        if file_format not in {"txt", "md", "docx"}:
            return "エラー: file_formatはtxt、md、docxのいずれかを指定してください。"

        await self._emit_status(
            __event_emitter__, f"{file_format.upper()}ファイルを作成しています…", False
        )

        try:
            safe_name = self._make_filename(filename, file_format)

            if file_format == "txt":
                file_bytes = content.encode("utf-8-sig")
                mime_type = "text/plain; charset=utf-8"
            elif file_format == "md":
                file_bytes = content.encode("utf-8")
                mime_type = "text/markdown; charset=utf-8"
            else:
                file_bytes = self._build_docx(content=content, title=title)
                mime_type = (
                    "application/vnd.openxmlformats-officedocument."
                    "wordprocessingml.document"
                )

            file_info = await self._store_file(
                request=__request__,
                user_info=__user__,
                filename=safe_name,
                file_bytes=file_bytes,
                mime_type=mime_type,
            )
            await self._attach_file(
                event_emitter=__event_emitter__,
                file_id=str(file_info["id"]),
                filename=str(file_info.get("filename") or safe_name),
            )
            await self._emit_status(
                __event_emitter__, f"{safe_name}を出力しました。", True
            )
            return f"ファイルを出力しました: {safe_name}（{len(file_bytes):,}バイト）"
        except Exception as exc:
            message = self._format_exception(exc)
            await self._emit_status(
                __event_emitter__, f"ファイル出力に失敗しました: {message}", True
            )
            return f"エラー: ファイル出力に失敗しました。{message}"

    async def export_table(
        self,
        headers: list[str],
        rows: list[list[Any]],
        filename: str = "table",
        file_format: TableFormat = "xlsx",
        sheet_name: str = "Sheet1",
        protect_formulas: bool = True,
        __request__: Optional[Request] = None,
        __user__: Optional[dict[str, Any]] = None,
        __event_emitter__: Optional[EventEmitter] = None,
    ) -> str:
        """
        表データをCSVまたはExcelファイルとして出力します。
        全行をheadersとrowsに渡してください。

        :param headers: 列名の一覧。例: ["氏名", "年齢", "所属"]
        :param rows: データ行の二次元配列。各行の要素数はheadersと一致させます。
        :param filename: 拡張子を除いたファイル名。
        :param file_format: csvまたはxlsx。
        :param sheet_name: Excelのシート名。CSVでは使用しません。
        :param protect_formulas: 数式インジェクションを防ぐ安全対策。
        :return: ファイル出力結果。
        """
        if file_format not in {"csv", "xlsx"}:
            return "エラー: file_formatはcsvまたはxlsxを指定してください。"
        if not isinstance(headers, list) or not headers:
            return "エラー: headersには1列以上の列名を指定してください。"
        if not isinstance(rows, list):
            return "エラー: rowsは二次元配列で指定してください。"

        expected_columns = len(headers)
        invalid_rows = [
            index + 1
            for index, row in enumerate(rows)
            if not isinstance(row, list) or len(row) != expected_columns
        ]
        if invalid_rows:
            shown = ", ".join(map(str, invalid_rows[:10]))
            suffix = "…" if len(invalid_rows) > 10 else ""
            return (
                f"エラー: headersは{expected_columns}列ですが、"
                f"列数が一致しない行があります（行: {shown}{suffix}）。"
            )

        await self._emit_status(
            __event_emitter__, f"{file_format.upper()}ファイルを作成しています…", False
        )

        try:
            safe_name = self._make_filename(filename, file_format)
            normalized_headers = [str(value) for value in headers]
            normalized_rows = [
                [
                    self._protect_spreadsheet_value(value)
                    if protect_formulas
                    else value
                    for value in row
                ]
                for row in rows
            ]

            if file_format == "csv":
                file_bytes = self._build_csv(normalized_headers, normalized_rows)
                mime_type = "text/csv; charset=utf-8"
            else:
                file_bytes = self._build_xlsx(
                    normalized_headers,
                    normalized_rows,
                    sheet_name=sheet_name,
                )
                mime_type = (
                    "application/vnd.openxmlformats-officedocument."
                    "spreadsheetml.sheet"
                )

            file_info = await self._store_file(
                request=__request__,
                user_info=__user__,
                filename=safe_name,
                file_bytes=file_bytes,
                mime_type=mime_type,
            )
            await self._attach_file(
                event_emitter=__event_emitter__,
                file_id=str(file_info["id"]),
                filename=str(file_info.get("filename") or safe_name),
            )
            await self._emit_status(
                __event_emitter__, f"{safe_name}を出力しました。", True
            )
            return (
                f"ファイルを出力しました: {safe_name} "
                f"（{len(rows):,}行、{len(headers):,}列）"
            )
        except Exception as exc:
            message = self._format_exception(exc)
            await self._emit_status(
                __event_emitter__, f"ファイル出力に失敗しました: {message}", True
            )
            return f"エラー: ファイル出力に失敗しました。{message}"

    async def _store_file(
        self,
        request: Optional[Request],
        user_info: Optional[dict[str, Any]],
        filename: str,
        file_bytes: bytes,
        mime_type: str,
    ) -> dict[str, Any]:
        errors: list[str] = []

        if self.valves.prefer_internal_upload:
            try:
                return await self._upload_file_internal(
                    request=request,
                    user_info=user_info,
                    filename=filename,
                    file_bytes=file_bytes,
                    mime_type=mime_type,
                )
            except Exception as exc:
                errors.append(f"内部登録: {self._format_exception(exc)}")

        try:
            return await self._upload_file_http(
                request=request,
                filename=filename,
                file_bytes=file_bytes,
                mime_type=mime_type,
            )
        except Exception as exc:
            errors.append(f"HTTP登録: {self._format_exception(exc)}")

        raise RuntimeError(" / ".join(errors))

    @staticmethod
    async def _upload_file_internal(
        request: Optional[Request],
        user_info: Optional[dict[str, Any]],
        filename: str,
        file_bytes: bytes,
        mime_type: str,
    ) -> dict[str, Any]:
        if request is None:
            raise RuntimeError("__request__を取得できませんでした")
        if not user_info or not user_info.get("id"):
            raise RuntimeError("__user__からユーザーIDを取得できませんでした")

        # Open WebUI自身のバックエンド処理を直接呼ぶため、
        # Docker内からHTTPで自己接続する必要がありません。
        from open_webui.models.users import Users
        from open_webui.routers.files import upload_file_handler
        from starlette.datastructures import Headers, UploadFile

        user_model = Users.get_user_by_id(str(user_info["id"]))
        if inspect.isawaitable(user_model):
            user_model = await user_model
        if user_model is None:
            raise RuntimeError("Open WebUIのユーザー情報を読み込めませんでした")

        upload_kwargs: dict[str, Any] = {
            "file": io.BytesIO(file_bytes),
            "filename": filename,
            "headers": Headers({"content-type": mime_type}),
        }
        if "size" in inspect.signature(UploadFile).parameters:
            upload_kwargs["size"] = len(file_bytes)
        upload_file = UploadFile(**upload_kwargs)

        handler_values: dict[str, Any] = {
            "file": upload_file,
            "metadata": {
                "source": "universal_file_export",
                "generated": True,
            },
            "process": False,
            "process_in_background": False,
            "user": user_model,
            "background_tasks": None,
            "db": None,
        }
        supported_parameters = inspect.signature(upload_file_handler).parameters
        handler_kwargs = {
            key: value
            for key, value in handler_values.items()
            if key in supported_parameters
        }

        try:
            result = upload_file_handler(request, **handler_kwargs)
            if inspect.isawaitable(result):
                result = await result
        finally:
            close_result = upload_file.close()
            if inspect.isawaitable(close_result):
                await close_result

        if isinstance(result, dict):
            data = result
        elif hasattr(result, "model_dump"):
            data = result.model_dump()
        elif hasattr(result, "dict"):
            data = result.dict()
        else:
            raise RuntimeError("内部登録結果を解釈できませんでした")

        if not data.get("id"):
            raise RuntimeError("内部登録結果にファイルIDがありません")
        return data

    async def _upload_file_http(
        self,
        request: Optional[Request],
        filename: str,
        file_bytes: bytes,
        mime_type: str,
    ) -> dict[str, Any]:
        authorization = request.headers.get("authorization") if request else None
        if not authorization and self.valves.api_key.strip():
            authorization = f"Bearer {self.valves.api_key.strip()}"
        if not authorization:
            raise RuntimeError(
                "HTTPフォールバック用の認証情報がありません。"
                "Valvesのapi_keyを設定してください"
            )

        candidates: list[str] = []
        configured = self.valves.open_webui_base_url.strip()
        if configured:
            candidates.append(configured.rstrip("/"))
        if request is not None:
            candidates.append(str(request.base_url).rstrip("/"))
        candidates.extend(
            [
                "http://127.0.0.1:8080",
                "http://localhost:8080",
            ]
        )
        candidates = list(dict.fromkeys(candidates))

        headers = {
            "Authorization": authorization,
            "Accept": "application/json",
        }
        files = {"file": (filename, file_bytes, mime_type)}
        metadata = json.dumps(
            {"source": "universal_file_export", "generated": True},
            ensure_ascii=False,
        )
        timeout = httpx.Timeout(float(self.valves.timeout_seconds), connect=15.0)
        failures: list[str] = []

        async with httpx.AsyncClient(timeout=timeout, follow_redirects=True) as client:
            for base_url in candidates:
                upload_url = f"{base_url}/api/v1/files/"
                try:
                    response = await client.post(
                        upload_url,
                        params={
                            "process": "false",
                            "process_in_background": "false",
                        },
                        headers=headers,
                        data={"metadata": metadata},
                        files=files,
                    )
                    if response.is_error:
                        failures.append(
                            f"{base_url}: HTTP {response.status_code} "
                            f"{response.text[:200]}"
                        )
                        continue
                    result = response.json()
                    if not result.get("id"):
                        failures.append(f"{base_url}: ファイルIDなし")
                        continue
                    return result
                except Exception as exc:
                    failures.append(f"{base_url}: {self._format_exception(exc)}")

        raise RuntimeError("; ".join(failures))

    @staticmethod
    async def _attach_file(
        event_emitter: Optional[EventEmitter], file_id: str, filename: str
    ) -> None:
        if event_emitter is None:
            raise RuntimeError("__event_emitter__を取得できませんでした")

        # 短縮名 files を使うと、添付情報がチャットDBにも保存されます。
        await event_emitter(
            {
                "type": "files",
                "data": {
                    "files": [
                        {
                            "type": "file",
                            "id": file_id,
                            "name": filename,
                            "url": f"/api/v1/files/{file_id}/content?attachment=true",
                        }
                    ]
                },
            }
        )

    @staticmethod
    async def _emit_status(
        event_emitter: Optional[EventEmitter], description: str, done: bool
    ) -> None:
        if event_emitter is None:
            return
        await event_emitter(
            {
                "type": "status",
                "data": {
                    "description": description,
                    "done": done,
                    "hidden": False,
                },
            }
        )

    @staticmethod
    def _build_docx(content: str, title: str = "") -> bytes:
        document = Document()
        normal_style = document.styles["Normal"]
        normal_style.font.name = "Yu Gothic"
        normal_style.font.size = Pt(10.5)

        if title.strip():
            document.add_heading(Tools._strip_markdown(title.strip()), level=0)

        lines = content.replace("\r\n", "\n").replace("\r", "\n").split("\n")
        for line in lines:
            heading_match = re.match(r"^(#{1,6})\s+(.+)$", line)
            bold_heading_match = re.match(r"^\*\*(.+?)\*\*$", line.strip())
            bullet_match = re.match(r"^\s*[-*+]\s+(.+)$", line)
            numbered_match = re.match(r"^\s*\d+[.)]\s+(.+)$", line)

            if heading_match:
                level = min(len(heading_match.group(1)), 6)
                document.add_heading(
                    Tools._strip_markdown(heading_match.group(2)), level=level
                )
            elif bold_heading_match:
                document.add_heading(
                    Tools._strip_markdown(bold_heading_match.group(1)), level=1
                )
            elif bullet_match:
                document.add_paragraph(
                    Tools._strip_markdown(bullet_match.group(1)), style="List Bullet"
                )
            elif numbered_match:
                document.add_paragraph(
                    Tools._strip_markdown(numbered_match.group(1)), style="List Number"
                )
            else:
                document.add_paragraph(Tools._strip_markdown(line))

        buffer = io.BytesIO()
        document.save(buffer)
        return buffer.getvalue()

    @staticmethod
    def _build_csv(headers: list[str], rows: list[list[Any]]) -> bytes:
        buffer = io.StringIO(newline="")
        writer = csv.writer(buffer, lineterminator="\r\n")
        writer.writerow(headers)
        writer.writerows(rows)
        return buffer.getvalue().encode("utf-8-sig")

    @staticmethod
    def _build_xlsx(
        headers: list[str], rows: list[list[Any]], sheet_name: str = "Sheet1"
    ) -> bytes:
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = Tools._sanitize_sheet_name(sheet_name)
        worksheet.freeze_panes = "A2"
        worksheet.append(headers)

        for cell in worksheet[1]:
            cell.font = Font(bold=True)

        for row in rows:
            worksheet.append(row)

        for column_index, header in enumerate(headers, start=1):
            values = [header]
            values.extend(row[column_index - 1] for row in rows)
            max_length = max(
                len(str(value)) if value is not None else 0 for value in values
            )
            worksheet.column_dimensions[get_column_letter(column_index)].width = min(
                max(max_length + 2, 8), 60
            )

        buffer = io.BytesIO()
        workbook.save(buffer)
        return buffer.getvalue()

    @staticmethod
    def _make_filename(filename: str, extension: str) -> str:
        raw_name = Path((filename or "output").strip()).name
        stem = Path(raw_name).stem if Path(raw_name).suffix else raw_name
        stem = re.sub(r'[<>:"/\\|?*\x00-\x1F]', "_", stem).strip(" .")
        if not stem:
            stem = "output"
        return f"{stem[:120]}.{extension}"

    @staticmethod
    def _sanitize_sheet_name(sheet_name: str) -> str:
        value = re.sub(r"[\\/*?:\[\]]", "_", (sheet_name or "Sheet1")).strip()
        return (value or "Sheet1")[:31]

    @staticmethod
    def _protect_spreadsheet_value(value: Any) -> Any:
        if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
            return "'" + value
        return value

    @staticmethod
    def _strip_markdown(value: str) -> str:
        value = re.sub(r"\*\*(.+?)\*\*", r"\1", value)
        value = re.sub(r"__(.+?)__", r"\1", value)
        value = re.sub(r"`(.+?)`", r"\1", value)
        return value

    @staticmethod
    def _format_exception(exc: Exception) -> str:
        text = str(exc).strip()
        return text or exc.__class__.__name__
