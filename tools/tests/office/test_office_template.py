from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document
from docx.shared import RGBColor
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from tools.office_template import (
    SAMPLE_MINUTES_DOCUMENT,
    TemplateError,
    check_minutes_template,
    create_demo,
    deploy_minutes_template,
    evaluate_expression,
    load_minutes_document,
    render_minutes_preview,
    render_template,
)


def test_evaluate_expression_reads_nested_field_and_threshold() -> None:
    values = {"header": {"title": "運営会議"}, "rate": 92.4}

    assert evaluate_expression("field:header.title", values) == "運営会議"
    assert (
        evaluate_expression("threshold:rate|>=|90|要確認|基準内", values)
        == "要確認"
    )


def test_excel_template_keeps_style_and_numeric_type(tmp_path: Path) -> None:
    source = tmp_path / "template.xlsx"
    target = tmp_path / "rendered.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "レポート"
    sheet["B2"] = "{{field:rate}}"
    sheet["B2"].font = Font(size=18, bold=True, color="C65911")
    sheet["B2"].fill = PatternFill("solid", fgColor="FFF2CC")
    sheet["B2"].number_format = '0.0"%"'
    sheet["B3"] = "判定: {{threshold:rate|>=|90|要確認|基準内}}"
    workbook.save(source)
    workbook.close()
    original = source.read_bytes()

    count = render_template(source, {"rate": 92.4}, target)

    assert count == 2
    assert source.read_bytes() == original
    result = load_workbook(target)
    try:
        assert result["レポート"]["B2"].value == 92.4
        assert result["レポート"]["B2"].font.size == 18
        assert result["レポート"]["B2"].font.bold is True
        assert result["レポート"]["B2"].fill.fgColor.rgb == "00FFF2CC"
        assert result["レポート"]["B2"].number_format == '0.0"%"'
        assert result["レポート"]["B3"].value == "判定: 要確認"
    finally:
        result.close()


def test_word_template_replaces_split_runs_with_first_run_style(tmp_path: Path) -> None:
    source = tmp_path / "template.docx"
    target = tmp_path / "rendered.docx"
    document = Document()
    paragraph = document.add_paragraph("氏名: ")
    first = paragraph.add_run("{{field:")
    first.font.color.rgb = RGBColor(192, 0, 0)
    paragraph.add_run("name}}")
    document.save(source)
    original = source.read_bytes()

    count = render_template(source, {"name": "乙野"}, target)

    assert count == 1
    assert source.read_bytes() == original
    result = Document(target)
    assert result.paragraphs[0].text == "氏名: 乙野"
    replacement_run = next(run for run in result.paragraphs[0].runs if "乙野" in run.text)
    assert replacement_run.font.color.rgb == RGBColor(192, 0, 0)


def test_invalid_template_does_not_create_output(tmp_path: Path) -> None:
    source = tmp_path / "template.docx"
    target = tmp_path / "rendered.docx"
    document = Document()
    document.add_paragraph("{{field:missing}}")
    document.save(source)

    with pytest.raises(TemplateError, match="入力値に項目がありません"):
        render_template(source, {}, target)

    assert not target.exists()


def test_template_itself_cannot_be_overwritten(tmp_path: Path) -> None:
    source = tmp_path / "template.docx"
    document = Document()
    document.add_paragraph("{{field:name}}")
    document.save(source)

    with pytest.raises(TemplateError, match="テンプレート自身"):
        render_template(source, {"name": "乙野"}, source, force=True)


def test_create_demo_makes_editable_templates_and_results(tmp_path: Path) -> None:
    created = create_demo(tmp_path)

    assert len(created) == 6
    assert all(path.exists() for path in created)
    excel = load_workbook(tmp_path / "demo_result.xlsx")
    try:
        assert excel["レポート"]["B3"].value == 92.4
        assert "確認が必要" in excel["レポート"]["B4"].value
    finally:
        excel.close()
    word = Document(tmp_path / "demo_result.docx")
    assert "92.4" in "\n".join(
        paragraph.text
        for table in word.tables
        for row in table.rows
        for cell in row.cells
        for paragraph in cell.paragraphs
    )


# ------------------------------------------------------------------
# 議事録テンプレートの編集導線(check / preview / deploy)
# ------------------------------------------------------------------


def _minutes_like_template(path: Path, body_rows: int = 21) -> None:
    """実物と同じ構造(タイトル行+ヘッダ表+21行の本文表)の合成テンプレを作る。"""
    document = Document()
    document.add_paragraph("○○委員会議事録")
    document.add_paragraph("日時：")
    document.add_paragraph("場所：")
    document.add_table(rows=2, cols=6)
    document.add_table(rows=body_rows, cols=2)
    document.save(path)


def test_check_accepts_a_structurally_sound_template(tmp_path: Path) -> None:
    template = tmp_path / "template.docx"
    _minutes_like_template(template)

    assert check_minutes_template(template) == []


def test_check_names_deleted_rows_in_plain_language(tmp_path: Path) -> None:
    """行を消して壊した場合に、何行必要かまで言えないと直しようがない。"""
    template = tmp_path / "template.docx"
    _minutes_like_template(template, body_rows=6)

    problems = check_minutes_template(template)

    assert problems
    assert "21行必要" in problems[0]
    assert "6行しかありません" in problems[0]


def test_check_notices_missing_title_line(tmp_path: Path) -> None:
    template = tmp_path / "template.docx"
    document = Document()
    document.add_paragraph("タイトルなし")
    document.add_table(rows=2, cols=6)
    document.add_table(rows=21, cols=2)
    document.save(template)

    problems = check_minutes_template(template)

    assert problems and "議事録" in problems[0]


def test_check_explains_non_docx_files(tmp_path: Path) -> None:
    template = tmp_path / "template.docx"
    template.write_bytes(b"not a docx")

    problems = check_minutes_template(template)

    assert problems and "開けません" in problems[0]


def test_preview_renders_sample_minutes_with_template_style(tmp_path: Path) -> None:
    """previewはサンプル議事録を実運用レンダラーで描く(書式確認の入口)。"""
    template = tmp_path / "template.docx"
    _minutes_like_template(template)
    target = tmp_path / "preview.docx"

    render_minutes_preview(template, SAMPLE_MINUTES_DOCUMENT, target)

    rendered = Document(target)
    texts = "\n".join(
        paragraph.text
        for table in rendered.tables
        for row in table.rows
        for cell in row.cells
        for paragraph in cell.paragraphs
    )
    assert "サンプル議題（書式確認用）" in texts
    assert "委員A" in texts


def test_deploy_refuses_broken_template_before_touching_docker(tmp_path: Path) -> None:
    """壊れたテンプレートは、dockerへ触る前に理由つきで止める。"""
    template = tmp_path / "template.docx"
    _minutes_like_template(template, body_rows=6)

    with pytest.raises(TemplateError, match="反映を中止しました"):
        deploy_minutes_template(template)


def test_load_minutes_document_accepts_api_response_wrapper(tmp_path: Path) -> None:
    path = tmp_path / "minutes.json"
    path.write_text(
        '{"document":{"ヘッダ":{},"議題":[],"その他":[]},"warnings":[]}',
        encoding="utf-8",
    )

    assert load_minutes_document(path) == {
        "ヘッダ": {},
        "議題": [],
        "その他": [],
    }
