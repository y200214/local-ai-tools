from pathlib import Path
from urllib.parse import unquote

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Border, PatternFill, Side

from excel_comment_report import build_comment_report, detect_reference_range


def _add_reference_tables(sheet, start_row: int) -> tuple[int, int]:
    thin = Side(style="thin", color="222222")
    table_border = Border(left=thin, right=thin, top=thin, bottom=thin)
    headings = [
        (start_row, "(1)入院関係の実績"),
        (start_row + 5, "途中に挟まった注記"),
        (start_row + 7, "(2)外来関係の実績"),
        (start_row + 14, "(3)医業収入（入外合算）実績"),
    ]
    for row, text in headings:
        sheet.cell(row, 1, text)
    end_row = start_row + 18
    for row in range(start_row, end_row + 1):
        for column in range(1, 12):
            cell = sheet.cell(row, column)
            if cell.value is None:
                cell.value = f"値{row}-{column}"
            cell.border = table_border
    sheet.cell(start_row + 1, 1).fill = PatternFill("solid", fgColor="D9D9D9")
    sheet.merge_cells(start_row=start_row, start_column=1, end_row=start_row, end_column=5)
    return start_row, end_row


def test_detect_reference_range_keeps_all_three_tables_with_intervening_content() -> None:
    workbook = Workbook()
    sheet = workbook.active
    expected = _add_reference_tables(sheet, 10)
    sheet.cell(expected[1] + 1, 1).value = None

    assert detect_reference_range(sheet) == expected


def test_build_comment_report_stacks_comments_and_static_tables(tmp_path: Path) -> None:
    source_path = tmp_path / "科別分析_updated.xlsm"
    workbook = Workbook()
    first = workbook.active
    first.title = "内科（案）"
    first.merge_cells("AG5:AY9")
    first["AG5"] = "内科コメント"
    first_start, first_end = _add_reference_tables(first, 10)
    first["K12"] = "=1+1"
    second = workbook.create_sheet("外科（案）")
    second.merge_cells("AG6:AY11")
    second["AG6"] = "外科コメント"
    second_start, second_end = _add_reference_tables(second, 16)
    workbook.save(source_path)

    records = [
        {
            "sheet": "内科（案）",
            "cell": "AG7",
            "value": "《入院》\n長いコメントを全文表示します。\n《まとめ》\n確認事項です。",
            "template": "入院・外来分析",
        },
        {
            "sheet": "外科（案）",
            "cell": "AG8",
            "value": "外科のコメント本文",
            "template": "入院・外来分析",
        },
    ]
    report_path = build_comment_report(source_path, records)

    assert report_path == tmp_path / "科別分析_updated_コメント確認一覧.xlsx"
    assert report_path.exists()
    report = load_workbook(report_path, data_only=False)
    sheet = report["コメント確認一覧"]
    assert sheet.sheet_view.showGridLines is False

    texts = {
        cell.value: cell.coordinate
        for row in sheet.iter_rows()
        for cell in row
        if isinstance(cell.value, str)
    }
    assert "No.1  内科（案）!AG7" in texts
    assert "No.2  外科（案）!AG8" in texts
    assert records[0]["value"] in texts
    assert records[1]["value"] in texts
    assert f"A{first_start}:K{first_end}" in texts
    assert f"A{second_start}:K{second_end}" in texts

    destination = next(
        cell
        for row in sheet.iter_rows()
        for cell in row
        if cell.value == "内科（案）!AG5"
    )
    assert destination.hyperlink is not None
    assert "科別分析_updated.xlsm" in unquote(destination.hyperlink.target)
    assert "AG5" in destination.hyperlink.target

    copied_heading = next(
        cell
        for row in sheet.iter_rows()
        for cell in row
        if cell.value == "(1)入院関係の実績"
    )
    assert copied_heading.column == 7
    copied_fill = sheet.cell(copied_heading.row + 1, 7)
    assert copied_fill.fill.fgColor.rgb == "00D9D9D9"
    assert copied_fill.border.left.style == "thin"
    assert not any(
        cell.data_type == "f"
        for row in sheet.iter_rows()
        for cell in row
    )
    report.close()

    original = load_workbook(source_path, keep_vba=True)
    assert original.sheetnames == ["内科（案）", "外科（案）"]
    assert original["内科（案）"]["AG5"].value == "内科コメント"
    original.close()
