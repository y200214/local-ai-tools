from pathlib import Path

import pytest
from openpyxl import Workbook

from excel_comment_context import (
    _analysis_value,
    _comparison_fact_records,
    _comparison_facts,
    _display_metric,
    collect_comment_contexts,
    prepare_comment_context,
    prepare_comment_contexts,
)


def _write_comment_sheet(workbook, title: str, *, with_example: bool = True):
    """コメント見出しと(任意で)記入例を持つシートを作る。"""
    sheet = workbook.create_sheet(title)
    sheet["AG4"] = "コメント"
    if with_example:
        # 記入例が無いと、直下セルが空のまま残り「記入例が見つかりません」になる
        sheet["AG5"] = "変更前"
        sheet["AG61"] = "コメント"
        sheet["AG62"] = "【例】\n・文章例"
    return sheet


def test_analysis_value_prefers_cached_formula_result() -> None:
    assert _analysis_value("=", 1_395_000_000) == 1_395_000_000
    assert _analysis_value("=A1+B1", 250) == 250
    assert _analysis_value("見出し", "見出し") == "見出し"


def test_display_metric_applies_excel_million_scale() -> None:
    assert _display_metric("①収入（百万円）", 1_395_000_000, "#,##0,,") == (
        1395.0,
        "1,395百万円",
        "百万円",
    )


def test_display_metric_recognizes_reverse_referral_rate_without_format() -> None:
    assert _display_metric("逆紹介率", 53.2, "General") == (53.2, "53.2‰", "‰")


def test_comparison_facts_use_latest_years_and_target() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet["B10"] = "(1)入院関係の実績"
    sheet["H11"] = "H30年度"
    sheet["I11"] = "R6年度"
    sheet["J11"] = "R7年度"
    sheet["K11"] = "R8年度"
    sheet["B12"] = "①収入"
    sheet["H12"] = 80
    sheet["I12"] = 90
    sheet["J12"] = 100
    sheet["K12"] = 110
    sheet["B13"] = "稼働率"
    sheet["J13"] = 0.934
    sheet["K13"] = 0.941
    sheet["B24"] = "(2)外来関係の実績"
    sheet["H25"] = "H30年度"
    sheet["J25"] = "R7年度"
    sheet["K25"] = "R8年度"
    sheet["B26"] = "外来患者数"
    sheet["J26"] = 200
    sheet["K26"] = 220
    sheet["B27"] = "逆紹介率"
    sheet["J27"] = 0.397
    sheet["K27"] = 0.418
    sheet["J27"].number_format = '#,##0.0"‰"'
    sheet["K27"].number_format = '#,##0.0"‰"'
    sheet["B39"] = "(3)医業収入（入外合算）実績"
    sheet["B40"] = "項目"
    sheet["J40"] = "R7年度"
    sheet["K40"] = "R8年度"
    sheet["B41"] = "合算収入"
    sheet["J41"] = 300
    sheet["K41"] = 330
    sheet["M5"] = "◆全病床稼働率93％以上"

    facts = _comparison_facts(sheet, sheet)
    records = _comparison_fact_records(sheet, sheet)

    assert "入院・①収入: R7年度=100, R8年度=110, 変化=+10.0%" in facts
    assert "入院・稼働率: R7年度=93.4%, R8年度=94.1%, 変化=+0.7ポイント" in facts
    assert "外来・外来患者数: R7年度=200, R8年度=220, 変化=+10.0%, 増減数=+20人" in facts
    assert "外来・逆紹介率: R7年度=39.7‰, R8年度=41.8‰, 変化=+2.1‰" in facts
    assert "入外合計・合算収入: R7年度=300, R8年度=330, 変化=+10.0%" in facts
    assert "入院・稼働率目標: 目標=93%以上, 最新=94.1%, 判定=達成, 差=+1.1ポイント" in facts
    assert "外来・逆紹介率基準: 基準=50‰, 最新=41.8‰, 判定=未達, 差=-8.2‰" in facts

    occupancy = next(
        record
        for record in records
        if record["kind"] == "comparison" and record["metric"] == "稼働率"
    )
    assert occupancy["section"] == "入院"
    assert occupancy["previous_year"] == "R7年度"
    assert occupancy["previous_value"] == 93.4
    assert occupancy["latest_year"] == "R8年度"
    assert occupancy["latest_value"] == 94.1
    assert occupancy["value_unit"] == "%"
    assert occupancy["absolute_change"] == pytest.approx(0.7)
    assert occupancy["change"] == pytest.approx(0.7)
    assert occupancy["change_text"] == "+0.7ポイント"
    assert occupancy["change_unit"] == "ポイント"
    trend = next(
        record
        for record in records
        if record["kind"] == "trend_comparison" and record["metric"] == "①収入"
    )
    assert trend["previous_year"] == "R6年度"
    long_term = next(
        record
        for record in records
        if record["kind"] == "long_comparison" and record["metric"] == "①収入"
    )
    assert long_term["previous_year"] == "H30年度"
    reverse_target = next(
        record
        for record in records
        if record["kind"] == "target" and record["metric"] == "逆紹介率"
    )
    assert reverse_target["target"] == 50.0
    assert reverse_target["gap"] == pytest.approx(-8.2)
    assert reverse_target["gap_unit"] == "‰"


def test_target_facts_use_latest_value_when_previous_year_is_blank() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet["M5"] = "◆全病床稼働率93％以上の達成・維持"
    sheet["B10"] = "(1)入院関係の実績"
    sheet["J11"] = "R7年度"
    sheet["K11"] = "R8年度"
    sheet["B12"] = "稼働率"
    sheet["J12"] = "-"
    sheet["K12"] = 0.8676
    sheet["B24"] = "(2)外来関係の実績"
    sheet["J25"] = "R7年度"
    sheet["K25"] = "R8年度"
    sheet["B26"] = "逆紹介率"
    sheet["J26"] = "-"
    sheet["K26"] = 53.2

    records = _comparison_fact_records(sheet, sheet)

    occupancy = next(record for record in records if record["kind"] == "target" and record["metric"] == "稼働率")
    reverse = next(record for record in records if record["kind"] == "target" and record["metric"] == "逆紹介率")
    assert occupancy["latest_value"] == pytest.approx(86.76)
    assert occupancy["gap"] == pytest.approx(-6.24)
    assert reverse["latest_value"] == pytest.approx(53.2)
    assert reverse["status"] == "達成"
    assert reverse["gap"] == pytest.approx(3.2)
    assert reverse["target_source"] == "fallback"


def test_reverse_referral_target_prefers_workbook_value() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet["B24"] = "(2)外来関係の実績"
    sheet["J25"] = "R7年度"
    sheet["K25"] = "R8年度"
    sheet["B26"] = "逆紹介率"
    sheet["J26"] = 42
    sheet["K26"] = 43
    sheet["BB31"] = "目標逆紹介率"
    sheet["BE31"] = 50
    sheet["BF31"] = 55

    target = next(
        record for record in _comparison_fact_records(sheet, sheet)
        if record["kind"] == "target" and record["metric"] == "逆紹介率"
    )
    assert target["target"] == 55
    assert target["gap"] == -12
    assert target["target_source"] == "workbook"


def test_prepare_comment_context_finds_data_template_and_target(tmp_path: Path) -> None:
    path = tmp_path / "source.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "診療科"
    sheet["B10"] = "入院収入"
    sheet["K10"] = 120
    sheet.merge_cells("AG4:AY4")
    sheet["AG4"] = "コメント"
    sheet.merge_cells("AG5:AY9")
    sheet.merge_cells("AG61:AY61")
    sheet["AG61"] = "コメント"
    sheet.merge_cells("AG62:AY69")
    sheet["AG62"] = "【例】\n（入院）\n・収入は増加している。"
    workbook.save(path)

    result = prepare_comment_context(path)

    assert result["sheet"] == "診療科"
    assert result["target"] == "AG5"
    assert "収入は増加" in result["template"]
    assert {"cell": "K10", "value": "120"} in result["data"]
    assert result["facts"] == []
    assert result["fact_records"] == []


def test_prepare_comment_context_skips_visible_non_comment_sheet(tmp_path: Path) -> None:
    path = tmp_path / "source.xlsx"
    workbook = Workbook()
    workbook.active.title = "メッセージ"
    workbook.active["A1"] = "説明"
    sheet = workbook.create_sheet("診療科（案）")
    sheet.merge_cells("AG4:AY4")
    sheet["AG4"] = "コメント"
    sheet.merge_cells("AG5:AY9")
    sheet.merge_cells("AG61:AY61")
    sheet["AG61"] = "コメント"
    sheet.merge_cells("AG62:AY69")
    sheet["AG62"] = "【例】\n（入院）\n・実績を記載する。"
    workbook.save(path)

    result = prepare_comment_context(path)

    assert result["sheet"] == "診療科（案）"
    assert result["target"] == "AG5"


def test_prepare_comment_contexts_finds_all_visible_targets_with_shifted_rows(tmp_path: Path) -> None:
    path = tmp_path / "source.xlsx"
    workbook = Workbook()
    workbook.active.title = "メッセージ"
    for title, header_row, target_row, example_row in (
        ("内科（案）", 4, 5, 62),
        ("外科（案）", 5, 6, 64),
    ):
        sheet = workbook.create_sheet(title)
        sheet[f"AG{header_row}"] = "分析内容（案）" if title.startswith("外科") else "コメント"
        sheet[f"AG{target_row}"] = "変更前"
        sheet[f"AG{example_row - 1}"] = "コメント"
        sheet[f"AG{example_row}"] = "【例】\n・文章例"
    hidden = workbook.create_sheet("非表示（案）")
    hidden.sheet_state = "hidden"
    hidden["AG4"] = "コメント"
    hidden["AG5"] = "対象外"
    workbook.save(path)

    results = prepare_comment_contexts(path)

    assert [(item["sheet"], item["target"]) for item in results] == [
        ("内科（案）", "AG5"),
        ("外科（案）", "AG6"),
    ]


def test_broken_sheet_does_not_take_down_the_others(tmp_path: Path) -> None:
    path = tmp_path / "source.xlsx"
    workbook = Workbook()
    workbook.active.title = "メッセージ"
    _write_comment_sheet(workbook, "内科（案）")
    _write_comment_sheet(workbook, "糖内（案）", with_example=False)
    _write_comment_sheet(workbook, "外科（案）")
    workbook.save(path)

    contexts, skipped = collect_comment_contexts(path)

    # 記入例の無い1枚のために残り2枚まで消えてはいけない
    assert [item["sheet"] for item in contexts] == ["内科（案）", "外科（案）"]
    assert [item["sheet"] for item in skipped] == ["糖内（案）"]
    assert "記入例" in skipped[0]["reason"]


def test_all_sheets_broken_still_fails_with_reasons(tmp_path: Path) -> None:
    path = tmp_path / "source.xlsx"
    workbook = Workbook()
    workbook.active.title = "メッセージ"
    _write_comment_sheet(workbook, "内科（案）", with_example=False)
    workbook.save(path)

    with pytest.raises(RuntimeError) as error:
        collect_comment_contexts(path)

    assert "1枚も抽出できませんでした" in str(error.value)
    assert "内科（案）" in str(error.value)


def test_named_sheet_failure_is_not_swallowed(tmp_path: Path) -> None:
    path = tmp_path / "source.xlsx"
    workbook = Workbook()
    workbook.active.title = "メッセージ"
    _write_comment_sheet(workbook, "内科（案）", with_example=False)
    workbook.save(path)

    # シートを名指しされた場合は、握らずそのまま失敗を伝える
    with pytest.raises(RuntimeError, match="記入例"):
        collect_comment_contexts(path, "内科（案）")
