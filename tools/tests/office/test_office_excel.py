import json
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from office_excel import (
    _replace_cell_xml,
    cmd_set,
    cmd_set_batch,
    rejected_update_reason,
)


def test_cmd_set_redirects_merged_cell_to_top_left(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.chdir(tmp_path)
    source_dir = tmp_path / "input"
    source_dir.mkdir()
    source = source_dir / "source.xlsx"

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "集計"
    sheet.merge_cells("B2:D4")
    workbook.save(source)

    cmd_set(source, "集計", "C3", "コメント本文", None)

    result = load_workbook(tmp_path / "output" / "source_updated.xlsx")
    assert result["集計"]["B2"].value == "コメント本文"
    assert load_workbook(source)["集計"]["B2"].value is None


def test_cmd_set_reads_multiline_value_from_output(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.chdir(tmp_path)
    source_dir = tmp_path / "input"
    source_dir.mkdir()
    source = source_dir / "source.xlsx"
    Workbook().save(source)
    value_file = tmp_path / "output" / "comment.txt"
    value_file.parent.mkdir()
    value_file.write_text("1行目\n2行目\n", encoding="utf-8")

    cmd_set(source, "Sheet", "A1", None, str(value_file))

    result = load_workbook(tmp_path / "output" / "source_updated.xlsx")
    assert result["Sheet"]["A1"].value == "1行目\n2行目"


def test_cmd_set_preserves_formula_cached_values(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.chdir(tmp_path)
    source_dir = tmp_path / "input"
    source_dir.mkdir()
    source = source_dir / "source.xlsx"

    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"] = "変更前"
    sheet["B2"] = "=1+1"
    workbook.save(source)

    patched = tmp_path / "cached.xlsx"
    with ZipFile(source, "r") as source_archive, ZipFile(
        patched, "w", compression=ZIP_DEFLATED
    ) as target_archive:
        for info in source_archive.infolist():
            data = source_archive.read(info.filename)
            if info.filename == "xl/worksheets/sheet1.xml":
                data = data.replace(b"<v></v>", b"<v>2</v>")
            target_archive.writestr(info, data)
    patched.replace(source)

    assert load_workbook(source, data_only=True)["Sheet"]["B2"].value == 2
    cmd_set(source, "Sheet", "A1", "変更後", None)

    result_path = tmp_path / "output" / "source_updated.xlsx"
    formula_result = load_workbook(result_path, data_only=False)
    value_result = load_workbook(result_path, data_only=True)
    assert formula_result["Sheet"]["A1"].value == "変更後"
    assert formula_result["Sheet"]["B2"].value == "=1+1"
    assert value_result["Sheet"]["B2"].value == 2


def test_replace_cell_xml_supports_prefixed_spreadsheet_namespace() -> None:
    source = (
        b'<x:worksheet xmlns:x="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
        b'<x:sheetData><x:row r="5"><x:c r="AG5" s="17" /></x:row></x:sheetData>'
        b'</x:worksheet>'
    )

    result = _replace_cell_xml(source, "AG5", "1行目\n2行目")

    assert b'<x:c r="AG5" s="17" t="inlineStr">' in result
    assert "1行目\n2行目" in result.decode("utf-8")


def test_cmd_set_batch_writes_multiple_sheets_to_one_output(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.chdir(tmp_path)
    source_dir = tmp_path / "input"
    source_dir.mkdir()
    source = source_dir / "source.xlsm"

    workbook = Workbook()
    first = workbook.active
    first.title = "内科（案）"
    first.merge_cells("AG5:AY9")
    second = workbook.create_sheet("外科（案）")
    second.merge_cells("AG6:AY11")
    workbook.save(source)

    cmd_set_batch(
        source,
        '[{"sheet":"内科（案）","cell":"AG7","value":"内科コメント"},'
        '{"sheet":"外科（案）","cell":"AG8","value":"外科コメント"}]',
    )

    result_path = tmp_path / "output" / "source_updated.xlsm"
    result = load_workbook(result_path, keep_vba=True)
    assert result["内科（案）"]["AG5"].value == "内科コメント"
    assert result["外科（案）"]["AG6"].value == "外科コメント"
    assert load_workbook(source, keep_vba=True)["内科（案）"]["AG5"].value is None


def test_cmd_set_batch_reads_long_updates_from_utf8_file(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.chdir(tmp_path)
    source_dir = tmp_path / "input"
    source_dir.mkdir()
    source = source_dir / "source.xlsx"

    workbook = Workbook()
    first = workbook.active
    first.title = "内科（案）"
    second = workbook.create_sheet("外科（案）")
    workbook.save(source)

    updates_file = tmp_path / "work" / "updates.json"
    updates_file.parent.mkdir()
    updates_file.write_text(
        '[{"sheet":"内科（案）","cell":"A1","value":"内科コメント",'
        '"template":"入院・外来分析"},'
        '{"sheet":"外科（案）","cell":"A1","value":"外科コメント"}]',
        encoding="utf-8",
    )

    cmd_set_batch(source, updates_file=str(updates_file))

    result = load_workbook(tmp_path / "output" / "source_updated.xlsx")
    assert result["内科（案）"]["A1"].value == "内科コメント"
    assert result["外科（案）"]["A1"].value == "外科コメント"


def test_rejected_update_reason_names_the_problem() -> None:
    sheetnames = ["内科（案）"]

    assert rejected_update_reason({"sheet": "内科（案）", "cell": "A1", "value": "本文"}, sheetnames) == ""
    assert "シートがありません" in rejected_update_reason(
        {"sheet": "糖内（案）", "cell": "A1", "value": "本文"}, sheetnames
    )
    assert "セル座標が不正" in rejected_update_reason(
        {"sheet": "内科（案）", "cell": "A", "value": "本文"}, sheetnames
    )
    assert "設定値が不正" in rejected_update_reason(
        {"sheet": "内科（案）", "cell": "A1", "value": None}, sheetnames
    )
    assert rejected_update_reason("文字列", sheetnames)


def test_set_batch_writes_the_valid_rows_and_reports_the_rest(
    tmp_path: Path, monkeypatch, capsys
) -> None:
    monkeypatch.chdir(tmp_path)
    source_dir = tmp_path / "input"
    source_dir.mkdir()
    source = source_dir / "source.xlsx"
    workbook = Workbook()
    workbook.active.title = "内科（案）"
    workbook.create_sheet("外科（案）")
    workbook.save(source)

    cmd_set_batch(
        source,
        updates_json=(
            '[{"sheet":"内科（案）","cell":"A1","value":"内科コメント"},'
            '{"sheet":"糖内（案）","cell":"A1","value":"消えるはずの1件"},'
            '{"sheet":"外科（案）","cell":"A1","value":"外科コメント"}]'
        ),
    )

    result = load_workbook(tmp_path / "output" / "source_updated.xlsx")
    # 存在しないシート1件のために、書ける2件まで捨ててはいけない
    assert result["内科（案）"]["A1"].value == "内科コメント"
    assert result["外科（案）"]["A1"].value == "外科コメント"
    printed = capsys.readouterr().out
    assert "書き込めなかった更新 (1件):" in printed
    assert "糖内（案）" in printed


def test_set_batch_still_fails_when_nothing_is_writable(
    tmp_path: Path, monkeypatch
) -> None:
    monkeypatch.chdir(tmp_path)
    source_dir = tmp_path / "input"
    source_dir.mkdir()
    source = source_dir / "source.xlsx"
    workbook = Workbook()
    workbook.active.title = "内科（案）"
    workbook.save(source)

    with pytest.raises(SystemExit, match="書き込める更新が1件もありません"):
        cmd_set_batch(
            source, updates_json='[{"sheet":"糖内（案）","cell":"A1","value":"本文"}]'
        )


def test_comment_write_uses_template_font_without_resizing_cells(
    tmp_path: Path, monkeypatch
) -> None:
    monkeypatch.chdir(tmp_path)
    source_dir = tmp_path / "input"
    source_dir.mkdir()
    source = source_dir / "source.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "分析"
    sheet.merge_cells("B2:F5")
    sheet.column_dimensions["B"].width = 12
    sheet.row_dimensions[2].height = 24
    sheet["B2"].font = Font(name="Yu Gothic", size=15, color="C00000")
    workbook.save(source)

    cmd_set_batch(
        source,
        updates_json='[{"sheet":"分析","cell":"B2","value":"コメント本文"}]',
    )

    target = tmp_path / "output" / "source_updated.xlsx"
    result = load_workbook(target)
    try:
        assert result["分析"]["B2"].font.name == "Yu Gothic"
        assert result["分析"]["B2"].font.size == 15
        assert result["分析"]["B2"].font.color.rgb == "00C00000"
        assert result["分析"].column_dimensions["B"].width == 12
        assert result["分析"].row_dimensions[2].height == 24
    finally:
        result.close()


# ------------------------------------------------------------------
# 書式テンプレートの移植(利用者がExcelで整えた書式を出力へ反映する)
# ------------------------------------------------------------------


COMMENT = "《入院》\nⅠ．入院収入は増加している。\n《外来》\nⅠ．外来収入も増加。\n《まとめ》\n・良好。"


def _style_template_file(tmp_path: Path, rich: bool = False, decorated: bool = False) -> Path:
    """
    comment_template 名前付きセルを持つひな型Excelを作る。

    decorated=True のときは罫線と塗りつぶしも設定する(利用者がExcelの画面で
    枠と色を付けた状態を再現する)。
    """
    from openpyxl.cell.rich_text import CellRichText, TextBlock
    from openpyxl.cell.text import InlineFont
    from openpyxl.styles import Alignment, Border, PatternFill, Side
    from openpyxl.workbook.defined_name import DefinedName

    name = "style_template_rich.xlsx" if rich else "style_template.xlsx"
    if decorated:
        name = "style_template_decorated.xlsx"
    path = tmp_path / name
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "コメント"
    if rich:
        # 利用者の実物と同じ形: 見出しだけ太字、本文は通常
        bold = InlineFont(rFont="ＭＳ Ｐゴシック", sz=14, b=True)
        plain = InlineFont(rFont="ＭＳ Ｐゴシック", sz=14, b=False)
        sheet["A1"] = CellRichText(
            [
                TextBlock(bold, "《入院》"),
                TextBlock(plain, "\n見本の本文がここに入る。長いほうが本文とみなされる。"),
                TextBlock(bold, "《まとめ》"),
                TextBlock(plain, "\n見本のまとめ。"),
            ]
        )
    else:
        sheet["A1"] = "書式見本"
    sheet["A1"].font = Font(name="ＭＳ Ｐゴシック", size=14, bold=True, color="2F5597")
    sheet["A1"].alignment = Alignment(wrap_text=True, vertical="top")
    if decorated:
        sheet["A1"].border = Border(
            left=Side(style="medium", color="C00000"),
            right=Side(style="medium", color="C00000"),
            top=Side(style="dashed"),
            bottom=Side(style="double"),
        )
        sheet["A1"].fill = PatternFill("solid", fgColor="D9E1F2")
    workbook.defined_names["comment_template"] = DefinedName(
        "comment_template", attr_text="コメント!$A$1"
    )
    workbook.save(path)
    workbook.close()
    return path


@pytest.fixture(autouse=True)
def _no_repo_template(monkeypatch):
    """テストが実物のひな型を拾わないようにする(既定パスを無効化)。"""
    import office_excel

    monkeypatch.setattr(
        office_excel, "DEFAULT_STYLE_TEMPLATE", Path("存在しないひな型.xlsx")
    )


def test_style_template_font_is_applied_to_written_cells(
    tmp_path: Path, monkeypatch, capsys
) -> None:
    """ひな型Excelの書式が、書き込んだコメント欄のセルに移る。"""
    monkeypatch.chdir(tmp_path)
    source_dir = tmp_path / "input"
    source_dir.mkdir()
    source = source_dir / "source.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "分析"
    sheet.merge_cells("B2:F5")
    sheet["B2"].font = Font(name="Yu Gothic", size=11)
    sheet["A1"] = "見出し"
    sheet["A1"].font = Font(name="Yu Gothic", size=9)
    workbook.save(source)
    template = _style_template_file(tmp_path)

    cmd_set_batch(
        source,
        updates_json=json.dumps(
            [{"sheet": "分析", "cell": "B2", "value": COMMENT}], ensure_ascii=False
        ),
        style_template=str(template),
    )

    result = load_workbook(tmp_path / "output" / "source_updated.xlsx")
    try:
        written = result["分析"]["B2"]
        assert written.value == COMMENT
        assert written.font.name == "ＭＳ Ｐゴシック"
        assert written.font.size == 14
        assert written.font.bold is True
        assert written.font.color.rgb == "002F5597"
        assert written.alignment.wrap_text is True
        assert written.alignment.vertical == "top"
        # 書き込んでいないセルの書式は巻き込まない
        assert result["分析"]["A1"].font.name == "Yu Gothic"
        assert result["分析"]["A1"].font.size == 9
    finally:
        result.close()
    assert "書式テンプレートを適用します" in capsys.readouterr().out


def test_missing_style_name_still_writes_text_and_explains(
    tmp_path: Path, monkeypatch, capsys
) -> None:
    """名前付きセルが無いひな型でも、本文の書込みまで止めてはいけない。"""
    monkeypatch.chdir(tmp_path)
    source_dir = tmp_path / "input"
    source_dir.mkdir()
    source = source_dir / "source.xlsx"
    workbook = Workbook()
    workbook.active.title = "分析"
    workbook.save(source)
    template = tmp_path / "no_name.xlsx"
    Workbook().save(template)

    cmd_set_batch(
        source,
        updates_json='[{"sheet":"分析","cell":"B2","value":"コメント本文"}]',
        style_template=str(template),
    )

    result = load_workbook(tmp_path / "output" / "source_updated.xlsx")
    try:
        assert result["分析"]["B2"].value == "コメント本文"
    finally:
        result.close()
    printed = capsys.readouterr().out
    assert "書式テンプレートを適用できません" in printed
    assert "comment_template" in printed


def test_style_template_keeps_borders_of_the_target_cell(
    tmp_path: Path, monkeypatch
) -> None:
    """移植するのはフォント・配置だけ。分析表の罫線・塗りは元のまま残す。"""
    from openpyxl.styles import Border, PatternFill, Side

    monkeypatch.chdir(tmp_path)
    source_dir = tmp_path / "input"
    source_dir.mkdir()
    source = source_dir / "source.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "分析"
    sheet["B2"].border = Border(top=Side(style="thin"), left=Side(style="double"))
    sheet["B2"].fill = PatternFill("solid", fgColor="FFF2CC")
    workbook.save(source)
    template = _style_template_file(tmp_path)

    cmd_set_batch(
        source,
        updates_json=json.dumps(
            [{"sheet": "分析", "cell": "B2", "value": COMMENT}], ensure_ascii=False
        ),
        style_template=str(template),
    )

    result = load_workbook(tmp_path / "output" / "source_updated.xlsx")
    try:
        written = result["分析"]["B2"]
        assert written.font.name == "ＭＳ Ｐゴシック"
        assert written.border.top.style == "thin"
        assert written.border.left.style == "double"
        assert written.fill.fgColor.rgb == "00FFF2CC"
    finally:
        result.close()


def test_見出しだけ太字のひな型は出力でも見出しだけ太字になる(
    tmp_path: Path, monkeypatch
) -> None:
    """
    利用者は《入院》などの見出しだけを太字にして使っている。

    セル全体のフォントしか見ないと、この太字が丸ごと落ちる。
    「太字にしてるところもなってない」の再発防止(2026-08-17)。
    """
    from openpyxl.cell.rich_text import CellRichText

    monkeypatch.chdir(tmp_path)
    source_dir = tmp_path / "input"
    source_dir.mkdir()
    source = source_dir / "source.xlsx"
    workbook = Workbook()
    workbook.active.title = "分析"
    workbook.save(source)
    template = _style_template_file(tmp_path, rich=True)

    cmd_set_batch(
        source,
        updates_json=json.dumps(
            [{"sheet": "分析", "cell": "B2", "value": COMMENT}], ensure_ascii=False
        ),
        style_template=str(template),
    )

    result = load_workbook(tmp_path / "output" / "source_updated.xlsx", rich_text=True)
    try:
        written = result["分析"]["B2"]
        assert isinstance(written.value, CellRichText), "部分書式が失われている"
        bold = {str(b).strip() for b in written.value if getattr(getattr(b, "font", None), "b", False)}
        assert "《入院》" in bold and "《まとめ》" in bold
        # 本文まで太字にしてはいけない
        assert not any("入院収入" in text for text in bold)
        assert "".join(str(b) for b in written.value) == COMMENT
    finally:
        result.close()


def test_コメント以外の書込みには書式テンプレートを当てない(
    tmp_path: Path, monkeypatch
) -> None:
    """
    excel_write は任意のセルに書ける汎用ツール。

    数値や短文にまでコメント用のフォントを当てると、表が壊れる。
    """
    monkeypatch.chdir(tmp_path)
    source_dir = tmp_path / "input"
    source_dir.mkdir()
    source = source_dir / "source.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "分析"
    sheet["B2"].font = Font(name="游ゴシック", size=9)
    workbook.save(source)
    template = _style_template_file(tmp_path)

    cmd_set_batch(
        source,
        updates_json='[{"sheet":"分析","cell":"B2","value":"123という短い値"}]',
        style_template=str(template),
    )

    result = load_workbook(tmp_path / "output" / "source_updated.xlsx")
    try:
        assert result["分析"]["B2"].font.name == "游ゴシック"
        assert result["分析"]["B2"].font.size == 9
    finally:
        result.close()


def test_setでも書式テンプレートが効く(tmp_path: Path, monkeypatch) -> None:
    """
    Kiloの excel_write は set を通る。

    set_batch にしか対応していなかったため、Kilo経由だと書式が
    まったく効かなかった(2026-08-17に利用者から報告)。
    """
    monkeypatch.chdir(tmp_path)
    source_dir = tmp_path / "input"
    source_dir.mkdir()
    source = source_dir / "source.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "分析"
    sheet["B2"].font = Font(name="游ゴシック", size=16)
    workbook.save(source)
    template = _style_template_file(tmp_path)

    cmd_set(source, "分析", "B2", COMMENT, None, str(template))

    result = load_workbook(tmp_path / "output" / "source_updated.xlsx")
    try:
        written = result["分析"]["B2"]
        assert written.value == COMMENT
        assert written.font.name == "ＭＳ Ｐゴシック"
        assert written.font.size == 14
    finally:
        result.close()


def test_既定のひな型が無くても書込みは通る(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.chdir(tmp_path)
    source_dir = tmp_path / "input"
    source_dir.mkdir()
    source = source_dir / "source.xlsx"
    workbook = Workbook()
    workbook.active.title = "分析"
    workbook.save(source)

    cmd_set(source, "分析", "B2", COMMENT, None)

    result = load_workbook(tmp_path / "output" / "source_updated.xlsx")
    try:
        assert result["分析"]["B2"].value == COMMENT
    finally:
        result.close()


def _write_comment_with(template: Path, tmp_path: Path) -> Path:
    """罫線と塗りを持つ分析表へ、ひな型を当ててコメントを書き込む。"""
    from openpyxl.styles import Border, PatternFill, Side

    source_dir = tmp_path / "input"
    source_dir.mkdir(exist_ok=True)
    source = source_dir / "source.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "分析"
    sheet["B2"].border = Border(top=Side(style="thin"), left=Side(style="double"))
    sheet["B2"].fill = PatternFill("solid", fgColor="FFF2CC")
    workbook.save(source)
    workbook.close()

    cmd_set_batch(
        source,
        updates_json=json.dumps(
            [{"sheet": "分析", "cell": "B2", "value": COMMENT}], ensure_ascii=False
        ),
        style_template=str(template),
    )
    return tmp_path / "output" / "source_updated.xlsx"


def test_ひな型に引いた罫線が書き込み先へ移る(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.chdir(tmp_path)
    result = load_workbook(_write_comment_with(_style_template_file(tmp_path, decorated=True), tmp_path))
    try:
        written = result["分析"]["B2"]
        assert written.border.left.style == "medium"
        assert written.border.right.style == "medium"
        assert written.border.top.style == "dashed", "書き込み先のthinが残っている"
        assert written.border.bottom.style == "double"
        assert written.border.left.color.rgb == "00C00000", "罫線の色が移っていない"
        # 罫線を移しても文字の書式は今までどおり移ること
        assert written.font.name == "ＭＳ Ｐゴシック"
    finally:
        result.close()


def test_ひな型の塗りつぶしが書き込み先へ移る(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.chdir(tmp_path)
    result = load_workbook(_write_comment_with(_style_template_file(tmp_path, decorated=True), tmp_path))
    try:
        written = result["分析"]["B2"]
        assert written.fill.patternType == "solid"
        assert written.fill.fgColor.rgb == "00D9E1F2", "書き込み先のFFF2CCが残っている"
    finally:
        result.close()


def test_罫線を移しても他のセルの見た目は変わらない(tmp_path: Path, monkeypatch) -> None:
    """追記方式なので、既存の定義を書き換えて巻き添えにしてはいけない。"""
    from openpyxl.styles import Border, PatternFill, Side

    monkeypatch.chdir(tmp_path)
    source_dir = tmp_path / "input"
    source_dir.mkdir()
    source = source_dir / "source.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "分析"
    sheet["B2"].border = Border(top=Side(style="thin"))
    sheet["C5"].value = "隣の表"
    sheet["C5"].border = Border(bottom=Side(style="hair"))
    sheet["C5"].fill = PatternFill("solid", fgColor="E2EFDA")
    workbook.save(source)
    workbook.close()

    cmd_set_batch(
        source,
        updates_json=json.dumps(
            [{"sheet": "分析", "cell": "B2", "value": COMMENT}], ensure_ascii=False
        ),
        style_template=str(_style_template_file(tmp_path, decorated=True)),
    )

    result = load_workbook(tmp_path / "output" / "source_updated.xlsx")
    try:
        neighbour = result["分析"]["C5"]
        assert neighbour.border.bottom.style == "hair", "隣のセルの罫線が変わった"
        assert neighbour.fill.fgColor.rgb == "00E2EFDA", "隣のセルの塗りが変わった"
        assert neighbour.value == "隣の表"
    finally:
        result.close()


def test_ひな型に罫線も塗りも無ければ書き込み先のものを残す(tmp_path: Path, monkeypatch) -> None:
    """既存のひな型(枠を引いていない)の挙動を変えないことの確認。"""
    monkeypatch.chdir(tmp_path)
    result = load_workbook(_write_comment_with(_style_template_file(tmp_path), tmp_path))
    try:
        written = result["分析"]["B2"]
        assert written.border.top.style == "thin"
        assert written.border.left.style == "double"
        assert written.fill.fgColor.rgb == "00FFF2CC"
    finally:
        result.close()


def test_適用した書式の内訳を画面に出す(tmp_path: Path, capsys) -> None:
    """ひな型で引いたのに反映されない、を黙って起こさないための表示。"""
    import office_excel

    office_excel.resolve_style_template(str(_style_template_file(tmp_path, decorated=True)))
    decorated = capsys.readouterr().out
    assert "罫線" in decorated and "塗りつぶし" in decorated

    office_excel.resolve_style_template(str(_style_template_file(tmp_path)))
    plain = capsys.readouterr().out
    assert "罫線" not in plain and "塗りつぶし" not in plain
