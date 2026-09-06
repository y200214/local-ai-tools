from pathlib import Path
from types import SimpleNamespace
import pandas as pd
from openpyxl import Workbook

import pytest
import excel_reader

from excel_reader import _visible_sheet_names, find_values, print_index, print_range


def test_visible_sheet_names_excludes_hidden_and_very_hidden(tmp_path: Path) -> None:
    path = tmp_path / "sheets.xlsm"
    workbook = Workbook()
    workbook.active.title = "表示"
    workbook.create_sheet("非表示").sheet_state = "hidden"
    workbook.create_sheet("VeryHidden").sheet_state = "veryHidden"
    workbook.save(path)

    with pd.ExcelFile(path) as xls:
        assert _visible_sheet_names(xls) == ["表示"]


def _sample_book(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "集計"
    sheet["A1"] = "診療科"
    sheet["B2"] = "内科"
    sheet["C2"] = 42
    workbook.create_sheet("非表示").sheet_state = "hidden"
    workbook.save(path)


def test_index_prints_structure_without_values(tmp_path: Path, capsys) -> None:
    path = tmp_path / "book.xlsm"
    _sample_book(path)

    print_index(path)

    output = capsys.readouterr().out
    assert "集計" in output
    assert "A1:C2" in output
    assert "内科" not in output
    assert "非表示" not in output


def test_index_handles_sheet_without_dimension(monkeypatch, capsys) -> None:
    class EmptySheet:
        title = "空シート"
        max_row = None
        max_column = None
        sheet_state = "visible"

        def reset_dimensions(self) -> None:
            pass

        def calculate_dimension(self, *, force: bool) -> str:
            assert force is True
            self.max_row = 4
            self.max_column = 3
            return "A1:C4"

    class EmptyWorkbook:
        worksheets = [EmptySheet()]

        @staticmethod
        def close() -> None:
            pass

    monkeypatch.setattr(excel_reader, "_open_workbook", lambda path: EmptyWorkbook())

    excel_reader.print_index(Path("empty.xlsx"))

    output = capsys.readouterr().out
    assert "空シート: A1:C4 (4 行 × 3 列, visible)" in output


def test_find_returns_coordinates_and_honors_limit(tmp_path: Path, capsys) -> None:
    path = tmp_path / "book.xlsm"
    _sample_book(path)

    find_values(path, "内科", max_results=1)

    output = capsys.readouterr().out
    assert "集計!B2" in output
    assert "内科" in output
    assert "結果上限 1 件" in output


def test_find_can_include_values_below_match(tmp_path: Path, capsys) -> None:
    path = tmp_path / "book.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet["B2"] = "コメント"
    sheet["B3"] = "記入例"
    workbook.save(path)

    find_values(path, "コメント", context_rows=2)

    output = capsys.readouterr().out
    assert "Sheet!B2" in output
    assert "直下1行: B3\t記入例" in output


def test_range_prints_coordinates_and_enforces_cell_limit(tmp_path: Path, capsys) -> None:
    path = tmp_path / "book.xlsm"
    _sample_book(path)

    print_range(path, "集計", "B2:C2", max_cells=2)
    output = capsys.readouterr().out
    assert "B2=内科" in output
    assert "C2=42" in output

    with pytest.raises(SystemExit, match="上限 2 セル"):
        print_range(path, "集計", "A1:C2", max_cells=2)


def test_range_handles_empty_cells_without_coordinate(monkeypatch, capsys) -> None:
    class Sheet:
        @staticmethod
        def iter_rows(**kwargs):
            return [[SimpleNamespace(value=None), SimpleNamespace(value="値")]]

    class Book:
        sheetnames = ["集計"]

        def __getitem__(self, name: str):
            assert name == "集計"
            return Sheet()

        @staticmethod
        def close() -> None:
            pass

    monkeypatch.setattr(excel_reader, "_open_workbook", lambda path, data_only=False: Book())

    excel_reader.print_range(Path("book.xlsx"), "集計", "B2:C2")

    output = capsys.readouterr().out
    assert "B2=" in output
    assert "C2=値" in output


def test_compact_range_omits_empty_cells(tmp_path: Path, capsys) -> None:
    path = tmp_path / "book.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"] = "項目"
    sheet["D4"] = 42
    workbook.save(path)

    excel_reader.print_range(path, "Sheet", "A1:D4", max_cells=16, compact=True)

    output = capsys.readouterr().out
    assert "A1=項目" in output
    assert "D4=42" in output
    assert "B1=" not in output
