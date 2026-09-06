from __future__ import annotations
from repo_paths import ROOT as PROJECT_ROOT, TOOLS as TOOLS_ROOT

import subprocess
import sys
from pathlib import Path


RUNNER = (TOOLS_ROOT / "safe_task_runner.py")


def _run(workspace: Path, source: str) -> subprocess.CompletedProcess[str]:
    work = workspace / "work"
    work.mkdir()
    script = work / "task.py"
    script.write_text(source, encoding="utf-8")
    return subprocess.run(
        [sys.executable, str(RUNNER), str(script)],
        cwd=workspace,
        capture_output=True,
        text=True,
        encoding="utf-8",
        check=False,
    )


def test_allows_reading_input_and_writing_output(tmp_path: Path) -> None:
    (tmp_path / "input").mkdir()
    (tmp_path / "input" / "source.txt").write_text("元データ", encoding="utf-8")

    result = _run(
        tmp_path,
        """from pathlib import Path
text = Path('input/source.txt').read_text(encoding='utf-8')
Path('output/result.txt').write_text(text + ' 加工済み', encoding='utf-8')
""",
    )

    assert result.returncode == 0, result.stderr
    assert (tmp_path / "output" / "result.txt").read_text(encoding="utf-8") == "元データ 加工済み"


def test_blocks_writing_input(tmp_path: Path) -> None:
    (tmp_path / "input").mkdir()
    source = tmp_path / "input" / "source.txt"
    source.write_text("元データ", encoding="utf-8")

    result = _run(
        tmp_path,
        "from pathlib import Path\nPath('input/source.txt').write_text('破壊', encoding='utf-8')\n",
    )

    assert result.returncode != 0
    assert "書き込み" in result.stderr
    assert source.read_text(encoding="utf-8") == "元データ"


def test_blocks_deleting_input(tmp_path: Path) -> None:
    (tmp_path / "input").mkdir()
    source = tmp_path / "input" / "source.txt"
    source.write_text("元データ", encoding="utf-8")

    result = _run(
        tmp_path,
        "from pathlib import Path\nPath('input/source.txt').unlink()\n",
    )

    assert result.returncode != 0
    assert "os.remove" in result.stderr
    assert source.read_text(encoding="utf-8") == "元データ"


def test_blocks_sensitive_reads(tmp_path: Path) -> None:
    (tmp_path / ".env").write_text("SECRET=test", encoding="utf-8")

    result = _run(
        tmp_path,
        "from pathlib import Path\nprint(Path('.env').read_text(encoding='utf-8'))\n",
    )

    assert result.returncode != 0
    assert "機密ファイル" in result.stderr
    assert "SECRET=test" not in result.stdout


def test_blocks_subprocesses(tmp_path: Path) -> None:
    result = _run(
        tmp_path,
        "import subprocess\nsubprocess.run(['python', '--version'], check=True)\n",
    )

    assert result.returncode != 0
    assert "subprocess.Popen" in result.stderr


def test_can_edit_excel_copy_with_openpyxl(tmp_path: Path) -> None:
    from openpyxl import Workbook, load_workbook

    input_dir = tmp_path / "input"
    input_dir.mkdir()
    source = input_dir / "source.xlsx"
    workbook = Workbook()
    workbook.active["A1"] = "変更前"
    workbook.save(source)

    result = _run(
        tmp_path,
        """from openpyxl import load_workbook
book = load_workbook('input/source.xlsx')
book.active['A1'] = '変更後'
book.save('output/result.xlsx')
""",
    )

    assert result.returncode == 0, result.stderr
    assert load_workbook(source).active["A1"].value == "変更前"
    assert load_workbook(tmp_path / "output" / "result.xlsx").active["A1"].value == "変更後"


def test_can_edit_word_copy_with_python_docx(tmp_path: Path) -> None:
    from docx import Document

    input_dir = tmp_path / "input"
    input_dir.mkdir()
    source = input_dir / "source.docx"
    document = Document()
    document.add_paragraph("変更前")
    document.save(source)

    result = _run(
        tmp_path,
        """from docx import Document
document = Document('input/source.docx')
document.paragraphs[0].text = '変更後'
document.save('output/result.docx')
""",
    )

    assert result.returncode == 0, result.stderr
    assert Document(source).paragraphs[0].text == "変更前"
    assert Document(tmp_path / "output" / "result.docx").paragraphs[0].text == "変更後"
