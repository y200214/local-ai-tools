"""toolpack_runner(専用ランナー第一段階)の統合テスト。

監査フックはプロセス単位でしか効かないため、実際に隔離子プロセスを起動して確かめる。
禁止操作の試験は、tmp_path 内に作った**架空の秘密ファイル**だけで行う。
本物の data/・templates/・.env・管理APIキーには一切触れない。
"""

from __future__ import annotations

import json
import sys
import tempfile
from pathlib import Path

import pytest

import toolpack_runner as runner

# Windows のジョブオブジェクト前提。この環境は win32 固定だが保険で明示
pytestmark = pytest.mark.skipif(sys.platform != "win32", reason="Windows専用")


def run(
    tmp_path: Path, main_src: str, *, tool_over: dict | None = None,
    inputs=None, config=None, instruction="やって",
    isolation: str = runner.ISOLATION_LOW,
) -> runner.RunReport:
    # 1テスト内で複数回呼んでも衝突しないよう、毎回ユニークな作業領域を作る
    base = Path(tempfile.mkdtemp(dir=tmp_path))
    package = base / "pkg"
    package.mkdir()
    (package / "main.py").write_text(main_src, encoding="utf-8")
    tool_json = {
        "permissions": {"ollama": False},
        "outputs": {"produces": [".txt"], "may_be_empty": True},
        "timeout_seconds": 60,
    }
    if tool_over:
        tool_json.update(tool_over)
    return runner.run_tool_package(
        package, tool_json, inputs=inputs or [], instruction=instruction,
        config=config, job_root=base / "job", isolation=isolation,
    )


def assert_blocked(report: runner.RunReport) -> None:
    """禁止操作が監査フックで弾かれたことを確かめる。

    attempt() のツールは、PermissionError を捕まえたときだけ status ok(message=blocked)を返す。
    弾かれなければ非ゼロ終了(E-TOOL)になる。
    """
    assert report.ok and report.result is not None and report.result.message == "blocked", (
        f"ブロックされなかった: ok={report.ok} code={report.code} detail={report.detail[:200]}"
    )


# 成果物なしで ok を返す最小ツール
BENIGN = '''\
import argparse, json
p = argparse.ArgumentParser(); p.add_argument("--request"); a = p.parse_args()
json.loads(open(a.request, encoding="utf-8").read())
print(json.dumps({"status": "ok", "message": "done", "files": [], "skipped": [], "notes": []}))
'''


def attempt(body: str) -> str:
    """禁止操作 body を試し、PermissionError なら ok を返すツール。

    ブロックされずに通ってしまったら、非ゼロ終了でそれを知らせる。
    """
    return (
        "import argparse, json, os, sys\n"
        "p = argparse.ArgumentParser(); p.add_argument('--request'); a = p.parse_args()\n"
        "try:\n"
        f"    {body}\n"
        "except PermissionError:\n"
        "    print(json.dumps({'status':'ok','message':'blocked','files':[],'skipped':[],'notes':[]})); sys.exit(0)\n"
        "sys.stderr.write('NOT BLOCKED'); sys.exit(7)\n"
    )


# ---------------------------------------------------------------------------
# 正常系・セルフチェック
# ---------------------------------------------------------------------------
def test_正常なツールが成果物を返す(tmp_path) -> None:
    main = '''\
import argparse, json, os
p = argparse.ArgumentParser(); p.add_argument("--request"); a = p.parse_args()
req = json.loads(open(a.request, encoding="utf-8").read())
out = os.path.join(os.path.dirname(a.request), req["outdir"], "result.txt")
open(out, "w", encoding="utf-8").write("done")
print(json.dumps({"status":"ok","message":"ok","files":["result.txt"],"skipped":[],"notes":[]}))
'''
    report = run(tmp_path, main)
    assert report.ok, report.detail
    assert [p.name for p in report.result.files] == ["result.txt"]


def test_self_checkが両方式で真を返す() -> None:
    assert runner.self_check() is True  # 既定(低整合性)
    assert runner.self_check(isolation=runner.ISOLATION_BASIC) is True


# ---------------------------------------------------------------------------
# 架空秘密ファイルは読めない(実データには触れない試験)
# ---------------------------------------------------------------------------
def test_架空秘密ファイルを読めない(tmp_path) -> None:
    secrets_root = tmp_path / "outside"
    fakes = {
        "fake-data/patient.txt": "架空の患者データ",
        "fake-home/admin_api_key.txt": "架空のキー",
        "fake-project/.env": "SECRET=xxx",
        "fake-outside/secret.txt": "架空の秘密",
    }
    for rel, text in fakes.items():
        path = secrets_root / rel
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(text, encoding="utf-8")

    for rel in fakes:
        target = (secrets_root / rel).resolve()
        report = run(tmp_path, attempt(f"open(r'{target}', encoding='utf-8').read()"))
        assert_blocked(report)  # 読み取りが弾かれた(架空秘密は読めない)


# ---------------------------------------------------------------------------
# 書き込み・入力保護・リンク
# ---------------------------------------------------------------------------
def test_outdir外への書き込みを拒否する(tmp_path) -> None:
    target = (tmp_path / "escape.txt").resolve()
    report = run(tmp_path, attempt(f"open(r'{target}', 'w').write('x')"))
    assert_blocked(report)
    assert not target.exists()


def test_入力コピーの書き換えを拒否する(tmp_path) -> None:
    src = tmp_path / "src.txt"
    src.write_text("original", encoding="utf-8")
    main = attempt(
        "req = json.loads(open(a.request, encoding='utf-8').read());"
        " ip = os.path.join(os.path.dirname(a.request), req['inputs'][0]['path']);"
        " open(ip, 'w').write('tampered')"
    )
    report = run(tmp_path, main, inputs=[("in-1", None, src, "src.txt")])
    assert_blocked(report)
    assert src.read_text(encoding="utf-8") == "original"  # 元ファイルも無事


def test_リンク作成を拒否する(tmp_path) -> None:
    report = run(tmp_path, attempt(
        "req = json.loads(open(a.request, encoding='utf-8').read());"
        " out = os.path.join(os.path.dirname(a.request), req['outdir'], 'link.txt');"
        " os.link(a.request, out)"
    ))
    assert_blocked(report)


# ---------------------------------------------------------------------------
# 子プロセス・通信
# ---------------------------------------------------------------------------
def test_子プロセス起動を拒否する(tmp_path) -> None:
    report = run(tmp_path, attempt("import subprocess; subprocess.run(['cmd', '/c', 'echo hi'])"))
    assert_blocked(report)


def test_通信を拒否する(tmp_path) -> None:
    report = run(tmp_path, attempt("import socket; socket.create_connection(('192.0.2.1', 80), timeout=1)"))
    assert_blocked(report)


def test_ollama許可でも別ポートは拒否する(tmp_path) -> None:
    report = run(
        tmp_path,
        attempt("import socket; socket.create_connection(('127.0.0.1', 9999), timeout=1)"),
        tool_over={"permissions": {"ollama": True}, "timeout_seconds": 60},
    )
    assert_blocked(report)


# ---------------------------------------------------------------------------
# 契約違反・タイムアウト・env
# ---------------------------------------------------------------------------
def test_JSON以外の標準出力を拒否する(tmp_path) -> None:
    report = run(tmp_path, "print('ただのテキスト')")
    assert not report.ok
    assert report.code in ("E-JSON", "E-MULTILINE")


def test_成果物を出さずexit0を拒否する(tmp_path) -> None:
    report = run(tmp_path, "import sys; sys.exit(0)")
    assert not report.ok and report.code == "E-EMPTY"


def test_タイムアウトで打ち切る(tmp_path) -> None:
    report = run(tmp_path, "while True:\n    pass", tool_over={"timeout_seconds": 2})
    assert not report.ok and report.code == "E-TIMEOUT"


def test_最小envに鍵を渡さない(tmp_path, monkeypatch) -> None:
    # 親に架空の鍵を置いても、子の環境には渡らないことを確認する
    monkeypatch.setenv("OPEN_WEBUI_ADMIN_API_KEY", "FAKE-KEY-SHOULD-NOT-LEAK")
    monkeypatch.setenv("USERPROFILE", str(tmp_path / "home"))
    main = '''\
import argparse, json, os
p = argparse.ArgumentParser(); p.add_argument("--request"); a = p.parse_args()
req = json.loads(open(a.request, encoding="utf-8").read())
out = os.path.join(os.path.dirname(a.request), req["outdir"], "env.txt")
open(out, "w", encoding="utf-8").write(os.environ.get("OPEN_WEBUI_ADMIN_API_KEY", "ABSENT"))
print(json.dumps({"status":"ok","message":"ok","files":["env.txt"],"skipped":[],"notes":[]}))
'''
    report = run(tmp_path, main)
    assert report.ok, report.detail
    assert report.result.files[0].read_text(encoding="utf-8") == "ABSENT"


def test_最小envの構築で鍵名の変数を落とす(tmp_path) -> None:
    env = runner.build_min_env(tmp_path / "tmp", tmp_path / "mpl")
    assert not any(
        any(bad in name.upper() for bad in runner._ENV_DENY_SUBSTRINGS) for name in env
    )
    assert "MPLCONFIGDIR" in env  # フォントキャッシュはジョブ内へ


# ---------------------------------------------------------------------------
# 段階2(制限トークン + 低整合性)— 監査フックを迂回してもOSが止めるか
# ---------------------------------------------------------------------------
# 監査フックは Python レベルのI/Oしか見えない。ctypes で CreateFileW を直接呼べば
# フックは素通りする。段階1ではそれで書けてしまい、段階2では OS が拒否する。
# この差が工程4の存在理由そのものなので、両方を突き合わせて確かめる。
CTYPES_ESCAPE = '''\
import argparse, json, ctypes
from ctypes import wintypes
p = argparse.ArgumentParser(); p.add_argument("--request"); p.parse_args()
k = ctypes.WinDLL("kernel32", use_last_error=True)
k.CreateFileW.restype = wintypes.HANDLE
k.CreateFileW.argtypes = [wintypes.LPCWSTR, wintypes.DWORD, wintypes.DWORD,
                          ctypes.c_void_p, wintypes.DWORD, wintypes.DWORD, wintypes.HANDLE]
k.WriteFile.argtypes = [wintypes.HANDLE, ctypes.c_char_p, wintypes.DWORD,
                        ctypes.POINTER(wintypes.DWORD), ctypes.c_void_p]
handle = k.CreateFileW(r"{target}", 0x40000000, 0, None, 2, 0x80, None)
invalid = ctypes.cast(ctypes.c_void_p(-1), ctypes.c_void_p).value
opened = handle is not None and handle != invalid
if opened:
    written = wintypes.DWORD()
    k.WriteFile(wintypes.HANDLE(handle), b"escaped", 7, ctypes.byref(written), None)
    k.CloseHandle(wintypes.HANDLE(handle))
print(json.dumps({{"status": "ok", "message": "opened=%s" % opened,
                   "files": [], "skipped": [], "notes": []}}))
'''


def test_段階1ではctypes直呼びの書き込みを止められない(tmp_path) -> None:
    # 段階2が要る理由の実証(この穴があるから低整合性を重ねる)
    target = (tmp_path / "escaped_basic.txt").resolve()
    report = run(
        tmp_path, CTYPES_ESCAPE.format(target=target),
        isolation=runner.ISOLATION_BASIC,
    )
    assert report.ok, report.detail
    assert target.exists(), "段階1でctypes書き込みが素通りする前提が崩れている"


def test_段階2ではctypes直呼びでもOSが書き込みを拒否する(tmp_path) -> None:
    target = (tmp_path / "escaped_low.txt").resolve()
    report = run(tmp_path, CTYPES_ESCAPE.format(target=target))  # 既定=低整合性
    assert report.ok, report.detail
    assert report.result.message == "opened=False"
    assert not target.exists(), "低整合性でも outdir 外へ書けてしまった"


def test_低整合性でopenpyxlとmatplotlibが使える(tmp_path) -> None:
    # 低整合性は書き込みだけを縛る。site-packages からの import は通る必要がある
    main = '''\
import argparse, json, os
p = argparse.ArgumentParser(); p.add_argument("--request"); a = p.parse_args()
req = json.loads(open(a.request, encoding="utf-8").read())
out_dir = os.path.join(os.path.dirname(a.request), req["outdir"])
import openpyxl
import matplotlib; matplotlib.use("Agg")
import matplotlib.pyplot as plt, matplotlib_fontja
fig, ax = plt.subplots(); ax.set_title("日本語")
fig.savefig(os.path.join(out_dir, "chart.png"))
book = openpyxl.Workbook(); book.active["A1"] = "値"
book.save(os.path.join(out_dir, "book.xlsx"))
print(json.dumps({"status":"ok","message":"ok","files":["chart.png","book.xlsx"],
                  "skipped":[],"notes":[]}))
'''
    report = run(
        tmp_path, main,
        tool_over={"outputs": {"produces": [".png", ".xlsx"], "may_be_empty": False},
                   "timeout_seconds": 180},
    )
    assert report.ok, report.detail
    assert sorted(p.name for p in report.result.files) == ["book.xlsx", "chart.png"]


def test_未知の隔離方式を拒否する(tmp_path) -> None:
    with pytest.raises(runner.RunnerError) as caught:
        run(tmp_path, BENIGN, isolation="none")
    assert caught.value.code == "E-ISOLATION"


def test_段階1でも監査フックの拒否は効く(tmp_path) -> None:
    # 切戻し用の段階1でも、Python経由の逸脱は従来どおり止まる
    target = (tmp_path / "escape_basic_python.txt").resolve()
    report = run(
        tmp_path, attempt(f"open(r'{target}', 'w').write('x')"),
        isolation=runner.ISOLATION_BASIC,
    )
    assert_blocked(report)
    assert not target.exists()


def test_ジョブを閉じると子プロセスが死ぬ() -> None:
    # KILL_ON_JOB_CLOSE の保証を、proc.kill のフォールバック無しで確かめる。
    # venv の python.exe はリダイレクタで孫プロセスを生むため、ランナーと同じく
    # ベースインタプリタを直接使う(1プロセスで判定を素直にする)
    import subprocess
    import toolpack_winjob

    interpreter, _ = runner.resolve_interpreter(runner.VENV_PYTHON)
    proc = subprocess.Popen(
        [str(interpreter), "-c", "import time; time.sleep(30)"],
        creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0),
    )
    job = toolpack_winjob.create_job()
    job.assign(int(proc._handle))
    assert proc.poll() is None  # close 前は sleep 中で生きている
    job.close()  # kill せずジョブを閉じるだけ
    # sleep(30) の途中でも即座に落ちる。TimeoutExpired が出れば殺せていない
    # (ジョブ終了時の終了コードは0になりうるので、値ではなく「即死したか」で見る)
    proc.wait(timeout=10)


def test_短縮名のジョブでもresolveした先へ書ける(tmp_path) -> None:
    r"""8.3 短縮名(C:\Users\KELCRE~1\...)のジョブ領域でも許可が効くこと。

    許可した領域を短縮名のまま覚えていると、ツール側が Path.resolve() した
    長い形と文字列が一致せず、**許可したはずの書き込みが拒否される**。
    pytest は basetemp を resolve するため、これで tmp_path を使う
    パッケージ同梱テストが全部落ちていた(2026-08-31)。
    """
    src = """
import argparse, json, os
from pathlib import Path
p = argparse.ArgumentParser(); p.add_argument("--request"); a = p.parse_args()
request = json.loads(Path(a.request).read_text(encoding="utf-8"))
job = Path(a.request).parent
# ツール側が素直に resolve() してから書く、というよくある書き方
out = (job / request["outdir"]).resolve()
(out / "結果.txt").write_text("ok", encoding="utf-8")
tmp = Path(os.environ["TEMP"]).resolve()
(tmp / "作業.txt").write_text("ok", encoding="utf-8")
print(json.dumps({"status": "ok", "message": "書けた", "files": ["結果.txt"],
                  "skipped": [], "notes": []}, ensure_ascii=False))
"""
    # 短縮名を含む領域を用意する(この環境の TEMP がまさにその形)
    short = Path(tempfile.mkdtemp(prefix="short_"))
    package = short / "pkg"
    package.mkdir()
    (package / "main.py").write_text(src, encoding="utf-8")
    report = runner.run_tool_package(
        package,
        {"permissions": {"ollama": False},
         "outputs": {"produces": [".txt"], "may_be_empty": False},
         "timeout_seconds": 60},
        inputs=[], instruction="書いて", job_root=short / "job",
    )
    assert report.ok, f"code={report.code} detail={report.detail[:400]}"


def test_ジョブ領域は長い形で子へ渡す(tmp_path) -> None:
    src = """
import argparse, json, os
from pathlib import Path
p = argparse.ArgumentParser(); p.add_argument("--request"); a = p.parse_args()
job = Path(a.request).parent
same = str(job) == str(job.resolve())
print(json.dumps({"status": "ok", "message": "同じ" if same else f"違う: {job}",
                  "files": [], "skipped": [], "notes": []}, ensure_ascii=False))
"""
    short = Path(tempfile.mkdtemp(prefix="short_"))
    package = short / "pkg"
    package.mkdir()
    (package / "main.py").write_text(src, encoding="utf-8")
    report = runner.run_tool_package(
        package,
        {"permissions": {"ollama": False},
         "outputs": {"produces": [], "may_be_empty": True},
         "timeout_seconds": 60},
        inputs=[], instruction="見て", job_root=short / "job",
    )
    assert report.ok and report.result is not None
    assert report.result.message == "同じ", report.result.message


# ---------------------------------------------------------------------------
# 添付ファイルの読み取り部品(toolpack_textio)
# ---------------------------------------------------------------------------
def test_読み取り部品を隔離の中でimportできる(tmp_path) -> None:
    src = """
import argparse, json
from pathlib import Path
import toolpack_textio
p = argparse.ArgumentParser(); p.add_argument("--request"); a = p.parse_args()
request = json.loads(Path(a.request).read_text(encoding="utf-8"))
text, name = toolpack_textio.read_request_input(request, str(Path(a.request).parent))
print(json.dumps({"status": "ok", "message": f"{name}:{text}", "files": [],
                  "skipped": [], "notes": []}, ensure_ascii=False))
"""
    source = tmp_path / "資料.txt"
    source.write_text("本文です", encoding="utf-8")
    report = run(tmp_path, src, inputs=[("in-1", None, source, "資料.txt")],
                 tool_over={"outputs": {"produces": [], "may_be_empty": True}})
    assert report.ok, f"code={report.code} detail={report.detail[:300]}"
    assert report.result.message == "資料.txt:本文です"


def test_隔離の中でWordを読める(tmp_path) -> None:
    """実物の文字起こしは .docx だった(台帳 5-20)。隔離の中で解析する。"""
    import docx

    document = docx.Document()
    document.add_paragraph("会議の本文")
    source = tmp_path / "議事録.docx"
    document.save(source)

    src = """
import argparse, json
from pathlib import Path
import toolpack_textio
p = argparse.ArgumentParser(); p.add_argument("--request"); a = p.parse_args()
request = json.loads(Path(a.request).read_text(encoding="utf-8"))
text, _ = toolpack_textio.read_request_input(request, str(Path(a.request).parent))
print(json.dumps({"status": "ok", "message": text, "files": [],
                  "skipped": [], "notes": []}, ensure_ascii=False))
"""
    report = run(tmp_path, src, inputs=[("in-1", None, source, "議事録.docx")],
                 tool_over={"outputs": {"produces": [], "may_be_empty": True}})
    assert report.ok, f"code={report.code} detail={report.detail[:300]}"
    assert report.result.message == "会議の本文"


def test_読み取り部品はツールから差し替えられない(tmp_path) -> None:
    """ジョブ直下は読めるが書けない。**書き換えられたら読み取りの意味が無い。**"""
    src = """
import argparse, json
from pathlib import Path
p = argparse.ArgumentParser(); p.add_argument("--request"); a = p.parse_args()
target = Path(a.request).parent / "toolpack_textio.py"
try:
    target.write_text("x", encoding="utf-8")
    message = "書けてしまった"
except PermissionError:
    message = "blocked"
except OSError:
    message = "blocked"
print(json.dumps({"status": "ok", "message": message, "files": [],
                  "skipped": [], "notes": []}, ensure_ascii=False))
"""
    assert_blocked(run(tmp_path, src, tool_over={"outputs": {"produces": [], "may_be_empty": True}}))


def test_同名を同梱してもコア側が優先される(tmp_path) -> None:
    """パッケージが偽の toolpack_textio.py を入れても、コアの方が読まれる。"""
    base = Path(tempfile.mkdtemp(dir=tmp_path))
    package = base / "pkg"
    package.mkdir()
    (package / "toolpack_textio.py").write_text(
        "def read_request_input(request, job_dir):\n    return ('偽物', 'x')\n",
        encoding="utf-8",
    )
    (package / "main.py").write_text("""
import argparse, json
from pathlib import Path
import toolpack_textio
p = argparse.ArgumentParser(); p.add_argument("--request"); a = p.parse_args()
request = json.loads(Path(a.request).read_text(encoding="utf-8"))
text, _ = toolpack_textio.read_request_input(request, str(Path(a.request).parent))
print(json.dumps({"status": "ok", "message": text, "files": [],
                  "skipped": [], "notes": []}, ensure_ascii=False))
""", encoding="utf-8")
    source = base / "資料.txt"
    source.write_text("本物の本文", encoding="utf-8")
    report = runner.run_tool_package(
        package,
        {"permissions": {"ollama": False},
         "outputs": {"produces": [], "may_be_empty": True},
         "timeout_seconds": 60},
        inputs=[("in-1", None, source, "資料.txt")], instruction="読んで",
        job_root=base / "job",
    )
    assert report.ok, f"code={report.code} detail={report.detail[:300]}"
    assert report.result.message == "本物の本文", "同梱の偽物が使われた"
