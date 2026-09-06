"""添付ファイルの読み取り(toolpack_textio)のテスト。

**読めることより、読めないときに何を返すかを重く見る。**
利用者は「失敗しました」だけでは動けない。何の形式で、どうすれば読めるかを返す。

実在の資料は使わない。すべてこの場で作った架空のファイルで確かめる。
"""

from __future__ import annotations

import zipfile

import pytest

import toolpack_textio as textio


def write(tmp_path, name: str, data: bytes):
    target = tmp_path / name
    target.write_bytes(data)
    return target


# ---------------------------------------------------------------------------
# 文字として読める形式
# ---------------------------------------------------------------------------
def test_UTF8のテキストを読む(tmp_path) -> None:
    path = write(tmp_path, "a.txt", "日本語の本文".encode("utf-8"))
    assert textio.read_text(str(path)) == "日本語の本文"


def test_cp932のテキストを読む(tmp_path) -> None:
    path = write(tmp_path, "a.txt", "日本語の本文".encode("cp932"))
    assert textio.read_text(str(path)) == "日本語の本文"


def test_BOM付きとUTF16を読む(tmp_path) -> None:
    assert textio.read_text(str(write(tmp_path, "a.txt", "あ".encode("utf-8-sig")))) == "あ"
    assert textio.read_text(str(write(tmp_path, "b.txt", "あ".encode("utf-16")))) == "あ"


def test_字幕形式は時刻と連番を落とす(tmp_path) -> None:
    body = (
        "WEBVTT\n\n"
        "1\n00:00:01.000 --> 00:00:03.000\nこんにちは\n\n"
        "2\n00:00:03.000 --> 00:00:05.000\n会議を始めます\n"
    )
    path = write(tmp_path, "a.vtt", body.encode("utf-8"))
    assert textio.read_text(str(path)) == "こんにちは\n会議を始めます"


def test_字幕の連続した重複は畳む(tmp_path) -> None:
    body = (
        "1\n00:00:01,000 --> 00:00:03,000\n同じ発言\n\n"
        "2\n00:00:03,000 --> 00:00:05,000\n同じ発言\n\n"
        "3\n00:00:05,000 --> 00:00:07,000\n次の発言\n"
    )
    path = write(tmp_path, "a.srt", body.encode("utf-8"))
    assert textio.read_text(str(path)) == "同じ発言\n次の発言"


def test_HTMLはタグと台本を落とす(tmp_path) -> None:
    body = "<html><script>x=1</script><style>p{}</style><body><h1>題</h1><p>本文</p></body></html>"
    path = write(tmp_path, "a.html", body.encode("utf-8"))
    text = textio.read_text(str(path))
    assert "題" in text and "本文" in text
    assert "x=1" not in text and "p{}" not in text


def test_Word文書は段落と表を読む(tmp_path) -> None:
    import docx

    document = docx.Document()
    document.add_paragraph("段落の本文")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "左"
    table.rows[0].cells[1].text = "右"
    path = tmp_path / "a.docx"
    document.save(path)

    text = textio.read_text(str(path))
    assert "段落の本文" in text
    assert "左\t右" in text


def test_発表資料はページとノートを読む(tmp_path) -> None:
    from pptx import Presentation

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "発表の題"
    slide.notes_slide.notes_text_frame.text = "話す内容"
    path = tmp_path / "a.pptx"
    presentation.save(path)

    text = textio.read_text(str(path))
    assert "1ページ" in text and "発表の題" in text and "話す内容" in text


def test_Excelはシート名と値を読む(tmp_path) -> None:
    from openpyxl import Workbook

    book = Workbook()
    book.active.title = "実績"
    book.active["A1"] = "科"
    book.active["B1"] = 10
    path = tmp_path / "a.xlsx"
    book.save(path)

    text = textio.read_text(str(path))
    assert "--- 実績 ---" in text and "科\t10" in text


# ---------------------------------------------------------------------------
# 読めないときに何を返すか
# ---------------------------------------------------------------------------
def test_PDFは理由と代わりの出し方を返す(tmp_path) -> None:
    path = write(tmp_path, "a.pdf", b"%PDF-1.4 dummy")
    with pytest.raises(textio.UnreadableFile) as caught:
        textio.read_text(str(path))
    message = str(caught.value)
    assert "PDF" in message
    assert "Word" in message or "テキスト" in message  # 代わりの出し方を示す


def test_古いWordは変換の仕方を返す(tmp_path) -> None:
    path = write(tmp_path, "a.doc", b"\xd0\xcf\x11\xe0dummy")
    with pytest.raises(textio.UnreadableFile) as caught:
        textio.read_text(str(path))
    assert ".docx" in str(caught.value)


def test_拡張子を付け替えたWordを見抜く(tmp_path) -> None:
    """実際に起きる。.docx を .txt へ変えて添付されると、
    「文字コードを判別できません」では利用者が直せない。"""
    import docx

    document = docx.Document()
    document.add_paragraph("本文")
    original = tmp_path / "a.docx"
    document.save(original)
    renamed = write(tmp_path, "a.txt", original.read_bytes())

    with pytest.raises(textio.UnreadableFile) as caught:
        textio.read_text(str(renamed))
    assert "Word" in str(caught.value) and ".docx" in str(caught.value)


def test_中身がPDFのdocxを見抜く(tmp_path) -> None:
    path = write(tmp_path, "a.docx", b"%PDF-1.4 dummy")
    with pytest.raises(textio.UnreadableFile) as caught:
        textio.read_text(str(path))
    assert "PDF" in str(caught.value)


def test_知らない拡張子は対応一覧を示す(tmp_path) -> None:
    path = write(tmp_path, "a.dat", b"\x01\x02\x03")
    with pytest.raises(textio.UnreadableFile) as caught:
        textio.read_text(str(path))
    assert ".docx" in str(caught.value)  # 何なら読めるかを出す


def test_空のファイルはそう言う(tmp_path) -> None:
    path = write(tmp_path, "a.txt", b"")
    with pytest.raises(textio.UnreadableFile) as caught:
        textio.read_text(str(path))
    assert "空" in str(caught.value)


def test_文字が取れなければ推測しない(tmp_path) -> None:
    """画像だけのWord文書などは「読めた」ことにしない。"""
    import docx

    document = docx.Document()
    path = tmp_path / "a.docx"
    document.save(path)
    with pytest.raises(textio.UnreadableFile) as caught:
        textio.read_text(str(path))
    assert "文字を取り出せません" in str(caught.value)


def test_圧縮ファイルは取り出すよう促す(tmp_path) -> None:
    path = tmp_path / "a.zip"
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("中身.txt", "x")
    with pytest.raises(textio.UnreadableFile) as caught:
        textio.read_text(str(path))
    assert "取り出" in str(caught.value)


def test_無い書類はそう言う(tmp_path) -> None:
    with pytest.raises(textio.UnreadableFile) as caught:
        textio.read_text(str(tmp_path / "無い.txt"))
    assert "見つかりません" in str(caught.value)


# ---------------------------------------------------------------------------
# request.json からの読み取り
# ---------------------------------------------------------------------------
def test_requestの入力を読む(tmp_path) -> None:
    (tmp_path / "in").mkdir()
    (tmp_path / "in" / "x.txt").write_text("本文", encoding="utf-8")
    request = {"inputs": [{"id": "in-1", "path": "in/x.txt", "filename": "議事録.txt"}]}
    text, display = textio.read_request_input(request, str(tmp_path))
    assert text == "本文" and display == "議事録.txt"


def test_添付が無ければそう言う(tmp_path) -> None:
    with pytest.raises(textio.UnreadableFile) as caught:
        textio.read_request_input({"inputs": []}, str(tmp_path))
    assert "添付" in str(caught.value)


def test_利用者へ見せる名前で理由を書く(tmp_path) -> None:
    """ジョブ内の内部名(x.txt)ではなく、元のファイル名で知らせる。"""
    (tmp_path / "in").mkdir()
    (tmp_path / "in" / "x.pdf").write_bytes(b"%PDF-1.4")
    request = {"inputs": [{"id": "in-1", "path": "in/x.pdf", "filename": "会議資料.pdf"}]}
    with pytest.raises(textio.UnreadableFile) as caught:
        textio.read_request_input(request, str(tmp_path))
    assert "会議資料.pdf" in str(caught.value)
