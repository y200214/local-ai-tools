"""複数シートへ書き込んだコメントの確認用Excelを別ファイルで作る。

元のExcelにはシート・数式・VBAを追加しない。確認用ブックには、貼付先、
コメント全文、参照表の静的コピー、書込結果を一件ずつ縦に並べる。
"""

from __future__ import annotations

import argparse
import json
import math
import sys
import unicodedata
from copy import copy
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
sys.stderr.reconfigure(encoding="utf-8", errors="replace")

REFERENCE_MIN_COLUMN = 1
REFERENCE_MAX_COLUMN = 11  # A:K
REPORT_REFERENCE_COLUMN = 7  # G
REPORT_MAX_COLUMN = 17  # Q
SEARCH_MAX_ROW = 500


def _normalized(value: object) -> str:
    return unicodedata.normalize("NFKC", str(value or "")).replace(" ", "")


def _row_text(sheet, row: int) -> str:
    return "".join(
        _normalized(sheet.cell(row, column).value)
        for column in range(REFERENCE_MIN_COLUMN, REFERENCE_MAX_COLUMN + 1)
    )


def detect_reference_range(sheet) -> tuple[int, int]:
    """(1)入院から(3)医業収入までを、間に別要素があっても検出する。"""
    scan_end = min(max(sheet.max_row, 1), SEARCH_MAX_ROW)
    first_heading = None
    third_heading = None
    for row in range(1, scan_end + 1):
        text = _row_text(sheet, row)
        if first_heading is None and "(1)" in text and "入院関係" in text:
            first_heading = row
        if first_heading is not None and "(3)" in text and "医業収入" in text:
            third_heading = row
            break

    if first_heading is None:
        nonempty_rows = [
            row
            for row in range(1, scan_end + 1)
            if _row_text(sheet, row)
        ]
        if not nonempty_rows:
            return 1, 1
        return nonempty_rows[0], min(nonempty_rows[-1], nonempty_rows[0] + 79)

    if third_heading is None:
        return first_heading, min(scan_end, first_heading + 79)

    last_nonempty = third_heading
    for row in range(third_heading + 1, scan_end + 2):
        if row > scan_end or not _row_text(sheet, row):
            if last_nonempty > third_heading:
                break
        else:
            last_nonempty = row
    return first_heading, max(third_heading, last_nonempty)


def _static_value(formula_cell, value_cell) -> Any:
    if formula_cell.data_type == "f":
        return value_cell.value if value_cell.value is not None else "（未計算）"
    return formula_cell.value


def _copy_cell_style(source, target) -> None:
    if source.has_style:
        target.font = copy(source.font)
        target.fill = copy(source.fill)
        target.border = copy(source.border)
        target.alignment = copy(source.alignment)
        target.number_format = source.number_format
        target.protection = copy(source.protection)


def _copy_reference_table(
    source_sheet,
    value_sheet,
    target_sheet,
    source_start: int,
    source_end: int,
    target_start: int,
) -> None:
    for source_column in range(REFERENCE_MIN_COLUMN, REFERENCE_MAX_COLUMN + 1):
        target_column = REPORT_REFERENCE_COLUMN + source_column - REFERENCE_MIN_COLUMN
        source_width = source_sheet.column_dimensions[get_column_letter(source_column)].width
        if source_width is not None:
            target_sheet.column_dimensions[get_column_letter(target_column)].width = max(
                target_sheet.column_dimensions[get_column_letter(target_column)].width or 0,
                min(source_width, 32),
            )

    for source_row in range(source_start, source_end + 1):
        target_row = target_start + source_row - source_start
        source_height = source_sheet.row_dimensions[source_row].height
        if source_height is not None:
            target_sheet.row_dimensions[target_row].height = max(
                target_sheet.row_dimensions[target_row].height or 0,
                source_height,
            )
        for source_column in range(REFERENCE_MIN_COLUMN, REFERENCE_MAX_COLUMN + 1):
            source_cell = source_sheet.cell(source_row, source_column)
            if isinstance(source_cell, MergedCell):
                continue
            target_column = REPORT_REFERENCE_COLUMN + source_column - REFERENCE_MIN_COLUMN
            target_cell = target_sheet.cell(target_row, target_column)
            target_cell.value = _static_value(
                source_cell,
                value_sheet.cell(source_row, source_column),
            )
            _copy_cell_style(source_cell, target_cell)

    for merged_range in source_sheet.merged_cells.ranges:
        if (
            merged_range.min_row >= source_start
            and merged_range.max_row <= source_end
            and merged_range.min_col >= REFERENCE_MIN_COLUMN
            and merged_range.max_col <= REFERENCE_MAX_COLUMN
        ):
            target_sheet.merge_cells(
                start_row=target_start + merged_range.min_row - source_start,
                end_row=target_start + merged_range.max_row - source_start,
                start_column=REPORT_REFERENCE_COLUMN + merged_range.min_col - REFERENCE_MIN_COLUMN,
                end_column=REPORT_REFERENCE_COLUMN + merged_range.max_col - REFERENCE_MIN_COLUMN,
            )


def _comment_row_count(comment: str) -> int:
    visual_lines = sum(
        max(1, math.ceil(len(line) / 48))
        for line in comment.splitlines() or [""]
    )
    return max(8, math.ceil(visual_lines * 1.15))


def _merged_anchor(sheet, coordinate: str) -> str:
    for merged_range in sheet.merged_cells.ranges:
        if coordinate in merged_range:
            return sheet.cell(merged_range.min_row, merged_range.min_col).coordinate
    return coordinate


def _validated_records(records: list[dict[str, Any]]) -> list[dict[str, str]]:
    if not isinstance(records, list) or not records:
        raise ValueError("recordsには1件以上の配列を指定してください")
    validated: list[dict[str, str]] = []
    for record in records:
        if not isinstance(record, dict):
            raise ValueError("各recordはオブジェクトで指定してください")
        sheet = record.get("sheet")
        cell = record.get("cell")
        value = record.get("value")
        if not all(isinstance(item, str) and item for item in (sheet, cell, value)):
            raise ValueError("各recordにはsheet、cell、valueの文字列が必要です")
        validated.append(
            {
                "sheet": sheet,
                "cell": cell.upper(),
                "value": value,
                "template": str(record.get("template") or ""),
            }
        )
    return validated


def build_comment_report(
    workbook_path: Path,
    records: list[dict[str, Any]],
    output_path: Path | None = None,
) -> Path:
    workbook_path = Path(workbook_path).resolve()
    validated = _validated_records(records)
    keep_vba = workbook_path.suffix.lower() == ".xlsm"
    source = load_workbook(workbook_path, data_only=False, keep_vba=keep_vba, keep_links=False)
    values = load_workbook(workbook_path, data_only=True, keep_vba=keep_vba, keep_links=False)
    try:
        missing = [record["sheet"] for record in validated if record["sheet"] not in source.sheetnames]
        if missing:
            raise ValueError(f"参照先シートがありません: {', '.join(missing)}")

        report = Workbook()
        sheet = report.active
        sheet.title = "コメント確認一覧"
        sheet.sheet_view.showGridLines = False
        sheet.freeze_panes = "A3"
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.sheet_view.zoomScale = 90

        dark_fill = PatternFill("solid", fgColor="3F3F3F")
        medium_fill = PatternFill("solid", fgColor="D9D9D9")
        light_fill = PatternFill("solid", fgColor="F2F2F2")
        white_font = Font(name="Meiryo UI", size=11, bold=True, color="FFFFFF")
        label_font = Font(name="Meiryo UI", size=9, bold=True, color="333333")
        body_font = Font(name="Meiryo UI", size=9, color="222222")
        thin_gray = Side(style="thin", color="B7B7B7")
        frame = Border(left=thin_gray, right=thin_gray, top=thin_gray, bottom=thin_gray)

        for column in range(1, 7):
            sheet.column_dimensions[get_column_letter(column)].width = 12
        sheet.merge_cells("A1:Q1")
        sheet["A1"] = "コメント確認一覧"
        sheet["A1"].fill = dark_fill
        sheet["A1"].font = Font(name="Meiryo UI", size=14, bold=True, color="FFFFFF")
        sheet["A1"].alignment = Alignment(horizontal="center", vertical="center")
        sheet.row_dimensions[1].height = 28

        written_at = datetime.fromtimestamp(workbook_path.stat().st_mtime).strftime("%Y/%m/%d %H:%M:%S")
        current_row = 3
        for number, record in enumerate(validated, start=1):
            source_sheet = source[record["sheet"]]
            value_sheet = values[record["sheet"]]
            reference_start, reference_end = detect_reference_range(source_sheet)
            reference_rows = reference_end - reference_start + 1
            content_rows = max(reference_rows, _comment_row_count(record["value"]))
            content_start = current_row + 3
            content_end = content_start + content_rows - 1

            sheet.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=17)
            header = sheet.cell(current_row, 1)
            header.value = f"No.{number}  {record['sheet']}!{record['cell']}"
            header.fill = dark_fill
            header.font = white_font
            header.alignment = Alignment(vertical="center")
            sheet.row_dimensions[current_row].height = 23

            metadata = [
                ("A", "B", "貼付先"),
                ("G", "H", "参照範囲"),
                ("L", "M", "状態"),
            ]
            for start_column, end_column, label in metadata:
                sheet.merge_cells(f"{start_column}{current_row + 1}:{end_column}{current_row + 1}")
                cell = sheet[f"{start_column}{current_row + 1}"]
                cell.value = label
                cell.fill = medium_fill
                cell.font = label_font
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.border = frame

            sheet.merge_cells(f"C{current_row + 1}:F{current_row + 1}")
            destination = sheet[f"C{current_row + 1}"]
            anchor = _merged_anchor(source_sheet, record["cell"])
            destination.value = f"{record['sheet']}!{anchor}"
            destination.hyperlink = f"{workbook_path.as_uri()}#'{record['sheet'].replace(chr(39), chr(39) * 2)}'!{anchor}"
            destination.style = "Hyperlink"
            destination.border = frame

            sheet.merge_cells(f"I{current_row + 1}:K{current_row + 1}")
            reference_cell = sheet[f"I{current_row + 1}"]
            reference_cell.value = f"A{reference_start}:K{reference_end}"
            reference_cell.font = body_font
            reference_cell.alignment = Alignment(horizontal="center", vertical="center")
            reference_cell.border = frame

            sheet.merge_cells(f"N{current_row + 1}:Q{current_row + 1}")
            status = sheet[f"N{current_row + 1}"]
            status.value = f"書込済み  {written_at}"
            status.font = body_font
            status.alignment = Alignment(horizontal="center", vertical="center")
            status.border = frame

            sheet.merge_cells(f"A{current_row + 2}:B{current_row + 2}")
            template_label = sheet[f"A{current_row + 2}"]
            template_label.value = "定型文"
            template_label.fill = medium_fill
            template_label.font = label_font
            template_label.alignment = Alignment(horizontal="center", vertical="center")
            template_label.border = frame
            sheet.merge_cells(f"C{current_row + 2}:F{current_row + 2}")
            template_cell = sheet[f"C{current_row + 2}"]
            template_cell.value = record["template"]
            template_cell.font = body_font
            template_cell.alignment = Alignment(vertical="center", wrap_text=True)
            template_cell.border = frame
            sheet.merge_cells(f"G{current_row + 2}:Q{current_row + 2}")
            table_label = sheet[f"G{current_row + 2}"]
            table_label.value = "参照表（値・書式の静的コピー）"
            table_label.fill = light_fill
            table_label.font = label_font
            table_label.alignment = Alignment(horizontal="center", vertical="center")
            table_label.border = frame

            sheet.merge_cells(
                start_row=content_start,
                start_column=1,
                end_row=content_end,
                end_column=6,
            )
            comment_cell = sheet.cell(content_start, 1)
            comment_cell.value = record["value"]
            comment_cell.font = body_font
            comment_cell.fill = light_fill
            comment_cell.alignment = Alignment(vertical="top", wrap_text=True)
            comment_cell.border = frame
            for row in range(content_start, content_end + 1):
                if sheet.row_dimensions[row].height is None:
                    sheet.row_dimensions[row].height = 18

            _copy_reference_table(
                source_sheet,
                value_sheet,
                sheet,
                reference_start,
                reference_end,
                content_start,
            )
            current_row = content_end + 2

        output_path = output_path or workbook_path.with_name(
            f"{workbook_path.stem}_コメント確認一覧.xlsx"
        )
        output_path = Path(output_path).resolve()
        output_path.parent.mkdir(parents=True, exist_ok=True)
        report.save(output_path)
        report.close()
    finally:
        source.close()
        values.close()
    return output_path


def _load_records(records_json: str | None, records_file: str | None) -> list[dict[str, Any]]:
    try:
        raw = Path(records_file).read_text(encoding="utf-8") if records_file else records_json
        return json.loads(raw or "")
    except (OSError, json.JSONDecodeError) as error:
        raise SystemExit(f"コメント記録JSONを読み込めません: {error}") from error


def main() -> None:
    parser = argparse.ArgumentParser(description="Excelコメントの確認一覧を別ブックで作る")
    parser.add_argument("workbook", help="コメント書込後のExcel")
    records = parser.add_mutually_exclusive_group(required=True)
    records.add_argument("--records-json", help="コメント記録のJSON配列")
    records.add_argument("--records-file", help="コメント記録のUTF-8 JSONファイル")
    parser.add_argument("--output", help="確認一覧の出力先")
    args = parser.parse_args()

    workbook_path = Path(args.workbook)
    if not workbook_path.exists():
        raise SystemExit(f"ファイルが見つかりません: {workbook_path}")
    output = build_comment_report(
        workbook_path,
        _load_records(args.records_json, args.records_file),
        Path(args.output) if args.output else None,
    )
    print(f"確認一覧を保存しました: {output}")


if __name__ == "__main__":
    main()
