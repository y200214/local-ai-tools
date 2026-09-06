"""Excelからコメント生成に必要な実データ・記入例・書込先を抽出する。"""

from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path

from openpyxl import load_workbook

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
sys.stderr.reconfigure(encoding="utf-8", errors="replace")


def _merged_anchor(sheet, row: int, column: int):
    coordinate = sheet.cell(row, column).coordinate
    for merged_range in sheet.merged_cells.ranges:
        if coordinate in merged_range:
            return sheet.cell(merged_range.min_row, merged_range.min_col)
    return sheet.cell(row, column)


def _analysis_value(raw_value, cached_value):
    """分析用には、式そのものよりExcelに保存済みの計算結果を優先する。"""
    if cached_value is not None and (
        isinstance(raw_value, str) and raw_value.lstrip().startswith("=")
    ):
        return cached_value
    return raw_value if raw_value is not None else cached_value


def _display_metric(label: str, value: float, number_format: str = "") -> tuple[float, str, str]:
    if "‰" in number_format or "逆紹介率" in label:
        permille = value * 100 if abs(value) <= 1 else value
        return permille, f"{permille:.1f}‰", "‰"
    if ("%" in number_format or "稼働率" in label) and abs(value) <= 1:
        percent = value * 100
        return percent, f"{percent:.1f}%", "%"
    unquoted_format = re.sub(r'"[^"]*"', "", number_format.split(";", 1)[0])
    scale_commas = len(re.search(r",+$", unquoted_format).group(0)) if re.search(r",+$", unquoted_format) else 0
    display_value = value / (1000 ** scale_commas) if scale_commas else value
    display_unit = "百万円" if scale_commas == 2 else ""
    if float(display_value).is_integer():
        return display_value, f"{int(display_value):,}{display_unit}", display_unit
    return display_value, f"{display_value:,.1f}{display_unit}", display_unit


def _comparison_fact_records(sheet, value_sheet) -> list[dict]:
    """左表の直近2年度を比較し、LLMが例文を事実として写さないための根拠を作る。"""
    sections: list[tuple[str, int]] = []
    section_boundaries: list[int] = []
    for row in range(1, min(sheet.max_row or 1, 200) + 1):
        for column in range(1, min(sheet.max_column or 1, 12) + 1):
            value = sheet.cell(row, column).value
            if not isinstance(value, str):
                continue
            if re.match(r"^\s*[（(]\d+[）)]", value):
                section_boundaries.append(row)
            if "入院関係" in value:
                sections.append(("入院", row))
            elif "外来関係" in value:
                sections.append(("外来", row))
            elif "入外合算" in value or "入外合計" in value:
                sections.append(("入外合計", row))

    facts: list[dict] = []
    latest_values: list[dict] = []
    section_rows = sorted(sections, key=lambda item: item[1])
    for section_name, start_row in section_rows:
        next_boundary = next(
            (row for row in sorted(set(section_boundaries)) if row > start_row),
            None,
        )
        end_row = next_boundary - 1 if next_boundary else min(sheet.max_row or start_row, start_row + 25)
        header_row = None
        year_columns: list[tuple[int, str]] = []
        for row in range(start_row, min(end_row, start_row + 5) + 1):
            candidates = []
            for column in range(1, min(sheet.max_column or 1, 32) + 1):
                value = sheet.cell(row, column).value
                if isinstance(value, str) and re.fullmatch(r"[HR]\s*\d+年度", value.strip()):
                    candidates.append((column, value.strip().replace(" ", "")))
            if len(candidates) >= 2:
                header_row = row
                year_columns = candidates
                break
        if header_row is None:
            continue

        previous_column, previous_year = year_columns[-2]
        latest_column, latest_year = year_columns[-1]
        for row in range(header_row + 1, end_row + 1):
            labels = [
                str(sheet.cell(row, column).value).strip()
                for column in range(1, min(latest_column, 7) + 1)
                if sheet.cell(row, column).value not in (None, "")
            ]
            label = " ".join(labels)
            previous = value_sheet.cell(row, previous_column).value
            latest = value_sheet.cell(row, latest_column).value
            if not label or not isinstance(latest, (int, float)):
                continue

            latest_numeric, latest_text, latest_unit = _display_metric(
                label,
                float(latest),
                value_sheet.cell(row, latest_column).number_format,
            )
            latest_values.append({
                "section": section_name,
                "metric": label,
                "latest_year": latest_year,
                "latest_value": latest_numeric,
                "latest_text": latest_text,
                "value_unit": latest_unit,
            })
            if not isinstance(previous, (int, float)):
                continue

            previous_numeric, previous_text, previous_unit = _display_metric(
                label,
                float(previous),
                value_sheet.cell(row, previous_column).number_format,
            )
            absolute_change = latest_numeric - previous_numeric
            if latest_unit == "‰":
                change = latest_numeric - previous_numeric
                change_text = f"{change:+.1f}‰"
                change_unit = "‰"
            elif latest_unit == "%" or previous_unit == "%":
                change = latest_numeric - previous_numeric
                change_text = f"{change:+.1f}ポイント"
                change_unit = "ポイント"
            elif previous_numeric:
                change = (latest_numeric - previous_numeric) / previous_numeric * 100
                change_text = f"{change:+.1f}%"
                change_unit = "%"
            else:
                change = None
                change_text = "比較不能（前年度0）"
                change_unit = "%"
            facts.append({
                "kind": "comparison",
                "section": section_name,
                "metric": label,
                "previous_year": previous_year,
                "previous_value": previous_numeric,
                "previous_text": previous_text,
                "latest_year": latest_year,
                "latest_value": latest_numeric,
                "latest_text": latest_text,
                "value_unit": latest_unit or previous_unit,
                "absolute_change": absolute_change,
                "change": change,
                "change_text": change_text,
                "change_unit": change_unit,
            })

            # ひな形の手術件数などは、直近年度だけでなくその前年との推移も
            # 評価するため、R6→R8相当の比較を別レコードとして保持する。
            if len(year_columns) >= 3:
                trend_column, trend_year = year_columns[-3]
                trend_value = value_sheet.cell(row, trend_column).value
                if isinstance(trend_value, (int, float)):
                    trend_numeric, trend_text, trend_unit = _display_metric(
                        label,
                        float(trend_value),
                        value_sheet.cell(row, trend_column).number_format,
                    )
                    trend_absolute = latest_numeric - trend_numeric
                    if latest_unit == "‰":
                        trend_change = trend_absolute
                        trend_change_text = f"{trend_change:+.1f}‰"
                        trend_change_unit = "‰"
                    elif latest_unit == "%" or trend_unit == "%":
                        trend_change = trend_absolute
                        trend_change_text = f"{trend_change:+.1f}ポイント"
                        trend_change_unit = "ポイント"
                    elif trend_numeric:
                        trend_change = trend_absolute / trend_numeric * 100
                        trend_change_text = f"{trend_change:+.1f}%"
                        trend_change_unit = "%"
                    else:
                        trend_change = None
                        trend_change_text = f"比較不能（{trend_year}0）"
                        trend_change_unit = "%"
                    facts.append({
                        "kind": "trend_comparison",
                        "section": section_name,
                        "metric": label,
                        "previous_year": trend_year,
                        "previous_value": trend_numeric,
                        "previous_text": trend_text,
                        "latest_year": latest_year,
                        "latest_value": latest_numeric,
                        "latest_text": latest_text,
                        "value_unit": latest_unit or trend_unit,
                        "absolute_change": trend_absolute,
                        "change": trend_change,
                        "change_text": trend_change_text,
                        "change_unit": trend_change_unit,
                    })

            h30 = next(((column, year) for column, year in year_columns if year == "H30年度"), None)
            if h30:
                h30_column, h30_year = h30
                h30_value = value_sheet.cell(row, h30_column).value
                if isinstance(h30_value, (int, float)):
                    h30_numeric, h30_text, h30_unit = _display_metric(
                        label,
                        float(h30_value),
                        value_sheet.cell(row, h30_column).number_format,
                    )
                    if latest_unit == "‰":
                        long_change_value = latest_numeric - h30_numeric
                        long_change_unit = "‰"
                        long_change = f"{long_change_value:+.1f}‰"
                        long_rate_change = None
                    elif latest_unit == "%" or h30_unit == "%":
                        long_change_value = latest_numeric - h30_numeric
                        long_change_unit = "ポイント"
                        long_change = f"{long_change_value:+.1f}ポイント"
                        long_rate_change = None
                    elif h30_numeric:
                        absolute = latest_numeric - h30_numeric
                        rate = absolute / h30_numeric * 100
                        long_change_value = absolute
                        long_change_unit = latest_unit or h30_unit
                        long_rate_change = rate
                        long_change = f"{absolute:+,.1f}{latest_unit}（{rate:+.1f}%）"
                    else:
                        long_change_value = None
                        long_change_unit = latest_unit or h30_unit
                        long_rate_change = None
                        long_change = "比較不能（H30年度0）"
                    facts.append({
                        "kind": "long_comparison",
                        "section": section_name,
                        "metric": label,
                        "previous_year": h30_year,
                        "previous_value": h30_numeric,
                        "previous_text": h30_text,
                        "latest_year": latest_year,
                        "latest_value": latest_numeric,
                        "latest_text": latest_text,
                        "value_unit": latest_unit or h30_unit,
                        "absolute_change": latest_numeric - h30_numeric,
                        "change": long_change_value,
                        "change_unit": long_change_unit,
                        "rate_change": long_rate_change,
                        "change_text": long_change,
                    })

    target_text = "\n".join(
        str(sheet.cell(row, column).value)
        for row in range(1, min(sheet.max_row or 1, 20) + 1)
        for column in range(1, min(sheet.max_column or 1, 32) + 1)
        if isinstance(sheet.cell(row, column).value, str)
    )
    target_match = re.search(r"稼働率\s*(\d+(?:\.\d+)?)\s*[％%]以上", target_text)
    occupancy = next(
        (
            fact
            for fact in reversed(facts)
            if fact["kind"] == "comparison"
            and fact["section"] == "入院"
            and "稼働率" in fact["metric"]
        ),
        None,
    )
    if occupancy is None:
        occupancy = next(
            (
                value
                for value in reversed(latest_values)
                if value["section"] == "入院" and "稼働率" in value["metric"]
            ),
            None,
        )
    if target_match and occupancy:
        target = float(target_match.group(1))
        current = float(occupancy["latest_value"])
        status = "達成" if current >= target else "未達"
        facts.append({
            "kind": "target",
            "section": "入院",
            "metric": "稼働率",
            "target": target,
            "target_unit": "%",
            "latest_value": current,
            "value_unit": "%",
            "status": status,
            "gap": current - target,
            "gap_unit": "ポイント",
        })
    reverse_referral = next(
        (
            fact
            for fact in reversed(facts)
            if fact["kind"] == "comparison"
            and fact["section"] == "外来"
            and "逆紹介率" in fact["metric"]
            and fact["value_unit"] == "‰"
        ),
        None,
    )
    if reverse_referral is None:
        reverse_referral = next(
            (
                value
                for value in reversed(latest_values)
                if value["section"] == "外来"
                and "逆紹介率" in value["metric"]
                and value["value_unit"] == "‰"
            ),
            None,
        )
    if reverse_referral:
        # 様式内に「目標逆紹介率」があれば、その行の最新年度値を正とする。
        # 旧様式など目標行がない場合のみ、制度基準50‰へフォールバックする。
        reverse_target_value = None
        for row in range(1, (sheet.max_row or 1) + 1):
            label_cells = [
                sheet.cell(row, column).value
                for column in range(1, (sheet.max_column or 1) + 1)
            ]
            if not any(isinstance(value, str) and "目標逆紹介率" in value for value in label_cells):
                continue
            numeric_values = [
                value_sheet.cell(row, column).value
                for column in range(1, (sheet.max_column or 1) + 1)
                if isinstance(value_sheet.cell(row, column).value, (int, float))
            ]
            if numeric_values:
                reverse_target_value = float(numeric_values[-1])
                break
        target = reverse_target_value if reverse_target_value is not None else 50.0
        current = float(reverse_referral["latest_value"])
        status = "達成" if current >= target else "未達"
        facts.append({
            "kind": "target",
            "section": "外来",
            "metric": "逆紹介率",
            "target": target,
            "target_unit": "‰",
            "latest_value": current,
            "value_unit": "‰",
            "status": status,
            "gap": current - target,
            "gap_unit": "‰",
            "target_source": "workbook" if reverse_target_value is not None else "fallback",
        })
    return facts


def _format_fact(fact: dict) -> str:
    prefix = f'{fact["section"]}・{fact["metric"]}'
    if fact["kind"] in {"comparison", "trend_comparison", "long_comparison"}:
        long_suffix = (
            "・長期比較" if fact["kind"] == "long_comparison"
            else "・推移比較" if fact["kind"] == "trend_comparison"
            else ""
        )
        count_unit = "人" if "患者数" in fact["metric"] else "件" if "件数" in fact["metric"] else ""
        absolute_change = ""
        if fact["kind"] == "comparison" and count_unit:
            change = fact["absolute_change"]
            formatted = f"{change:+,.0f}" if float(change).is_integer() else f"{change:+,.1f}"
            absolute_change = f", 増減数={formatted}{count_unit}"
        return (
            f'{prefix}{long_suffix}: {fact["previous_year"]}={fact["previous_text"]}, '
            f'{fact["latest_year"]}={fact["latest_text"]}, 変化={fact["change_text"]}'
            f'{absolute_change}'
        )
    target_label = "目標" if fact["metric"] == "稼働率" else "基準"
    threshold = "以上" if fact["metric"] == "稼働率" else ""
    text = (
        f'{prefix}{target_label}: {target_label}={fact["target"]:g}'
        f'{fact["target_unit"]}{threshold}, 最新={fact["latest_value"]:g}'
        f'{fact["value_unit"]}, 判定={fact["status"]}'
    )
    if fact["metric"] in {"稼働率", "逆紹介率"}:
        text += f', 差={fact["gap"]:+.1f}{fact["gap_unit"]}'
    return text


def _comparison_facts(sheet, value_sheet) -> list[str]:
    """LLMへ渡す従来形式の事実文。構造化レコードと同じ計算結果から作る。"""
    return [_format_fact(fact) for fact in _comparison_fact_records(sheet, value_sheet)]


def _comment_labels(sheet) -> list:
    return [
        cell
        for row in sheet.iter_rows()
        for cell in row
        if isinstance(cell.value, str)
        and cell.value.strip().startswith(("コメント", "分析内容"))
    ]


def _build_comment_context(path: Path, sheet, value_sheet) -> dict:
    labels = _comment_labels(sheet)
    if not labels:
        raise RuntimeError(f"コメント見出しが見つかりません: {sheet.title}")

    upper_label = min(labels, key=lambda cell: cell.row)
    target = _merged_anchor(sheet, upper_label.row + 1, upper_label.column)

    template = ""
    for label in sorted(labels, key=lambda cell: cell.row, reverse=True):
        candidate = _merged_anchor(sheet, label.row + 1, label.column).value
        if isinstance(candidate, str) and candidate.strip():
            template = candidate.strip()
            break
    if not template:
        raise RuntimeError(f"コメント記入例が見つかりません: {sheet.title}")

    data = []
    max_row = min(sheet.max_row or 1, 200)
    max_column = min(sheet.max_column or 1, 32)  # AF列までを左側の分析表とする。
    for row in sheet.iter_rows(
        min_row=1,
        max_row=max_row,
        min_col=1,
        max_col=max_column,
    ):
        for cell in row:
            value = _analysis_value(cell.value, value_sheet[cell.coordinate].value)
            if value is not None:
                data.append({"cell": cell.coordinate, "value": str(value)})

    fact_records = _comparison_fact_records(sheet, value_sheet)
    return {
        "source": str(path),
        "sheet": sheet.title,
        "target": target.coordinate,
        "template": template,
        "facts": [_format_fact(fact) for fact in fact_records],
        "fact_records": fact_records,
        "data": data,
    }


def collect_comment_contexts(
    path: Path, sheet_name: str | None = None
) -> tuple[list[dict], list[dict]]:
    """
    コメント対象シートを抽出し、成功分と抽出できなかったシートを分けて返す。

    体裁の崩れた1シートのために残り全部の処理が消えるのが最も困るため、
    全シート走査では失敗を握って理由とともに残す。
    シートを名指しされた場合だけは、握らずそのまま失敗を伝える。
    """
    workbook = load_workbook(
        path,
        data_only=False,
        keep_vba=path.suffix.lower() == ".xlsm",
        keep_links=False,
    )
    value_workbook = load_workbook(
        path,
        data_only=True,
        keep_vba=path.suffix.lower() == ".xlsm",
        keep_links=False,
    )
    try:
        if sheet_name is not None:
            if sheet_name not in workbook.sheetnames:
                raise RuntimeError(f"シートがありません: {sheet_name}")
            sheet = workbook[sheet_name]
            return [_build_comment_context(path, sheet, value_workbook[sheet.title])], []

        sheets = [
            sheet
            for sheet in workbook.worksheets
            if sheet.sheet_state == "visible" and _comment_labels(sheet)
        ]
        if not sheets:
            raise RuntimeError("表示中のコメント対象シートが見つかりません")

        contexts: list[dict] = []
        skipped: list[dict] = []
        for sheet in sheets:
            try:
                contexts.append(
                    _build_comment_context(path, sheet, value_workbook[sheet.title])
                )
            except Exception as error:
                skipped.append({"sheet": sheet.title, "reason": str(error)})
        if not contexts:
            details = "、".join(
                f"{item['sheet']}: {item['reason']}" for item in skipped
            )
            raise RuntimeError(f"コメント対象シートを1枚も抽出できませんでした({details})")
        return contexts, skipped
    finally:
        workbook.close()
        value_workbook.close()


def prepare_comment_contexts(path: Path, sheet_name: str | None = None) -> list[dict]:
    """互換API。抽出できたシートだけを返す。"""
    contexts, _ = collect_comment_contexts(path, sheet_name)
    return contexts


def prepare_comment_context(path: Path, sheet_name: str | None = None) -> dict:
    """単一シート用の互換API。省略時は最初の表示中コメント対象シートを返す。"""
    return prepare_comment_contexts(path, sheet_name)[0]


def main() -> None:
    parser = argparse.ArgumentParser(description="Excelコメント生成用コンテキストを抽出")
    parser.add_argument("path")
    target_group = parser.add_mutually_exclusive_group()
    target_group.add_argument("--sheet")
    target_group.add_argument("--all", action="store_true", help="表示中の全コメント対象シートを抽出")
    args = parser.parse_args()
    if args.all:
        contexts, skipped = collect_comment_contexts(Path(args.path))
        result = {"contexts": contexts, "skipped": skipped}
    else:
        result = prepare_comment_context(Path(args.path), args.sheet)
    print(json.dumps(result, ensure_ascii=False))


if __name__ == "__main__":
    main()
