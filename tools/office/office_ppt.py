"""
PowerPointファイル(.pptx)を読み書きするツールの雛形。

使い方(リポジトリ直下で。対象ファイルは input/ に置いてもらう):
  文言の表示:   ...python.exe tools\\office_ppt.py show "input\\発表.pptx"
  文字列の置換: ...python.exe tools\\office_ppt.py replace "input\\発表.pptx" --old "旧文言" --new "新文言"
  新規作成:     ...python.exe tools\\office_ppt.py new "output\\発表.pptx" --from "input\\構成.txt"

(...python.exe = text-processing-bridge\\.venv\\Scripts\\python.exe)

新規作成の原稿は普通のテキストでよい。行頭の記号だけ解釈する:
  「# 」=新しいスライド(タイトル) / それ以外の行=そのスライドの箇条書き

守っている約束:
- input/ のファイルは絶対に上書きしない。置換結果は output/ へ
  「元名_updated.pptx」として保存し、新規作成は保存先が既にあると中止する
- 置換はラン(書式のかたまり)単位で行い、書式を保つ。語が泣き別れて
  置換できないときは件数が0になるので、showで文言を確認して選び直す
"""

from __future__ import annotations

import argparse
import sys
from pathlib import Path

from pptx import Presentation

# PowerShell経由(既定cp932)だと日本語出力が化けるため、常にUTF-8で出す
sys.stdout.reconfigure(encoding="utf-8", errors="replace")
sys.stderr.reconfigure(encoding="utf-8", errors="replace")


def _iter_text_frames(presentation):
    """全スライドのテキスト枠(表の中を含む)を辿る。"""
    for number, slide in enumerate(presentation.slides, 1):
        for shape in slide.shapes:
            if shape.has_text_frame:
                yield number, shape.text_frame
            if shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        yield number, cell.text_frame


def cmd_show(path: Path) -> None:
    presentation = Presentation(path)
    current = 0
    for number, frame in _iter_text_frames(presentation):
        if number != current:
            print(f"\n=== スライド {number} ===")
            current = number
        for paragraph in frame.paragraphs:
            text = "".join(run.text for run in paragraph.runs).strip()
            if text:
                print(f"  {text}")


def cmd_replace(path: Path, old: str, new: str) -> None:
    presentation = Presentation(path)
    count = 0
    for _, frame in _iter_text_frames(presentation):
        for paragraph in frame.paragraphs:
            for run in paragraph.runs:
                if old in run.text:
                    run.text = run.text.replace(old, new)
                    count += 1

    if count == 0:
        print(f"「{old}」は見つかりませんでした(書式の切れ目で泣き別れている可能性)")
        return

    output_dir = Path("output")
    output_dir.mkdir(parents=True, exist_ok=True)
    target = output_dir / f"{path.stem}_updated{path.suffix}"
    presentation.save(target)
    print(f"{count}箇所を置換しました")
    print(f"保存しました(元ファイルは無変更): {target}")


def cmd_new(target: Path, source_text: str) -> None:
    if target.exists():
        raise SystemExit(f"保存先が既にあります(上書きしない約束): {target}")

    presentation = Presentation()
    # レイアウト1 = 「タイトルとコンテンツ」。箇条書きに使う。
    layout = presentation.slide_layouts[1]
    slide = None

    for line in source_text.replace("\r\n", "\n").split("\n"):
        line = line.strip()
        if not line:
            continue
        if line.startswith("# "):
            slide = presentation.slides.add_slide(layout)
            slide.shapes.title.text = line[2:].strip()
            # 本文プレースホルダを空にしておく(既定の文言を残さない)。
            slide.placeholders[1].text_frame.clear()
            continue
        if slide is None:
            # タイトル行より前に本文が来た場合も落とさず、無題スライドに載せる。
            slide = presentation.slides.add_slide(layout)
            slide.shapes.title.text = ""
            slide.placeholders[1].text_frame.clear()
        frame = slide.placeholders[1].text_frame
        text = line.lstrip("-・ 　")
        if frame.paragraphs[0].text == "" and len(frame.paragraphs) == 1:
            frame.paragraphs[0].text = text
        else:
            frame.add_paragraph().text = text

    target.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(target)
    print(f"作成しました: {target}(スライド{len(presentation.slides)}枚)")


def main() -> None:
    parser = argparse.ArgumentParser(
        description="PowerPointの表示・置換・新規作成(既存ファイルは上書きしない)"
    )
    sub = parser.add_subparsers(dest="command", required=True)

    p_show = sub.add_parser("show", help="スライドごとの文言を表示する")
    p_show.add_argument("path")

    p_replace = sub.add_parser("replace", help="文字列を置換して別名保存する")
    p_replace.add_argument("path")
    p_replace.add_argument("--old", required=True, help="置換前の文字列")
    p_replace.add_argument("--new", required=True, help="置換後の文字列")

    p_new = sub.add_parser("new", help="テキスト原稿からpptxを新規作成する")
    p_new.add_argument("path", help="保存先(例 Exsel\\発表.pptx)")
    p_new.add_argument(
        "--from", dest="source", required=True,
        help="原稿テキストファイル(UTF-8)",
    )

    args = parser.parse_args()
    path = Path(args.path)

    if args.command == "new":
        source = Path(args.source)
        if not source.exists():
            raise SystemExit(f"原稿ファイルが見つかりません: {source}")
        cmd_new(path, source.read_text(encoding="utf-8-sig"))
        return

    if not path.exists():
        raise SystemExit(f"ファイルが見つかりません: {path}")

    if args.command == "show":
        cmd_show(path)
    elif args.command == "replace":
        cmd_replace(path, args.old, args.new)


if __name__ == "__main__":
    main()
