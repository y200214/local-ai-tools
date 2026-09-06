"""
Excelファイルの全シートを一覧表示し、CSVへ書き出す読み取り専用ツール。

使い方(リポジトリ直下で。対象ファイルは input/ に置いてもらう):
  text-processing-bridge\\.venv\\Scripts\\python.exe tools\\excel_reader.py "input\\R8第1四半期 科別分析.xlsm"

- 元ファイルには一切書き込まない
- CSVは output/ へ保存する。Excelで開いても文字化けしないよう
  BOM付きUTF-8で書く
- シート名にファイル名に使えない文字が入っていても安全な名前へ置き換える
"""

from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils.cell import get_column_letter, range_boundaries

# PowerShell経由(既定cp932)だと日本語出力が化けるため、常にUTF-8で出す
sys.stdout.reconfigure(encoding="utf-8", errors="replace")
sys.stderr.reconfigure(encoding="utf-8", errors="replace")


def _safe_stem(name: str) -> str:
    """Windowsのファイル名に使えない文字を置き換える。"""
    return re.sub(r'[<>:"/\\|?*\x00-\x1F]', "_", name).strip(" .") or "sheet"


def _visible_sheet_names(xls: pd.ExcelFile) -> list[str]:
    """表示中のシート名だけを返す。"""
    if xls.engine == "openpyxl":
        return [
            sheet.title
            for sheet in xls.book.worksheets
            if sheet.sheet_state == "visible"
        ]
    if xls.engine == "xlrd":
        return [
            sheet_name
            for index, sheet_name in enumerate(xls.sheet_names)
            if xls.book.sheet_by_index(index).visibility == 0
        ]
    return list(xls.sheet_names)


def _open_workbook(path: Path, *, data_only: bool = False):
    """Open modern Excel files read-only without loading all cell data into memory."""
    return load_workbook(path, read_only=True, data_only=data_only, keep_links=False)


def _selected_worksheets(workbook, include_hidden: bool):
    return [
        sheet
        for sheet in workbook.worksheets
        if include_hidden or sheet.sheet_state == "visible"
    ]


def _ensure_dimensions(sheet) -> tuple[int, int]:
    """使用範囲メタデータが無いシートは実データから範囲を再計算する。"""
    if sheet.max_row is None or sheet.max_column is None:
        sheet.reset_dimensions()
        sheet.calculate_dimension(force=True)
    return sheet.max_row or 1, sheet.max_column or 1


def print_index(path: Path, *, include_hidden: bool = False) -> None:
    """Print workbook structure without printing cell contents."""
    workbook = _open_workbook(path)
    print(f"ファイル: {path}")
    print("シート索引:")
    for index, sheet in enumerate(_selected_worksheets(workbook, include_hidden), 1):
        max_row, max_column = _ensure_dimensions(sheet)
        last_cell = f"{get_column_letter(max_column)}{max_row}"
        print(
            f"  {index}. {sheet.title}: A1:{last_cell} "
            f"({max_row} 行 × {max_column} 列, {sheet.sheet_state})"
        )
    workbook.close()


def print_range(
    path: Path,
    sheet_name: str,
    cell_range: str,
    *,
    data_only: bool = False,
    max_cells: int = 500,
    compact: bool = False,
) -> None:
    """Print a bounded cell range with coordinates for source verification."""
    min_col, min_row, max_col, max_row = range_boundaries(cell_range)
    cell_count = (max_col - min_col + 1) * (max_row - min_row + 1)
    if cell_count > max_cells:
        raise SystemExit(
            f"指定範囲は {cell_count} セルです。上限 {max_cells} セル以下に分割してください。"
        )
    workbook = _open_workbook(path, data_only=data_only)
    if sheet_name not in workbook.sheetnames:
        workbook.close()
        raise SystemExit(f"シートがありません: {sheet_name}")
    sheet = workbook[sheet_name]
    print(f"範囲: {sheet_name}!{cell_range} ({cell_count} セル)")
    for row_index, row in enumerate(sheet.iter_rows(
        min_row=min_row,
        max_row=max_row,
        min_col=min_col,
        max_col=max_col,
    ), start=min_row):
        if compact:
            for column_index, cell in enumerate(row, start=min_col):
                if cell.value is not None:
                    print(f"{get_column_letter(column_index)}{row_index}={cell.value}")
            continue
        print(
            "\t".join(
                f"{get_column_letter(column_index)}{row_index}="
                f"{'' if cell.value is None else cell.value}"
                for column_index, cell in enumerate(row, start=min_col)
            )
        )
    workbook.close()


def find_values(
    path: Path,
    query: str,
    *,
    include_hidden: bool = False,
    max_results: int = 50,
    context_rows: int = 0,
) -> None:
    """Search values but return only a bounded list of matching cells."""
    if not query:
        raise SystemExit("--find には空でない検索語を指定してください。")
    workbook = _open_workbook(path, data_only=True)
    needle = query.casefold()
    found = 0
    print(f"検索: {query!r} (最大 {max_results} 件)")
    for sheet in _selected_worksheets(workbook, include_hidden):
        _ensure_dimensions(sheet)
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value is None or needle not in str(cell.value).casefold():
                    continue
                print(f"{sheet.title}!{cell.coordinate}\t{cell.value}")
                for row_offset in range(1, context_rows + 1):
                    context_row = cell.row + row_offset
                    if context_row > (sheet.max_row or context_row):
                        break
                    context_cell = sheet.cell(context_row, cell.column)
                    if context_cell.value is not None:
                        coordinate = f"{get_column_letter(cell.column)}{context_row}"
                        print(f"  直下{row_offset}行: {coordinate}\t{context_cell.value}")
                found += 1
                if found >= max_results:
                    print(f"結果上限 {max_results} 件に達したため検索を停止しました。")
                    workbook.close()
                    return
    workbook.close()
    print(f"一致: {found} 件")


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Excelの全シートを一覧し、CSVへ書き出す(元ファイルは変更しない)"
    )
    parser.add_argument("path", help="対象のExcelファイル(.xlsx / .xlsm / .xls)")
    parser.add_argument(
        "--rows", type=int, default=10, help="画面に表示する先頭行数(既定10)"
    )
    parser.add_argument(
        "--no-csv", action="store_true", help="CSVを書き出さず表示だけ行う"
    )
    parser.add_argument(
        "--include-hidden",
        action="store_true",
        help="非表示・VeryHiddenシートも読み込む",
    )
    parser.add_argument(
        "--index-only",
        action="store_true",
        help="セル本文を読まず、シート名と使用範囲だけを表示する",
    )
    parser.add_argument("--sheet", help="範囲表示するシート名")
    parser.add_argument("--range", dest="cell_range", help="表示範囲 (例: A1:F20)")
    parser.add_argument("--find", help="全セルから値を検索し、一致セルだけを表示する")
    parser.add_argument("--max-results", type=int, default=50, help="検索結果上限")
    parser.add_argument("--max-cells", type=int, default=500, help="範囲表示セル数上限")
    parser.add_argument(
        "--context-rows",
        type=int,
        default=0,
        help="検索一致セルと同じ列の直下を追加表示する行数",
    )
    parser.add_argument(
        "--compact",
        action="store_true",
        help="範囲表示で空セルを省き、非空セルだけを座標付きで表示する",
    )
    parser.add_argument(
        "--data-only", action="store_true", help="数式でなく保存済み計算結果を表示する"
    )
    parser.add_argument(
        "--csv-sheet",
        help="CSV出力を1シートだけに絞る(未指定なら表示中の全シート)",
    )
    args = parser.parse_args()

    excel_path = Path(args.path)
    if not excel_path.exists():
        raise SystemExit(f"ファイルが見つかりません: {excel_path}")

    lightweight_mode = args.index_only or args.find is not None or args.sheet or args.cell_range
    if lightweight_mode and excel_path.suffix.lower() == ".xls":
        raise SystemExit("軽量な索引・検索・範囲表示は .xlsx / .xlsm のみ対応しています。")
    if bool(args.sheet) != bool(args.cell_range):
        raise SystemExit("範囲表示では --sheet と --range を両方指定してください。")
    selected_modes = sum(
        [args.index_only, args.find is not None, bool(args.sheet and args.cell_range)]
    )
    if selected_modes > 1:
        raise SystemExit("--index-only、--find、--sheet/--range は同時指定できません。")
    if args.index_only:
        print_index(excel_path, include_hidden=args.include_hidden)
        return
    if args.find is not None:
        find_values(
            excel_path,
            args.find,
            include_hidden=args.include_hidden,
            max_results=args.max_results,
            context_rows=args.context_rows,
        )
        return
    if args.sheet and args.cell_range:
        print_range(
            excel_path,
            args.sheet,
            args.cell_range,
            data_only=args.data_only,
            max_cells=args.max_cells,
            compact=args.compact,
        )
        return

    # 出力先は output/ に固定する。入力(input/)と結果を混ぜない。
    output_dir = Path("output")

    xls = pd.ExcelFile(excel_path)
    sheet_names = xls.sheet_names if args.include_hidden else _visible_sheet_names(xls)
    if args.csv_sheet:
        if args.csv_sheet not in sheet_names:
            raise SystemExit(f"シートがありません: {args.csv_sheet}(あるのは {sheet_names})")
        sheet_names = [args.csv_sheet]
    print(f"ファイル: {excel_path}")
    print("検出されたシート:")
    for i, sheet_name in enumerate(sheet_names, 1):
        print(f"  {i}. {sheet_name}")

    for sheet_name in sheet_names:
        df = pd.read_excel(excel_path, sheet_name=sheet_name)
        print(f"\n--- シート: {sheet_name} ({df.shape[0]} 行 × {df.shape[1]} 列) ---")
        print(df.head(args.rows))

        if args.no_csv:
            continue
        output_dir.mkdir(parents=True, exist_ok=True)
        csv_path = output_dir / f"{excel_path.stem}_{_safe_stem(sheet_name)}.csv"
        df.to_csv(csv_path, index=False, encoding="utf-8-sig")
        print(f"CSV出力完了: {csv_path}")


if __name__ == "__main__":
    main()
