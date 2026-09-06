"""
Wordファイル(.docx)を読み書きするツールの雛形。

使い方(リポジトリ直下で。対象ファイルは input/ に置いてもらう):
  本文の表示:   ...python.exe tools\\office_word.py show "input\\文書.docx"
  文字列の置換: ...python.exe tools\\office_word.py replace "input\\文書.docx" --old "旧文言" --new "新文言"
  新規作成:     ...python.exe tools\\office_word.py new "output\\新文書.docx" --from "input\\原稿.txt"

(...python.exe = text-processing-bridge\\.venv\\Scripts\\python.exe)

新規作成の原稿は普通のテキストでよい。行頭の記号だけ解釈する:
  「# 」=大見出し / 「## 」=小見出し / 「- 」「・」=箇条書き / それ以外=本文

守っている約束:
- input/ のファイルは絶対に上書きしない。置換結果は output/ へ
  「元名_updated.docx」として保存し、新規作成は保存先が既にあると中止する
- 置換はラン(書式のかたまり)単位で行い、フォント・太字などの書式を保つ。
  ただし置換対象の語が複数のランへ泣き別れしていると置換できないことがある。
  その場合は件数報告が0になるので、表示(show)で前後の文を確認して語を選び直す
"""

from __future__ import annotations

import argparse
import sys
from pathlib import Path

from docx import Document

# PowerShell経由(既定cp932)だと日本語出力が化けるため、常にUTF-8で出す
sys.stdout.reconfigure(encoding="utf-8", errors="replace")
sys.stderr.reconfigure(encoding="utf-8", errors="replace")


def _iter_paragraphs(document):
    """本文と表の中の段落をまとめて辿る。表の中の文言も置換対象にするため。"""
    for paragraph in document.paragraphs:
        yield paragraph
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    yield paragraph


def cmd_show(path: Path) -> None:
    document = Document(path)
    for i, paragraph in enumerate(_iter_paragraphs(document), 1):
        text = paragraph.text.strip()
        if text:
            print(f"{i:4}: {text}")


def cmd_replace(path: Path, old: str, new: str) -> None:
    document = Document(path)
    count = 0
    for paragraph in _iter_paragraphs(document):
        for run in paragraph.runs:
            if old in run.text:
                run.text = run.text.replace(old, new)
                count += 1

    if count == 0:
        print(f"「{old}」は見つかりませんでした(書式の切れ目で泣き別れている可能性)")
        return

    output_dir = Path("output")
    output_dir.mkdir(parents=True, exist_ok=True)
    target = output_dir / f"{path.stem}_updated{path.suffix}"
    document.save(target)
    print(f"{count}箇所を置換しました")
    print(f"保存しました(元ファイルは無変更): {target}")


def cmd_new(target: Path, source_text: str) -> None:
    from docx.shared import Pt

    if target.exists():
        raise SystemExit(f"保存先が既にあります(上書きしない約束): {target}")

    document = Document()
    normal = document.styles["Normal"]
    normal.font.name = "Yu Gothic"
    normal.font.size = Pt(10.5)

    for line in source_text.replace("\r\n", "\n").split("\n"):
        if line.startswith("## "):
            document.add_heading(line[3:].strip(), level=2)
        elif line.startswith("# "):
            document.add_heading(line[2:].strip(), level=1)
        elif line.startswith(("- ", "・")):
            document.add_paragraph(line.lstrip("-・ 　"), style="List Bullet")
        else:
            document.add_paragraph(line)

    target.parent.mkdir(parents=True, exist_ok=True)
    document.save(target)
    print(f"作成しました: {target}")


def main() -> None:
    parser = argparse.ArgumentParser(description="Wordの表示・置換・新規作成(既存ファイルは上書きしない)")
    sub = parser.add_subparsers(dest="command", required=True)

    p_show = sub.add_parser("show", help="本文(表の中身を含む)を表示する")
    p_show.add_argument("path")

    p_replace = sub.add_parser("replace", help="文字列を置換して別名保存する")
    p_replace.add_argument("path")
    p_replace.add_argument("--old", required=True, help="置換前の文字列")
    p_replace.add_argument("--new", required=True, help="置換後の文字列")

    p_new = sub.add_parser("new", help="テキスト原稿からdocxを新規作成する")
    p_new.add_argument("path", help="保存先(例 Exsel\\新文書.docx)")
    p_new.add_argument(
        "--from", dest="source", required=True,
        help="原稿テキストファイル(UTF-8)",
    )

    args = parser.parse_args()
    path = Path(args.path)

    if args.command == "new":
        source = Path(args.source)
        if not source.exists():
            raise SystemExit(f"原稿ファイルが見つかりません: {source}")
        cmd_new(path, source.read_text(encoding="utf-8-sig"))
        return

    if not path.exists():
        raise SystemExit(f"ファイルが見つかりません: {path}")

    if args.command == "show":
        cmd_show(path)
    elif args.command == "replace":
        cmd_replace(path, args.old, args.new)


if __name__ == "__main__":
    main()
