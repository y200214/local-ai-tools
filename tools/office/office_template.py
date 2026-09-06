r"""
Officeで見た目を編集できるExcel／Wordテンプレートへ値を差し込む試作ツール。

テンプレート内に次の目印を書く。目印のフォント・色・配置は生成後も維持される。

  {{field:occupancy_rate}}
  {{threshold:occupancy_rate|>=|90|基準を超えています|基準内です}}

使い方:
  ...python.exe tools\office_template.py demo
  ...python.exe tools\office_template.py render TEMPLATE --values VALUES.json --out OUTPUT
  ...python.exe tools\office_template.py minutes TEMPLATE.docx --document DOCUMENT.json --out OUTPUT.docx

議事録テンプレートの書式をWordで編集して本番へ反映する流れ:
  1. check   TEMPLATE.docx            編集後に壊れていないか確認(行の削除などを検出)
  2. preview TEMPLATE.docx            サンプル議事録で書式の見え方を確認
  3. deploy  TEMPLATE.docx            checkを通ったものだけOpen WebUIコンテナへ反映

VALUES.json はLLM処理後の構造化データを想定したJSONオブジェクト。fieldでは
入れ子のキーも ``header.title`` のようなドット区切りで参照できる。

入力テンプレートは上書きしない。Excelは対象セルのXMLだけを置き換えるため、
既存の書式・数式キャッシュ・マクロを保持する。
"""

from __future__ import annotations
from repo_paths import ROOT as PROJECT_ROOT

import argparse
import json
import operator
import re
import sys
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Iterator

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

if __package__:
    from .office_excel import _load, _save_surgical_copy_batch
else:
    from .office_excel import _load, _save_surgical_copy_batch


TOKEN_RE = re.compile(r"\{\{\s*(.+?)\s*\}\}", re.DOTALL)
COMPARISONS = {
    ">": operator.gt,
    ">=": operator.ge,
    "<": operator.lt,
    "<=": operator.le,
    "==": operator.eq,
    "!=": operator.ne,
}
SUPPORTED_SUFFIXES = {".xlsx", ".xlsm", ".docx"}


class TemplateError(ValueError):
    """テンプレートの目印または入力値が不正なときの説明可能なエラー。"""


def _lookup(values: dict[str, Any], key: str) -> Any:
    """完全一致を優先し、無ければドット区切りで入れ子を辿る。"""
    key = key.strip()
    if key in values:
        return values[key]

    current: Any = values
    for part in key.split("."):
        if not isinstance(current, dict) or part not in current:
            raise TemplateError(f"入力値に項目がありません: {key}")
        current = current[part]
    return current


def _scalar(value: Any, key: str) -> str | int | float:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "はい" if value else "いいえ"
    if isinstance(value, (str, int, float)):
        return value
    raise TemplateError(f"{key} は文字列または数値ではありません")


def _decimal(value: Any, key: str) -> Decimal:
    if isinstance(value, bool):
        raise TemplateError(f"{key} は数値ではありません")
    normalized = str(value).strip().replace(",", "")
    if normalized.endswith("%"):
        normalized = normalized[:-1].strip()
    try:
        return Decimal(normalized)
    except InvalidOperation as error:
        raise TemplateError(f"{key} は数値ではありません: {value}") from error


def evaluate_expression(expression: str, values: dict[str, Any]) -> str | int | float:
    """安全な2種類の目印(field/threshold)だけを評価する。"""
    expression = expression.strip()
    if expression.startswith("field:"):
        key = expression.removeprefix("field:").strip()
        if not key:
            raise TemplateError("fieldの項目名が空です")
        return _scalar(_lookup(values, key), key)

    if expression.startswith("threshold:"):
        parts = expression.split("|")
        if len(parts) != 5:
            raise TemplateError(
                "thresholdは「項目|比較記号|基準値|該当時コメント|非該当時コメント」"
                "の順で指定してください"
            )
        key = parts[0].removeprefix("threshold:").strip()
        comparison = parts[1].strip()
        threshold_text = parts[2].strip()
        if not key:
            raise TemplateError("thresholdの項目名が空です")
        if comparison not in COMPARISONS:
            raise TemplateError(f"未対応の比較記号です: {comparison}")
        actual = _decimal(_lookup(values, key), key)
        threshold = _decimal(threshold_text, "基準値")
        selected = parts[3] if COMPARISONS[comparison](actual, threshold) else parts[4]
        return selected.replace("{value}", str(actual)).replace(
            "{threshold}", str(threshold)
        )

    raise TemplateError(f"未対応の目印です: {expression}")


def render_text(text: str, values: dict[str, Any]) -> str:
    """文章内の全目印を文字列として差し替える。"""
    return TOKEN_RE.sub(
        lambda match: str(evaluate_expression(match.group(1), values)), text
    )


def _excel_value(text: str, values: dict[str, Any]) -> str | int | float:
    """セル全体がfieldなら数値型を保ち、それ以外は文章として置換する。"""
    exact = TOKEN_RE.fullmatch(text)
    if exact and exact.group(1).strip().startswith("field:"):
        return evaluate_expression(exact.group(1), values)
    return render_text(text, values)


def render_excel_template(
    template: Path, values: dict[str, Any], target: Path
) -> int:
    workbook = _load(template)
    updates: list[tuple[str, str, str | int | float]] = []
    try:
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows():
                for cell in row:
                    if not isinstance(cell.value, str) or "{{" not in cell.value:
                        continue
                    location = f"{sheet.title}!{cell.coordinate}"
                    if cell.data_type == "f":
                        raise TemplateError(
                            f"{location}: 数式の中には目印を置けません"
                        )
                    try:
                        rendered = _excel_value(cell.value, values)
                    except TemplateError as error:
                        raise TemplateError(f"{location}: {error}") from error
                    updates.append((sheet.title, cell.coordinate, rendered))
    finally:
        workbook.close()

    if not updates:
        raise TemplateError("Excel内に差し込み目印が見つかりません")
    _save_surgical_copy_batch(template, updates, target=target)
    return len(updates)


def _iter_container_paragraphs(container) -> Iterator:
    """本文・表・入れ子の表にある段落を辿る。"""
    yield from container.paragraphs
    for table in container.tables:
        seen_cells: set[Any] = set()
        for row in table.rows:
            for cell in row.cells:
                if cell._tc in seen_cells:
                    continue
                seen_cells.add(cell._tc)
                yield from _iter_container_paragraphs(cell)


def _iter_document_paragraphs(document: Document) -> Iterator:
    """本文に加えてヘッダー／フッターも重複なく辿る。"""
    seen: set[Any] = set()
    containers: list[Any] = [document]
    for section in document.sections:
        containers.extend(
            [
                section.header,
                section.footer,
                section.first_page_header,
                section.first_page_footer,
                section.even_page_header,
                section.even_page_footer,
            ]
        )
    for container in containers:
        for paragraph in _iter_container_paragraphs(container):
            if paragraph._p not in seen:
                seen.add(paragraph._p)
                yield paragraph


def _replace_run_span(paragraph, start: int, end: int, replacement: str) -> None:
    """複数ランに分断された目印を、先頭ランの書式を残して置換する。"""
    positions: list[tuple[int, int]] = []
    position = 0
    for run in paragraph.runs:
        positions.append((position, position + len(run.text)))
        position += len(run.text)

    first_index = next(
        index for index, (_, run_end) in enumerate(positions) if start < run_end
    )
    last_index = next(
        index for index, (run_start, run_end) in enumerate(positions)
        if run_start < end <= run_end
    )
    first = paragraph.runs[first_index]
    last = paragraph.runs[last_index]
    first_start = positions[first_index][0]
    last_start = positions[last_index][0]
    prefix = first.text[: start - first_start]
    suffix = last.text[end - last_start :]

    if first_index == last_index:
        first.text = prefix + replacement + suffix
        return

    first.text = prefix + replacement
    for index in range(first_index + 1, last_index):
        paragraph.runs[index].text = ""
    last.text = suffix


def _render_word_paragraph(paragraph, values: dict[str, Any]) -> int:
    text = "".join(run.text for run in paragraph.runs)
    matches = list(TOKEN_RE.finditer(text))
    for match in matches:
        try:
            evaluate_expression(match.group(1), values)
        except TemplateError as error:
            excerpt = text[:80].replace("\n", " ")
            raise TemplateError(f"Word「{excerpt}」: {error}") from error
    for match in reversed(matches):
        replacement = str(evaluate_expression(match.group(1), values))
        _replace_run_span(paragraph, match.start(), match.end(), replacement)
    return len(matches)


def render_word_template(
    template: Path, values: dict[str, Any], target: Path
) -> int:
    document = Document(template)
    count = sum(
        _render_word_paragraph(paragraph, values)
        for paragraph in _iter_document_paragraphs(document)
    )
    if not count:
        raise TemplateError("Word内に差し込み目印が見つかりません")
    document.save(target)
    return count


def _validated_target(template: Path, target: Path, force: bool) -> Path:
    if template.resolve() == target.resolve():
        raise TemplateError("入力テンプレート自身は上書きできません")
    if target.suffix.lower() != template.suffix.lower():
        raise TemplateError("出力ファイルの拡張子はテンプレートと揃えてください")
    if target.exists() and not force:
        raise TemplateError(f"出力先が既にあります: {target}（上書きは--force）")
    target.parent.mkdir(parents=True, exist_ok=True)
    return target


def render_template(
    template: Path,
    values: dict[str, Any],
    target: Path,
    *,
    force: bool = False,
) -> int:
    if not template.exists():
        raise TemplateError(f"テンプレートが見つかりません: {template}")
    suffix = template.suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        raise TemplateError(f"未対応の形式です: {suffix}")
    target = _validated_target(template, target, force)
    if suffix in {".xlsx", ".xlsm"}:
        return render_excel_template(template, values, target)
    return render_word_template(template, values, target)


def load_minutes_document(path: Path) -> dict[str, Any]:
    """構造化議事録APIの応答JSONまたはdocument本体を読み込む。"""
    try:
        payload = json.loads(path.read_text(encoding="utf-8-sig"))
    except (OSError, json.JSONDecodeError) as error:
        raise TemplateError(f"議事録JSONを読めません: {path}: {error}") from error
    if not isinstance(payload, dict):
        raise TemplateError("議事録JSONの最上位はオブジェクトにしてください")
    document = payload.get("document", payload)
    if not isinstance(document, dict) or not isinstance(document.get("議題"), list):
        raise TemplateError("議事録JSONにdocument.議題がありません")
    return document


def _load_renderer():
    """実運用の議事録レンダラー(ブリッジ側)を読み込む。"""
    bridge_root = PROJECT_ROOT / "text-processing-bridge"
    if str(bridge_root) not in sys.path:
        sys.path.insert(0, str(bridge_root))
    from app import render_minutes

    return render_minutes


def render_minutes_preview(
    template: Path,
    document: dict[str, Any],
    target: Path,
    *,
    force: bool = False,
) -> None:
    """実運用と同じrender_minutes.pyで実物Wordテンプレートをプレビューする。"""
    if not template.exists():
        raise TemplateError(f"テンプレートが見つかりません: {template}")
    if template.suffix.lower() != ".docx":
        raise TemplateError("議事録テンプレートは.docxを指定してください")
    target = _validated_target(template, target, force)

    try:
        rendered = _load_renderer().render(template.read_bytes(), document)
    except Exception as error:
        raise TemplateError(f"議事録テンプレートを描画できません: {error}") from error
    target.write_bytes(rendered)


# ---------------------------------------------------------------------------
# 議事録テンプレートの編集導線(check → preview → deploy)。
# 書式(タイトルのフォント・サイズ・色、表の見た目)はテンプレートdocxから
# そのまま複製されるため、Wordで編集して deploy すれば出力へ反映される。
# 危ないのは行の追加・削除(レンダラーは決まった行をひな型として使う)。
# それを反映前に check が捕まえる。
# ---------------------------------------------------------------------------

# 本番テンプレートのコンテナ内パス(PipeのMINUTES_TEMPLATE_PATHと同じ)
MINUTES_CONTAINER = "open-webui"
MINUTES_CONTAINER_PATH = "/app/backend/data/templates/minutes_template.docx"

# 書式確認用のサンプル議事録。実在の会議・人名は使わない
SAMPLE_MINUTES_DOCUMENT: dict[str, Any] = {
    "ヘッダ": {
        "タイトル": "書式確認用サンプル委員会議事録",
        "日時": "令和8年8月14日（金）17時00分",
        "場所": "第1会議室",
        "出席者": "委員A、委員B、委員C",
        "書記": "事務局D",
    },
    "議題": [
        {
            "番号": "1",
            "表題": "サンプル議題（書式確認用）",
            "項目": [
                {
                    "枝番": "①",
                    "表題": "経営指標の報告",
                    "資料": "資料1",
                    "担当": "事務局",
                    "説明": "これは書式確認用のサンプル本文です。"
                    "タイトルや表のフォント・サイズ・色の変更がここへ反映されます。",
                    "質疑応答": [
                        {"発言者": "委員A", "内容": "サンプルの質問です。"},
                        {"発言者": "", "内容": "発言者不明(赤マーカー)のサンプルです。"},
                    ],
                },
                {
                    "枝番": "②",
                    "表題": "運営状況の報告",
                    "資料": "資料2",
                    "担当": "総務課",
                    "説明": "2項目目のサンプルです。",
                    "質疑応答": [],
                },
            ],
        },
        {
            "番号": "2",
            "表題": "推定議題のサンプル",
            "推定": True,
            "項目": [
                {
                    "枝番": "①",
                    "表題": "規程の改正について",
                    "資料": "",
                    "担当": "",
                    "説明": "推定議題(黄マーカー)のサンプルです。",
                    "質疑応答": [],
                }
            ],
        },
    ],
    "その他": [{"発言者": "委員B", "内容": "その他欄のサンプル発言です。"}],
    "次回開催": "令和8年9月11日（金）17時00分　第1会議室",
}


def check_minutes_template(template: Path) -> list[str]:
    """
    テンプレートが実運用レンダラーで使えるかを確かめ、問題を平文で返す。

    空リスト = 問題なし。書式の変更(フォント・サイズ・色)はここでは
    一切問題にしない。構造(行・表・タイトル行)だけを見る。
    """
    problems: list[str] = []
    try:
        document = Document(template)
    except Exception as error:
        return [f"Wordファイルとして開けません: {error}"]

    renderer = _load_renderer()
    required_rows = renderer.PROTO_NEXT + 1
    if len(document.tables) < 2:
        problems.append(
            f"表が{len(document.tables)}個しかありません。"
            "1つ目がヘッダ、2つ目が本文という表構成が必要です"
        )
    elif len(document.tables[1].rows) < required_rows:
        problems.append(
            f"本文の表(2つ目)が{len(document.tables[1].rows)}行しかありません。"
            f"レンダラーがひな型として使う行を含めて{required_rows}行必要です。"
            "行を削除した場合は元のテンプレートから戻してください"
            "(書式だけの変更なら行数は変わりません)"
        )
    if not any(p.text.strip().endswith("議事録") for p in document.paragraphs):
        problems.append(
            "タイトル行(末尾が「議事録」で終わる段落)が見つかりません。"
            "タイトルの書式はこの行から取られます"
        )
    if problems:
        return problems

    try:
        renderer.render(template.read_bytes(), SAMPLE_MINUTES_DOCUMENT)
    except Exception as error:
        problems.append(f"サンプル議事録の試し描画に失敗しました: {error}")
    return problems


def deploy_minutes_template(template: Path, container: str = MINUTES_CONTAINER) -> None:
    """checkを通ったテンプレートだけを、Open WebUIコンテナへ反映する。"""
    problems = check_minutes_template(template)
    if problems:
        raise TemplateError("反映を中止しました:\n  - " + "\n  - ".join(problems))

    import hashlib
    import subprocess
    import tempfile

    no_window = getattr(subprocess, "CREATE_NO_WINDOW", 0)

    def run_docker(*arguments: str) -> None:
        result = subprocess.run(
            ["docker", *arguments],
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
            timeout=120,
            creationflags=no_window,
        )
        if result.returncode:
            raise TemplateError(
                f"docker {arguments[0]} に失敗しました: {(result.stderr or result.stdout).strip()[-500:]}"
            )

    run_docker("cp", str(template), f"{container}:{MINUTES_CONTAINER_PATH}")

    # 反映後のファイルを取り出してハッシュで確認する(置けたつもりを防ぐ)
    with tempfile.TemporaryDirectory() as verify_dir:
        pulled = Path(verify_dir) / "verify.docx"
        run_docker("cp", f"{container}:{MINUTES_CONTAINER_PATH}", str(pulled))
        local = hashlib.sha256(template.read_bytes()).hexdigest()
        remote = hashlib.sha256(pulled.read_bytes()).hexdigest()
        if local != remote:
            raise TemplateError("反映後の検証に失敗しました(コンテナ内のファイルが一致しません)")


def _make_demo_excel(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "レポート"
    sheet.column_dimensions["A"].width = 22
    sheet.column_dimensions["B"].width = 58
    sheet.row_dimensions[4].height = 48
    sheet.row_dimensions[6].height = 62

    sheet.merge_cells("A1:B1")
    sheet["A1"] = "GUI編集テンプレート・Excel試作"
    sheet["A1"].font = Font(name="Yu Gothic", size=18, bold=True, color="FFFFFF")
    sheet["A1"].fill = PatternFill("solid", fgColor="2F5597")
    sheet["A1"].alignment = Alignment(horizontal="center")

    sheet["A3"] = "病床稼働率"
    sheet["B3"] = "{{field:occupancy_rate}}"
    sheet["B3"].font = Font(name="Yu Gothic", size=18, bold=True, color="C65911")
    sheet["B3"].fill = PatternFill("solid", fgColor="FFF2CC")
    sheet["B3"].number_format = '0.0"%"'

    sheet["A4"] = "数値コメント"
    sheet["B4"] = (
        "{{threshold:occupancy_rate|>=|90|90%以上のため、稼働状況の確認が必要です。"
        "|90%未満で推移しています。}}"
    )
    sheet["B4"].font = Font(name="Yu Gothic", size=12, color="1F4E78")
    sheet["B4"].fill = PatternFill("solid", fgColor="DDEBF7")
    sheet["B4"].alignment = Alignment(wrap_text=True, vertical="center")

    sheet["A6"] = "LLMコメント"
    sheet["B6"] = "{{field:llm_comment}}"
    sheet["B6"].font = Font(name="Yu Gothic", size=12, color="375623")
    sheet["B6"].fill = PatternFill("solid", fgColor="E2F0D9")
    sheet["B6"].alignment = Alignment(wrap_text=True, vertical="center")

    sheet["A8"] = "黄色・青・緑のセルの書式や配置をExcelで変更し、保存後に再生成してください。"
    sheet.merge_cells("A8:B8")
    sheet["A8"].font = Font(name="Yu Gothic", size=10, italic=True, color="666666")
    workbook.save(path)
    workbook.close()


def _make_demo_word(path: Path) -> None:
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run("GUI編集テンプレート・Word試作")
    title_run.bold = True
    title_run.font.name = "Yu Gothic"
    title_run.font.size = Pt(18)
    title_run.font.color.rgb = RGBColor(47, 85, 151)

    table = document.add_table(rows=3, cols=2)
    table.style = "Table Grid"
    entries = [
        ("病床稼働率", "{{field:occupancy_rate}}", RGBColor(198, 89, 17), 16),
        (
            "数値コメント",
            "{{threshold:occupancy_rate|>=|90|90%以上のため、稼働状況の確認が必要です。|90%未満で推移しています。}}",
            RGBColor(31, 78, 121),
            11,
        ),
        ("LLMコメント", "{{field:llm_comment}}", RGBColor(55, 86, 35), 11),
    ]
    for row, (label, token, color, size) in zip(table.rows, entries):
        row.cells[0].text = label
        value_paragraph = row.cells[1].paragraphs[0]
        value_paragraph.clear()
        run = value_paragraph.add_run(token)
        run.font.name = "Yu Gothic"
        run.font.size = Pt(size)
        run.font.color.rgb = color
        if label == "病床稼働率":
            run.bold = True

    note = document.add_paragraph(
        "右側の欄のフォント、文字サイズ、色、表の幅などをWordで変更し、保存後に再生成してください。"
    )
    note.runs[0].font.name = "Yu Gothic"
    note.runs[0].font.size = Pt(9)
    note.runs[0].font.color.rgb = RGBColor(102, 102, 102)
    document.save(path)


def create_demo(directory: Path, *, force: bool = False) -> list[Path]:
    directory.mkdir(parents=True, exist_ok=True)
    excel_template = directory / "demo_template.xlsx"
    word_template = directory / "demo_template.docx"
    values_path = directory / "values.json"
    excel_result = directory / "demo_result.xlsx"
    word_result = directory / "demo_result.docx"
    readme_path = directory / "README.txt"
    paths = [
        excel_template,
        word_template,
        values_path,
        excel_result,
        word_result,
        readme_path,
    ]
    existing = [path for path in paths if path.exists()]
    if existing and not force:
        raise TemplateError(
            "デモ出力が既にあります（上書きは--force）: "
            + "、".join(str(path) for path in existing)
        )

    _make_demo_excel(excel_template)
    _make_demo_word(word_template)
    values = {
        "occupancy_rate": 92.4,
        "llm_comment": "前月から上昇しています。要因を確認し、来月も推移を確認します。",
    }
    values_path.write_text(
        json.dumps(values, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    render_template(excel_template, values, excel_result, force=True)
    render_template(word_template, values, word_result, force=True)
    readme_path.write_text(
        "Office GUIテンプレート試作\n\n"
        "1. demo_template.xlsx または demo_template.docx をOfficeで開きます。\n"
        "2. 目印が入った欄のフォント・サイズ・色・配置を変更して保存します。\n"
        "3. リポジトリ直下で次のように再生成します。\n\n"
        "text-processing-bridge\\.venv\\Scripts\\python.exe "
        "tools\\office_template.py render <テンプレート> "
        "--values <values.json> --out <新しい出力先>\n\n"
        "demo_result.xlsx / demo_result.docx は変更前の生成例です。\n"
        "threshold目印の90やコメント文もOffice上で変更できます。\n",
        encoding="utf-8",
    )
    return paths


def _load_values(path: Path) -> dict[str, Any]:
    try:
        values = json.loads(path.read_text(encoding="utf-8-sig"))
    except (OSError, json.JSONDecodeError) as error:
        raise TemplateError(f"入力JSONを読めません: {path}: {error}") from error
    if not isinstance(values, dict):
        raise TemplateError("入力JSONの最上位はオブジェクトにしてください")
    return values


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Officeで書式を編集できるテンプレートへの差し込み試作"
    )
    sub = parser.add_subparsers(dest="command", required=True)

    demo = sub.add_parser("demo", help="Excel／Wordの編集用デモ一式を作る")
    demo.add_argument(
        "--directory", default="output\\office_template_demo", help="デモの保存先"
    )
    demo.add_argument("--force", action="store_true", help="既存のデモ一式を上書きする")

    render = sub.add_parser("render", help="テンプレートへJSONの値を差し込む")
    render.add_argument("template", help=".xlsx/.xlsm/.docxテンプレート")
    render.add_argument("--values", required=True, help="差し込むJSONオブジェクト")
    render.add_argument("--out", help="出力先（省略時はoutput/<名前>_rendered）")
    render.add_argument("--force", action="store_true", help="既存の出力を上書きする")

    minutes = sub.add_parser(
        "minutes", help="実運用のレンダラーで議事録Wordテンプレートを試す"
    )
    minutes.add_argument("template", help="実物の議事録.docxテンプレート")
    minutes.add_argument("--document", required=True, help="構造化議事録のJSON")
    minutes.add_argument("--out", required=True, help="プレビューdocxの出力先")
    minutes.add_argument("--force", action="store_true", help="既存の出力を上書きする")

    check = sub.add_parser(
        "check", help="編集した議事録テンプレートが壊れていないか確認する"
    )
    check.add_argument("template", help="確認する議事録.docxテンプレート")

    preview = sub.add_parser(
        "preview", help="サンプル議事録でテンプレートの書式の見え方を確認する"
    )
    preview.add_argument("template", help="議事録.docxテンプレート")
    preview.add_argument(
        "--out",
        default=str(Path("output") / "minutes_template_preview.docx"),
        help="プレビューdocxの出力先",
    )
    preview.add_argument("--force", action="store_true", help="既存の出力を上書きする")

    deploy = sub.add_parser(
        "deploy", help="checkを通った議事録テンプレートをOpen WebUIへ反映する"
    )
    deploy.add_argument("template", help="反映する議事録.docxテンプレート")
    deploy.add_argument(
        "--container", default=MINUTES_CONTAINER, help="Open WebUIのコンテナ名"
    )

    args = parser.parse_args()
    try:
        if args.command == "demo":
            created = create_demo(Path(args.directory), force=args.force)
            print("デモを作成しました:")
            for path in created:
                print(f"  {path}")
            return

        if args.command == "minutes":
            render_minutes_preview(
                Path(args.template),
                load_minutes_document(Path(args.document)),
                Path(args.out),
                force=args.force,
            )
            print(f"保存しました（テンプレートは無変更）: {args.out}")
            return

        if args.command == "check":
            problems = check_minutes_template(Path(args.template))
            if problems:
                raise TemplateError(
                    "このままでは使えません:\n  - " + "\n  - ".join(problems)
                )
            print("問題ありません。preview で見え方を確認し、deploy で反映できます")
            return

        if args.command == "preview":
            render_minutes_preview(
                Path(args.template),
                SAMPLE_MINUTES_DOCUMENT,
                Path(args.out),
                force=args.force,
            )
            print(f"サンプル議事録を保存しました（テンプレートは無変更）: {args.out}")
            print("Wordで開いて、タイトル・表・本文の書式を確認してください")
            return

        if args.command == "deploy":
            deploy_minutes_template(Path(args.template), args.container)
            print(f"反映しました: {args.template} → {args.container}:{MINUTES_CONTAINER_PATH}")
            print("次の議事録作成から適用されます(コンテナの再起動は不要)")
            return

        template = Path(args.template)
        target = (
            Path(args.out)
            if args.out
            else Path("output") / f"{template.stem}_rendered{template.suffix}"
        )
        count = render_template(
            template,
            _load_values(Path(args.values)),
            target,
            force=args.force,
        )
        print(f"{count}箇所を差し込みました")
        print(f"保存しました（テンプレートは無変更）: {target}")
    except TemplateError as error:
        raise SystemExit(str(error)) from error


if __name__ == "__main__":
    main()
